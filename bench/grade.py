#!/usr/bin/env python3
"""Deterministic graders for desk-track tasks.

Usage: grade.py <task_dir> <workspace_dir>  -> prints JSON {passed, checks:[...]}

A task's task.yaml lists `checks`. Every check has `type`, optional `name`,
optional `required` (default true). A task passes when every required check
passes. Reference files live in <task_dir>/reference and are never visible to
the harness. Output paths are globs relative to the workspace; if a plain name
does not match at the top level it is searched recursively.
"""
from __future__ import annotations
import glob, json, os, re, sys, tempfile, importlib.util, io, contextlib
os.environ.setdefault('TQDM_DISABLE', '1')
import pandas as pd
import yaml

# ---------- helpers ----------

def _norm_col(c: str) -> str:
    return re.sub(r'[^a-z0-9]+', '_', str(c).strip().lower()).strip('_')

def find_file(ws: str, pattern: str) -> str | None:
    hits = sorted(glob.glob(os.path.join(ws, pattern)))
    if not hits:
        hits = sorted(glob.glob(os.path.join(ws, '**', pattern), recursive=True))
    hits = [h for h in hits if os.path.isfile(h) and '/.proto' not in h and '/.codex' not in h]
    if not hits:
        # file names are matched case-insensitively as a fallback: an agent that saves Memo.md for memo.md
        # delivered the file; the checks on its contents decide the rest
        import fnmatch
        pat = pattern.lower()
        for base, dirs, files in os.walk(ws):
            dirs[:] = [d for d in dirs if d not in ('.proto', '.codex', 'node_modules', '.git')]
            for f in files:
                rel = os.path.relpath(os.path.join(base, f), ws).lower()
                if fnmatch.fnmatch(rel, pat) or ('/' not in pat and fnmatch.fnmatch(f.lower(), pat)):
                    hits.append(os.path.join(base, f))
        hits.sort()
    return hits[0] if hits else None

# Abbreviations that end in a period but do not end a sentence.
_ABBREV = re.compile(r'\b(?:a\.m|p\.m|e\.g|i\.e|etc|vs|approx|dept|mr|mrs|ms|dr|st|ave|rd|blvd|inc|co|corp|ltd|llc|jr|sr|'
                     r'jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec|mon|tue|tues|wed|thu|thur|thurs|fri|sat|sun)\.$', re.I)

def split_sentences(t: str) -> list[str]:
    """Line breaks always end a sentence; within a line, a period ends one unless it closes an abbreviation."""
    out: list[str] = []
    for line in re.split(r'\n+', t):
        parts: list[str] = []
        for x in re.split(r'(?<=[.!?])\s+', line):
            x = x.strip()
            if not x: continue
            if parts and _ABBREV.search(parts[-1]):
                parts[-1] = parts[-1] + ' ' + x
            else:
                parts.append(x)
        out.extend(parts)
    return out

def read_table(path: str, sheet: str | None = None) -> pd.DataFrame:
    if path.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    else:
        # Normalise line endings first: CRLF and the CR-CR-LF that csv.writer
        # emits without newline='' are both common in real business files, and
        # pandas' C parser treats a lone CR as a record break.
        raw = open(path, 'rb').read().decode('utf-8', errors='replace')
        text = raw.replace('\r\r\n', '\n').replace('\r\n', '\n')
        try:
            df = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False)
        except pd.errors.ParserError:
            # A malformed CSV (ragged rows, unquoted commas) is a real defect, but
            # the run should fail on the checks that see it, not on a parser crash.
            df = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False,
                             engine='python', on_bad_lines='skip')
            df.attrs['malformed'] = True
    df.columns = [_norm_col(c) for c in df.columns]
    return df.fillna('')

def _malformed_note(df: pd.DataFrame) -> str:
    return ' [malformed CSV: ragged rows skipped]' if df.attrs.get('malformed') else ''

def col(df: pd.DataFrame, name: str) -> pd.Series:
    n = _norm_col(name)
    if n not in df.columns:
        raise KeyError(f'column {name!r} not found; have {list(df.columns)}')
    return df[n].astype(str)

def normalize(v: str, ops: list[str]) -> str:
    s = str(v)
    for op in ops or []:
        if op == 'lower': s = s.lower()
        elif op == 'strip': s = s.strip()
        elif op == 'digits': s = re.sub(r'\D', '', s)
        elif op == 'alnum': s = re.sub(r'[^a-z0-9]', '', s.lower())
        elif op == 'money': s = re.sub(r'[^0-9.\-]', '', s)
    return s

def to_num(v) -> float | None:
    try:
        return float(re.sub(r'[^0-9.\-eE]', '', str(v)))
    except Exception:
        return None

# A spreadsheet cell or CSV field that reads as one number: optional currency symbol or code, sign or
# parentheses, thousands separators, decimals, percent. Dates, ids and prose are not numbers
# ("10 Jan 2026" is not 102026, "cus_4fvssavh" is not 4).
_CCY = r'(?:USD|EUR|GBP|CAD|AUD|NZD|INR|JPY|CHF|MXN|SEK|NOK|DKK)'
_NUMERIC_TEXT = re.compile(r'^\s*(?:' + _CCY + r'\s*)?[\(\-\u2212+]?\s*[$\u20ac\u00a3]?\s*\d{1,3}(?:[,\s]\d{3})*(?:\.\d+)?\s*\)?\s*%?\s*(?:' + _CCY + r')?\s*$|^\s*[\(\-\u2212+]?\s*[$\u20ac\u00a3]?\s*\d+(?:\.\d+)?\s*\)?\s*%?\s*$')

def cell_num(v) -> float | None:
    """Strict number for grid scans: real numbers pass through; text must look like a single figure."""
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    t = str(v)
    if not _NUMERIC_TEXT.match(t): return None
    neg = '(' in t or '-' in t or '\u2212' in t
    x = to_num(re.sub(r'[A-Za-z%$\u20ac\u00a3(),\s\u2212\-+]', '', t))
    return None if x is None else (-x if neg else x)

def num_eq(a, b, tol: float) -> bool:
    # accounting negatives "(0.27)" are negative; fall back to the loose parse for text like "12 units"
    x = cell_num(a); x = to_num(a) if x is None else x
    y = cell_num(b); y = to_num(b) if y is None else y
    return x is not None and y is not None and abs(x - y) <= tol

def read_text(path: str) -> str:
    with open(path, encoding='utf-8', errors='replace') as f:
        return f.read()

_recalc_cache: dict[str, str] = {}

def _soffice_recalc(path: str, outdir: str) -> str | None:
    """Recalculate with LibreOffice: on PATH, or through the bench-recalc container
    when BENCH_RECALC_DOCKER_IMAGE is set. Returns the converted file or None."""
    import shutil, subprocess
    src = os.path.join(outdir, 'in.xlsx')
    # Builders often write cached results beside formulas. LibreOffice trusts those on
    # load and recalculates only the cells without one, which can leave a total that no
    # longer matches its own recalculated inputs. Re-saving through openpyxl drops every
    # cached value so LibreOffice must compute the whole workbook.
    try:
        from openpyxl import load_workbook
        load_workbook(path).save(src)
    except Exception:
        shutil.copyfile(path, src)
    image = os.environ.get('BENCH_RECALC_DOCKER_IMAGE')
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if image:
        cmd = ['docker', 'run', '--rm', '--user', f'{os.getuid()}:{os.getgid()}', '-e', 'HOME=/tmp',
               '-v', f'{outdir}:/x', image, 'soffice', '-env:UserInstallation=file:///tmp/lo-profile',
               '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', '/x/out', '/x/in.xlsx']
    elif soffice:
        # An explicit profile dir keeps parallel graders from fighting over one lock file.
        cmd = [soffice, f'-env:UserInstallation=file://{outdir}/lo-profile', '--headless', '--calc',
               '--convert-to', 'xlsx', '--outdir', os.path.join(outdir, 'out'), src]
    else:
        return None
    # Under load (parallel graders, agent containers) a conversion can time out or fail to
    # start. Retry before giving up: a silent fall-through to the Python engine has produced
    # false verdicts, so when a real engine is configured its failure is reported, not masked.
    out = os.path.join(outdir, 'out', 'in.xlsx')
    import time as _time
    for attempt in range(3):
        try:
            subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240, check=False,
                           env=dict(os.environ, HOME=os.environ.get('HOME', '/tmp')))
        except Exception:
            pass
        if os.path.exists(out):
            return out
        _time.sleep(5 * (attempt + 1))
    raise RuntimeError('LibreOffice recalculation failed three times; check BENCH_RECALC_DOCKER_IMAGE / soffice')

def _formulas_recalc(path: str, outdir: str) -> str | None:
    try:
        import formulas  # type: ignore
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            xl = formulas.ExcelModel().loads(path).finish()
            xl.calculate()
            xl.write(dirpath=outdir)
        cands = [q for q in glob.glob(os.path.join(outdir, '**', '*'), recursive=True) if q.lower().endswith('.xlsx')]
        return cands[0] if cands else None
    except Exception:
        return None

def recalculated_workbook(path: str) -> str:
    """Return a path to a copy of the workbook with formula results computed.
    LibreOffice first (real engine: structured references, EDATE, everything),
    then the `formulas` Python engine, then the original file (cached values only).
    BENCH_RECALC=formulas forces the Python engine."""
    if path in _recalc_cache:
        return _recalc_cache[path]
    outdir = tempfile.mkdtemp(prefix='recalc-'); os.chmod(outdir, 0o777)
    out = None
    if os.environ.get('BENCH_RECALC', 'auto') != 'formulas':
        out = _soffice_recalc(path, outdir)  # None only when no LibreOffice is configured at all
    if out is None:
        out = _formulas_recalc(path, tempfile.mkdtemp(prefix='recalc-py-'))
    out = out or path
    _recalc_cache[path] = out
    return out

# ---------- checks ----------

def c_file_exists(ws, ref, spec):
    p = find_file(ws, spec['path'])
    return bool(p), f'found {os.path.relpath(p, ws)}' if p else f'no file matching {spec["path"]}'

def c_csv_columns(ws, ref, spec):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    df = read_table(p, spec.get('sheet'))
    want = [_norm_col(c) for c in spec['columns']]
    have = list(df.columns)
    if spec.get('exact'):
        ok = have == want
        return ok, f'columns {have}' + ('' if ok else f' != {want}')
    missing = [c for c in want if c not in have]
    return not missing, (f'missing {missing}' if missing else 'all required columns present') + _malformed_note(df)

def c_csv_row_count(ws, ref, spec):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    n = len(read_table(p, spec.get('sheet')))
    if 'equals_ref' in spec:
        want = len(read_table(os.path.join(ref, spec['equals_ref'])))
    else:
        want = int(spec['equals'])
    tol = int(spec.get('tolerance', 0))
    df = read_table(p, spec.get('sheet'))
    return abs(n - want) <= tol, f'{n} rows, expected {want}' + _malformed_note(df)

def c_csv_set_equal(ws, ref, spec):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    ops = spec.get('normalize', ['strip', 'lower'])
    got = {normalize(v, ops) for v in col(read_table(p, spec.get('sheet')), spec['column']) if str(v).strip()}
    rdf = read_table(os.path.join(ref, spec['ref']))
    want = {normalize(v, ops) for v in col(rdf, spec.get('ref_column', spec['column'])) if str(v).strip()}
    extra, missing = sorted(got - want), sorted(want - got)
    ok = not extra and not missing
    return ok, f'{len(got)} got / {len(want)} expected; missing={missing[:8]} extra={extra[:8]}'

def c_csv_values_match(ws, ref, spec):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    ops = spec.get('normalize', ['strip', 'lower'])
    tol = float(spec.get('tolerance', 0.005))
    df = read_table(p, spec.get('sheet'))
    rdf = read_table(os.path.join(ref, spec['ref']))
    key = spec['key']
    keys = key if isinstance(key, list) else [key]   # a list keys on several columns (van + month, invoice + line)
    def keyed(frame):
        cols = [col(frame, k) for k in keys]
        return ['|'.join(normalize(v, ['strip', 'lower']) for v in parts) for parts in zip(*cols)]
    try:
        gmap = {k: row for k, row in zip(keyed(df), df.to_dict('records'))}
        ref_keys = keyed(rdf)
    except KeyError as e:
        return False, str(e)
    total = hits = 0
    wrong: list[str] = []
    must = {'|'.join(normalize(x, ['strip', 'lower']) for x in (k if isinstance(k, list) else str(k).split('|'))) for k in spec.get('must_match_keys', [])}
    must_bad: list[str] = []
    for kk, rrow in zip(ref_keys, rdf.to_dict('records')):
        total += 1
        grow = gmap.get(kk)
        ok = grow is not None
        if ok:
            for c in spec['columns']:
                gv = grow.get(_norm_col(c), '')
                rv = rrow.get(_norm_col(c), '')
                if spec.get('numeric'):
                    # a blank reference cell is matched by a blank deliverable cell, never by a number
                    if str(rv).strip() == '' or str(gv).strip() == '':
                        ok = ok and str(rv).strip() == '' and str(gv).strip() == ''
                    else:
                        ok = ok and num_eq(gv, rv, tol)
                else:
                    ok = ok and normalize(gv, ops) == normalize(rv, ops)
        hits += ok
        if not ok:
            wrong.append(kk)
            if kk in must: must_bad.append(kk)
    acc = hits / total if total else 0.0
    min_acc = float(spec.get('min_accuracy', 1.0))
    passed = acc >= min_acc and not must_bad
    return passed, f'accuracy {acc:.3f} (min {min_acc}); wrong={wrong[:8]}; required-wrong={must_bad}'

def _pick_sheets(wb, name):
    """The named sheet, matched case-insensitively (the Python recalculation engine upper-cases sheet names);
    every sheet when no name is given or none matches."""
    if name:
        hit = [sh for sh in wb.worksheets if sh.title.lower() == str(name).lower()]
        if hit: return hit
    return list(wb.worksheets)

def c_xlsx_has_formulas(ws, ref, spec):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    from openpyxl import load_workbook
    wb = load_workbook(p)  # formulas kept
    sheets = _pick_sheets(wb, spec.get('sheet'))
    n = sum(1 for sh in sheets for row in sh.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith('='))
    want = int(spec.get('min_count', 1))
    return n >= want, f'{n} formula cells (min {want})'

def c_xlsx_value_present(ws, ref, spec):
    """Some numeric cell (after recalculation) equals `expected` within rel_tol.
    If `near_text` is given, that text must appear in the same row OR the same column
    (so both regions-down-the-side and vendors-across-the-top layouts pass)."""
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    from openpyxl import load_workbook
    wb = load_workbook(recalculated_workbook(p), data_only=True)
    expected = float(spec['expected']); rel = float(spec.get('rel_tol', 0.01))
    near = str(spec.get('near_text', '')).lower()
    sheets = _pick_sheets(wb, spec.get('sheet'))
    tol = max(abs(expected) * rel, 0.01)
    for sh in sheets:
        grid = [[c for c in row] for row in sh.iter_rows()]
        if not grid: continue
        ncols = max(len(r) for r in grid)
        col_text = ['' for _ in range(ncols)]
        for r in grid:
            for j, c in enumerate(r):
                if c.value is not None: col_text[j] += ' ' + str(c.value).lower()
        for r in grid:
            row_text = ' '.join(str(c.value).lower() for c in r if c.value is not None)
            for j, c in enumerate(r):
                if c.value is None or isinstance(c.value, bool): continue
                v = cell_num(c.value)
                if v is None or abs(v - expected) > tol: continue
                if not near or near in row_text or near in col_text[j]:
                    return True, f'{sh.title}!{c.coordinate} = {v}'
    return False, f'no cell ≈ {expected}' + (f' in a row or column mentioning {near!r}' if near else '')

def c_xlsx_no_errors(ws, ref, spec):
    """No formula cell evaluates to an Excel error (#NAME?, #VALUE!, #REF!, #DIV/0!, ...) after
    recalculation. A workbook that opens with errors in it fails a business owner's first glance
    even when the checked cells happen to be right. `max_errors` (default 0) tolerates a few."""
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    from openpyxl import load_workbook
    wb_f = load_workbook(p)
    wb_v = load_workbook(recalculated_workbook(p), data_only=True)
    bad = []
    by_lower = {name.lower(): name for name in wb_v.sheetnames}
    for sh in wb_f.worksheets:
        if spec.get('sheet') and sh.title.lower() != str(spec['sheet']).lower(): continue
        vname = by_lower.get(sh.title.lower())
        if vname is None:
            bad.append(f'{sh.title}: sheet missing after recalculation'); continue
        vsh = wb_v[vname]
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    v = vsh[c.coordinate].value
                    if isinstance(v, str) and v.startswith('#'):
                        bad.append(f'{sh.title}!{c.coordinate}={v}')
    allowed = int(spec.get('max_errors', 0))
    return len(bad) <= allowed, (f'{len(bad)} error cells: ' + ', '.join(bad[:8])) if bad else 'no error cells'

def _text_check(ws, spec, mode):
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    t = read_text(p).lower()
    phrases = [str(x).lower() for x in spec['phrases']]
    found = [x for x in phrases if x in t]
    if mode == 'all':
        return len(found) == len(phrases), f'missing {[x for x in phrases if x not in found]}'
    if mode == 'any':
        return bool(found), f'found {found}' if found else f'none of {phrases}'
    return not found, f'forbidden present {found}' if found else 'none present'

def c_text_contains_all(ws, ref, spec): return _text_check(ws, spec, 'all')
def c_text_contains_any(ws, ref, spec): return _text_check(ws, spec, 'any')
def c_text_not_contains(ws, ref, spec): return _text_check(ws, spec, 'none')

def c_text_matches_all(ws, ref, spec):
    """Every regex in `patterns` matches the text (case-insensitive)."""
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    t = read_text(p)
    missing = [pat for pat in spec['patterns'] if not re.search(pat, t, re.I)]
    return not missing, f'missing {missing}' if missing else f'all {len(spec["patterns"])} patterns present'

def c_text_sentence_matches(ws, ref, spec):
    """Some single sentence matches every regex in `all` and none in `none` (case-insensitive).
    Sentence-level matching stops a memo passing on words scattered across unrelated
    sentences ("a gap of $123k between regions" + "North" elsewhere) and lets a negation
    guard reject "rather than a missing-data issue"."""
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    t = read_text(p)
    sentences = split_sentences(t)
    alls = spec.get('all', []); nones = spec.get('none', [])
    for sent in sentences:
        if all(re.search(a, sent, re.I) for a in alls) and not any(re.search(n, sent, re.I) for n in nones):
            return True, f'sentence: {sent[:160]!r}'
    return False, f'no single sentence matches all of {alls}' + (f' without {nones}' if nones else '')

def c_text_numbers_present(ws, ref, spec):
    """Every number in `numbers` appears in the text within rel_tol (commas, currency
    symbols, and parentheses are ignored). Grades facts a memo must state, not the
    judgment call it makes."""
    p = find_file(ws, spec['path'])
    if not p: return False, 'output file missing'
    t = read_text(p)
    found = []
    for m in re.finditer(r'\(?-?[$€£]?\s?\d[\d,]*(?:\.\d+)?\)?', t):
        v = to_num(m.group(0))
        if v is not None: found.append(v)
    rel = float(spec.get('rel_tol', 0.005))
    missing = []
    for want in spec['numbers']:
        want = float(want)
        if not any(abs(v - want) <= max(abs(want) * rel, 0.01) for v in found):
            missing.append(want)
    need = int(spec.get('min_count', len(spec['numbers'])))
    found = len(spec['numbers']) - len(missing)
    return found >= need, f'{found}/{len(spec["numbers"])} figures present (need {need})' + (f'; missing {missing}' if missing else '')

def c_custom(ws, ref, spec, task_dir=None):
    mod_path = os.path.join(task_dir, spec.get('module', 'check.py'))
    # Each task's check module loads in isolation: a unique module name, no bytecode written into the task
    # folder, and any helper modules or sys.path entries it adds are dropped afterwards, so one grader process
    # grading many tasks cannot hand task B a helper module that task A imported under the same name.
    saved_modules, saved_path, saved_flag = set(sys.modules), list(sys.path), sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        name = 'task_check_' + re.sub(r'\W', '_', os.path.relpath(mod_path, ROOT) if 'ROOT' in globals() else mod_path)
        s = importlib.util.spec_from_file_location(name, mod_path)
        m = importlib.util.module_from_spec(s); s.loader.exec_module(m)  # type: ignore
        res = m.check(ws, ref)
    finally:
        for k in set(sys.modules) - saved_modules:
            if k not in ('formulas', 'openpyxl', 'pandas', 'numpy') and not k.startswith(('formulas.', 'openpyxl.', 'pandas.', 'numpy.', 'grade')):
                sys.modules.pop(k, None)
        sys.path[:] = saved_path
        sys.dont_write_bytecode = saved_flag
    ok = all(r.get('passed') for r in res)
    return ok, '; '.join(f"{r.get('name')}={'ok' if r.get('passed') else 'FAIL'} {r.get('detail','')}" for r in res)

CHECKS = {
    'file_exists': c_file_exists,
    'csv_columns': c_csv_columns,
    'csv_row_count': c_csv_row_count,
    'csv_set_equal': c_csv_set_equal,
    'csv_values_match': c_csv_values_match,
    'xlsx_has_formulas': c_xlsx_has_formulas,
    'xlsx_value_present': c_xlsx_value_present,
    'xlsx_no_errors': c_xlsx_no_errors,
    'text_contains_all': c_text_contains_all,
    'text_contains_any': c_text_contains_any,
    'text_not_contains': c_text_not_contains,
    'text_numbers_present': c_text_numbers_present,
    'text_matches_all': c_text_matches_all,
    'text_sentence_matches': c_text_sentence_matches,
}

def grade(task_dir: str, ws: str) -> dict:
    task = yaml.safe_load(open(os.path.join(task_dir, 'task.yaml')))
    ref = os.path.join(task_dir, 'reference')
    out = []; grader_errors = []
    for spec in task.get('checks', []):
        t = spec['type']
        try:
            if t == 'custom':
                ok, detail = c_custom(ws, ref, spec, task_dir=task_dir)
            elif t not in CHECKS:
                # A check type this grader build does not know (a task edited after the runner
                # started, or a stale bundle) must never read as the agent failing.
                raise KeyError(f'unknown check type {t!r}; the grader process may predate the task file')
            else:
                ok, detail = CHECKS[t](ws, ref, spec)
        except Exception as e:  # a grader crash is recorded, never a crashed run
            msg = f'{type(e).__name__}: {e}'
            if re.search(r'Unable to read workbook|could not read stylesheet|not a zip file|BadZipFile|invalid XML|is not a valid', msg, re.I):
                # The deliverable itself cannot be opened by a standard reader: that is the agent's
                # defect, graded as a failed check, not a grader error to regrade later.
                ok, detail = False, f'deliverable unreadable by a standard reader: {msg}'
            else:
                ok, detail = False, f'grader error: {msg}'
                grader_errors.append(spec.get('name', t))
        out.append({'name': spec.get('name', t), 'type': t,
                    'required': bool(spec.get('required', True)), 'passed': bool(ok), 'detail': str(detail)[:500]})
    passed = all(c['passed'] for c in out if c['required'])
    # grader_errors non-empty means "ungraded until regraded": reports must not count it as a fail.
    return {'task': task.get('id'), 'passed': passed, 'checks': out, 'grader_errors': grader_errors}

if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__); sys.exit(2)
    print(json.dumps(grade(sys.argv[1], sys.argv[2]), indent=2))
