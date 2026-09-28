#!/usr/bin/env python3
"""Metamorphic testing of the desk grader: estimate how often it FAILS correct work.

Every desk task ships a reference solution that must pass. This tool applies transforms that
should not change whether that work is correct (CRLF line endings, all-quoted CSV fields, a
workbook re-saved through openpyxl, Markdown re-wrapped to another width, ...), grades the
transformed copy with bench/grade.py, and reports every transform that turns the pass into a
fail, with the failing check and its detail.

Transform classes
  invariant           correct work stays correct under any reasonable contract. A fail here is a
                      CANDIDATE FALSE NEGATIVE: either the grader is too strict or the transform
                      was misclassified for this task; a human decides which.
  contract-sensitive  may legitimately matter for some tasks (a BOM on a CSV meant for an import
                      template, row order the ask asks for, column order of a fixed template).
                      Fails are reported as information, never as grader defects.

Each transform has a default class; `classify()` promotes an invariant transform to
contract-sensitive for a given task when the task's own checks or ask make the property part
of the contract. The rules (see RULES below and classify()):
  csv_reverse_rows     sensitive when the ask, or the name of a check on the file, uses ordering
                       language (sorted, ranked, oldest first, chronological, ...), or a custom/plan
                       module reads the file (order is then not provably free); otherwise invariant
                       (no built-in check reads row order).
  csv_reverse_columns  sensitive when a csv_columns check on the file has `exact: true`, or the
                       ask mentions an import or a template; otherwise invariant.
  csv_bom / csv_header_case
                       sensitive for the same import/template/exact-columns reasons; otherwise
                       invariant (csv_bom is listed as contract-sensitive by default and demoted
                       to invariant only for a non-import CSV without an exact template).
  xlsx_sheet_order / xlsx_cover_sheet_first
                       sensitive when some check reads the workbook as one table without naming a
                       sheet (csv_* checks on an .xlsx path, or a custom module that reads the file
                       and uses `.active` / `worksheets[0]`): the first sheet is then the de facto
                       contract. Otherwise invariant.
  md_smart_quotes, move_to_subdir, rename_case
                       always contract-sensitive (literal characters and file location are what
                       the ask names).
  everything else      invariant.

Usage
  metamorphic.py [task ...] [--transforms a,b] [--json PATH] [--all] [--list-transforms]
  With no task ids and no --all, a representative sample of csv / xlsx / md / html / ics tasks
  runs (LibreOffice recalculation costs 2-5 s per workbook, so --all is slow).
Exit status: 0 no candidate false negatives, 1 some found, 2 a reference solution failed baseline.

Offline tooling only: nothing in the runner or the grader imports this module. It imports
bench/grade.py lazily, only when grading.
"""
from __future__ import annotations

import argparse
import csv
import fnmatch
import glob
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import textwrap
import time
from dataclasses import dataclass, field
from typing import Callable, Optional

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESK = os.path.join(ROOT, 'tasks', 'desk')

INVARIANT = 'invariant'
SENSITIVE = 'contract-sensitive'

# Representative default sample: csv imports (exact template, custom check, two files), a csv
# report, xlsx with memo, xlsx read as a table, xlsx with custom check, plain memos, a memo with
# a custom check, files in a subfolder, html pages (custom and plain), and an ics calendar.
DEFAULT_SAMPLE = [
    'zendesk-users-import', 'address-standardize', 'card-expense-coding', 'long-to-wide-attendance',
    'project-margin', 'price-list-update', 'bank-reconciliation',
    'investor-update', 'minutes-from-transcript', 'maintenance-notices',
    'menu-page-allergens', 'team-directory-page', 'schedule-to-ics',
]

# ---------------------------------------------------------------- file kinds

def file_kind(relpath: str) -> Optional[str]:
    ext = os.path.splitext(relpath)[1].lower()
    return {'.csv': 'csv', '.tsv': 'csv', '.xlsx': 'xlsx', '.xlsm': 'xlsx', '.md': 'md', '.markdown': 'md',
            '.txt': 'md', '.html': 'html', '.htm': 'html', '.ics': 'ics'}.get(ext)

# ---------------------------------------------------------------- CSV transforms (bytes -> bytes)

_BOM = b'\xef\xbb\xbf'


def _csv_parse(data: bytes):
    """-> (rows, bom, line_terminator, trailing_newline, delimiter). Blank lines are kept as []."""
    bom = data.startswith(_BOM)
    text = data[len(_BOM):].decode('utf-8') if bom else data.decode('utf-8')
    term = '\r\n' if '\r\n' in text else '\n'
    first = text.split('\n', 1)[0]
    delim = '\t' if first.count('\t') > first.count(',') else ','
    rows = list(csv.reader(io.StringIO(text, newline=''), delimiter=delim))
    return rows, bom, term, text.endswith(('\n', '\r')), delim


def _csv_write(rows, bom: bool, term: str, trailing: bool, delim: str, quoting=csv.QUOTE_MINIMAL) -> bytes:
    buf = io.StringIO(newline='')
    w = csv.writer(buf, delimiter=delim, lineterminator=term, quoting=quoting)
    for r in rows:
        w.writerow(r)
    out = buf.getvalue()
    if not trailing and out.endswith(term):
        out = out[:-len(term)]
    return (_BOM if bom else b'') + out.encode('utf-8')


def csv_crlf(data: bytes) -> bytes:
    rows, bom, _, trailing, d = _csv_parse(data)
    return _csv_write(rows, bom, '\r\n', trailing, d)


def csv_quote_all(data: bytes) -> bytes:
    rows, bom, term, trailing, d = _csv_parse(data)
    return _csv_write(rows, bom, term, trailing, d, quoting=csv.QUOTE_ALL)


def csv_reverse_rows(data: bytes) -> bytes:
    rows, bom, term, trailing, d = _csv_parse(data)
    if not rows:
        return data
    body = [r for r in rows[1:] if r]
    return _csv_write([rows[0]] + body[::-1], bom, term, trailing, d)


def csv_reverse_columns(data: bytes) -> bytes:
    rows, bom, term, trailing, d = _csv_parse(data)
    width = max((len(r) for r in rows), default=0)
    out = [(r + [''] * (width - len(r)))[::-1] if r else r for r in rows]
    return _csv_write(out, bom, term, trailing, d)


def csv_bom(data: bytes) -> bytes:
    return data if data.startswith(_BOM) else _BOM + data


def csv_no_final_newline(data: bytes) -> bytes:
    return data.rstrip(b'\r\n')


def csv_header_case(data: bytes) -> bytes:
    """Header cells to Title Case with spaces for underscores ("external_id" -> "External Id")."""
    rows, bom, term, trailing, d = _csv_parse(data)
    if not rows:
        return data
    rows[0] = [h.replace('_', ' ').title() for h in rows[0]]
    return _csv_write(rows, bom, term, trailing, d)

# ---------------------------------------------------------------- XLSX transforms (src path -> dst path)
# Each returns True when it wrote dst, False when it does not apply (dst untouched).


def xlsx_resave(src: str, dst: str) -> bool:
    from openpyxl import load_workbook
    load_workbook(src).save(dst)
    return True


def xlsx_extra_sheet(src: str, dst: str) -> bool:
    """Append an unrelated text-only sheet (no numbers, no formulas) after the existing ones."""
    from openpyxl import load_workbook
    wb = load_workbook(src)
    name = 'Notes'
    while name in wb.sheetnames:
        name += '_'
    sh = wb.create_sheet(name)
    sh['A1'] = 'Working notes'
    sh['A2'] = 'Prepared from the files in the folder; see the other sheets for the figures.'
    wb.save(dst)
    return True


def xlsx_sheet_order(src: str, dst: str) -> bool:
    """Reverse the order of the sheets (needs two or more)."""
    from openpyxl import load_workbook
    wb = load_workbook(src)
    if len(wb.worksheets) < 2:
        return False
    wb._sheets.reverse()  # openpyxl keeps sheet order in this list; formulas refer to sheets by name
    wb.active = 0
    wb.save(dst)
    return True


def xlsx_cover_sheet_first(src: str, dst: str) -> bool:
    """Insert a text-only cover sheet in front of the existing sheets and make it active."""
    from openpyxl import load_workbook
    wb = load_workbook(src)
    name = 'Cover'
    while name in wb.sheetnames:
        name += '_'
    sh = wb.create_sheet(name, 0)
    sh['A1'] = 'Workbook prepared for the owner'
    sh['A2'] = 'Figures are on the following sheets.'
    wb.active = 0
    wb.save(dst)
    return True

# ---------------------------------------------------------------- Markdown / text transforms (bytes -> bytes)

_LIST_MARK = re.compile(r'^(\s*)([-*+]|\d{1,3}[.)])(\s+)')
_BLOCK_START = re.compile(r'^\s*([-*+]|\d{1,3}[.)])\s|^\s*#|^\s*>|^\s*\||^\s*(```|~~~)|^\s*<|^\s*([-*_]\s*){3,}$|^\s*=+\s*$|^\s*-+\s*$')


def _decode(data: bytes) -> tuple[str, str, bool]:
    text = data.decode('utf-8')
    term = '\r\n' if '\r\n' in text else '\n'
    return text.replace('\r\n', '\n'), term, data.startswith(_BOM)


def _encode(text: str, term: str) -> bytes:
    return (text.replace('\n', term) if term != '\n' else text).encode('utf-8')


def _md_items(lines: list[str]):
    """Split a paragraph block into wrappable items: (prefix, subsequent_indent, words) or ('raw', line)."""
    items: list = []
    for ln in lines:
        m = _LIST_MARK.match(ln)
        if re.match(r'^\s*(#|>|\||<)', ln) or re.match(r'^\s*([-*_]\s*){3,}$', ln):
            items.append(('raw', ln)); continue
        if m:
            prefix = m.group(0)
            items.append(['item', prefix, ' ' * len(prefix), ln[len(prefix):].split(), ln.endswith('  ')])
        elif items and items[-1][0] == 'item' and not items[-1][4]:
            items[-1][3].extend(ln.split()); items[-1][4] = ln.endswith('  ')
        else:
            indent = re.match(r'^\s*', ln).group(0)
            items.append(['item', indent, indent, ln.split(), ln.endswith('  ')])
    return items


def _wrap_item(prefix: str, indent: str, words: list[str], width: int) -> list[str]:
    """Greedy wrap; tries wider widths if a continuation line would start with Markdown syntax."""
    if not words:
        return [prefix.rstrip()]
    for w in range(width, width + 40):
        lines, cur = [], prefix + words[0]
        for word in words[1:]:
            if len(cur) + 1 + len(word) > w:
                lines.append(cur); cur = indent + word
            else:
                cur += ' ' + word
        lines.append(cur)
        if not any(_BLOCK_START.match(x[len(indent):]) or re.match(r'^\d{1,3}[.)]$', x.strip()) for x in lines[1:]):
            return lines
    return [prefix + ' '.join(words)]


def _md_rewrap(data: bytes, width: int) -> bytes:
    text, term, _ = _decode(data)
    out: list[str] = []
    block: list[str] = []
    fence = None

    def flush():
        if not block:
            return
        if any(re.match(r'^\s*\|', ln) for ln in block) or any(re.match(r'^\s*[=-]+\s*$', ln) for ln in block):
            out.extend(block)  # tables and setext headings stay as they are
        else:
            for it in _md_items(block):
                if it[0] == 'raw':
                    out.append(it[1])
                else:
                    _, prefix, indent, words, hard = it
                    wrapped = _wrap_item(prefix, indent, words, width)
                    if hard:
                        wrapped[-1] += '  '
                    out.extend(wrapped)
        block.clear()

    for ln in text.split('\n'):
        if fence:
            out.append(ln)
            if ln.strip().startswith(fence):
                fence = None
            continue
        m = re.match(r'^\s*(```|~~~)', ln)
        if m:
            flush(); fence = m.group(1); out.append(ln); continue
        if not ln.strip():
            flush(); out.append(ln); continue
        if ln.startswith(('    ', '\t')) and not block:
            out.append(ln); continue  # indented code block
        block.append(ln)
    flush()
    return _encode('\n'.join(out), term)


def md_rewrap_72(data: bytes) -> bytes:
    return _md_rewrap(data, 72)


def md_rewrap_narrow(data: bytes) -> bytes:
    return _md_rewrap(data, 40)


def md_unwrap(data: bytes) -> bytes:
    return _md_rewrap(data, 100000)


def md_star_bullets(data: bytes) -> bytes:
    text, term, _ = _decode(data)
    lines, fence = [], None
    for ln in text.split('\n'):
        if re.match(r'^\s*(```|~~~)', ln):
            fence = None if fence else ln.strip()[:3]
        elif not fence:
            ln = re.sub(r'^(\s*)-(\s+)(?!-)', r'\1*\2', ln) if not re.match(r'^\s*([-*_]\s*){3,}$', ln) else ln
        lines.append(ln)
    return _encode('\n'.join(lines), term)


def md_trailing_ws(data: bytes) -> bytes:
    """One trailing space on every non-blank line (one, not two: two spaces is a Markdown hard break)."""
    text, term, _ = _decode(data)
    return _encode('\n'.join(ln + ' ' if ln.strip() and not ln.endswith(' ') else ln for ln in text.split('\n')), term)


def md_crlf(data: bytes) -> bytes:
    text, _, _ = _decode(data)
    return _encode(text, '\r\n')


def md_smart_quotes(data: bytes) -> bytes:
    """Straight quotes and apostrophes to typographic ones (what word processors emit)."""
    text, term, _ = _decode(data)
    text = re.sub(r"(?<=\w)'(?=\w)", '\u2019', text)
    text = re.sub(r"(^|[\s(\[])'", '\\1\u2018', text)
    text = text.replace("'", '\u2019')
    text = re.sub(r'(^|[\s(\[])"', '\\1\u201c', text)
    text = text.replace('"', '\u201d')
    return _encode(text, term)

# ---------------------------------------------------------------- HTML transforms (bytes -> bytes)

_TAG = re.compile(r'(<!--.*?-->|<![^>]*>|</?[A-Za-z][^>]*>)', re.S)
_RAW = {'script', 'style', 'pre', 'textarea'}
_HTML_BLOCK = {'address', 'article', 'aside', 'blockquote', 'body', 'caption', 'dd', 'div', 'dl', 'dt', 'figure',
               'footer', 'form', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'header', 'hr', 'li', 'main', 'nav', 'ol', 'p',
               'section', 'table', 'tbody', 'td', 'tfoot', 'th', 'thead', 'tr', 'ul', 'head', 'html', 'title', 'meta',
               'link'}


def _html_segments(text: str):
    """Yield (kind, segment, tagname) with kind in text | tag | raw. Content of script/style/pre/textarea
    (and everything that looks like a tag inside it) is 'raw' until the matching close tag."""
    raw_until = None
    for seg in _TAG.split(text):
        if not seg:
            continue
        is_tag = bool(_TAG.fullmatch(seg))
        name = ''
        if is_tag and not seg.startswith('<!'):
            name = re.match(r'</?([A-Za-z][A-Za-z0-9-]*)', seg).group(1).lower()
        if raw_until:
            if is_tag and seg.startswith('</') and name == raw_until:
                raw_until = None
                yield 'tag', seg, name
            else:
                yield 'raw', seg, ''
            continue
        if is_tag:
            yield 'tag', seg, name
            if name in _RAW and not seg.startswith('</') and not seg.rstrip('>').endswith('/'):
                raw_until = name
        else:
            yield 'text', seg, ''


def _html_map(data: bytes, text_fn=None, tag_fn=None) -> bytes:
    text, term, _ = _decode(data)
    out = []
    for kind, seg, name in _html_segments(text):
        if kind == 'text' and text_fn:
            seg = text_fn(seg)
        elif kind == 'tag' and tag_fn:
            seg = tag_fn(seg, name)
        out.append(seg)
    return _encode(''.join(out), term)


def html_reindent(data: bytes) -> bytes:
    """Every whitespace run that holds a line break becomes a line break plus a fixed 4-space indent."""
    return _html_map(data, text_fn=lambda s: re.sub(r'[ \t]*\n\s*', '\n    ', s))


def html_break_blocks(data: bytes) -> bytes:
    """A line break before every block-level tag (whitespace between blocks does not render)."""
    return _html_map(data, tag_fn=lambda s, n: '\n' + s if n in _HTML_BLOCK else s)


def html_rewrap_text(data: bytes, width: int = 40) -> bytes:
    """Break long text runs across lines at word boundaries (rendered identically)."""
    def wrap(s: str) -> str:
        if len(s) <= width or not s.strip():
            return s
        lead = re.match(r'^\s*', s).group(0); trail = re.search(r'\s*$', s).group(0)
        words = s.split()
        lines, cur = [], words[0]
        for w in words[1:]:
            if len(cur) + 1 + len(w) > width:
                lines.append(cur); cur = w
            else:
                cur += ' ' + w
        lines.append(cur)
        return lead + '\n'.join(lines) + trail
    return _html_map(data, text_fn=wrap)


def html_escape_quotes(data: bytes) -> bytes:
    """Apostrophes and double quotes in text written as character references (&#39; &quot;)."""
    return _html_map(data, text_fn=lambda s: s.replace("'", '&#39;').replace('"', '&quot;'))


def html_entities_nonascii(data: bytes) -> bytes:
    """Non-ASCII characters in text written as numeric character references (&#233; for e-acute)."""
    return _html_map(data, text_fn=lambda s: ''.join(c if ord(c) < 128 else f'&#{ord(c)};' for c in s))


def html_upper_tags(data: bytes) -> bytes:
    """Tag names in upper case (HTML tag names are case-insensitive)."""
    def up(s: str, n: str) -> str:
        if not n:
            return s
        return re.sub(r'^(</?)([A-Za-z][A-Za-z0-9-]*)', lambda m: m.group(1) + m.group(2).upper(), s)
    return _html_map(data, tag_fn=up)

# ---------------------------------------------------------------- ICS transforms (bytes -> bytes)


def ics_crlf(data: bytes) -> bytes:
    """RFC 5545 content lines end in CRLF."""
    text, _, _ = _decode(data)
    return _encode(text, '\r\n')


def ics_fold(data: bytes, width: int = 40) -> bytes:
    """RFC 5545 line folding: a content line may be split between any two characters by a line break
    followed by one space (section 3.1). Folding at 40 octets exercises what folding at 75 would on long
    DESCRIPTION lines."""
    text, term, _ = _decode(data)
    out = []
    for ln in text.split('\n'):
        if len(ln.encode('utf-8')) <= width:
            out.append(ln); continue
        parts, cur, limit = [], '', width
        for ch in ln:
            if len((cur + ch).encode('utf-8')) > limit:
                parts.append(cur); cur = ch; limit = width - 1
            else:
                cur += ch
        parts.append(cur)
        out.append('\n '.join(parts))
    return _encode('\n'.join(out), term)

# ---------------------------------------------------------------- registry


@dataclass(frozen=True)
class Transform:
    name: str
    kind: str                 # csv | xlsx | md | html | ics | workspace
    default_class: str
    fn: Callable
    doc: str = ''


TRANSFORMS: dict[str, Transform] = {t.name: t for t in [
    Transform('csv_crlf', 'csv', INVARIANT, csv_crlf, 'CRLF line endings'),
    Transform('csv_quote_all', 'csv', INVARIANT, csv_quote_all, 'every field quoted'),
    Transform('csv_no_final_newline', 'csv', INVARIANT, csv_no_final_newline, 'no newline after the last record'),
    Transform('csv_reverse_rows', 'csv', INVARIANT, csv_reverse_rows, 'data rows in reverse order'),
    Transform('csv_reverse_columns', 'csv', INVARIANT, csv_reverse_columns, 'columns in reverse order'),
    Transform('csv_header_case', 'csv', INVARIANT, csv_header_case, 'header names in Title Case with spaces'),
    Transform('csv_bom', 'csv', SENSITIVE, csv_bom, 'UTF-8 byte-order mark (Excel "CSV UTF-8")'),
    Transform('xlsx_resave', 'xlsx', INVARIANT, xlsx_resave, 're-saved through openpyxl (no cached values)'),
    Transform('xlsx_extra_sheet', 'xlsx', INVARIANT, xlsx_extra_sheet, 'unrelated text-only sheet appended'),
    Transform('xlsx_sheet_order', 'xlsx', INVARIANT, xlsx_sheet_order, 'sheet order reversed (2+ sheets)'),
    Transform('xlsx_cover_sheet_first', 'xlsx', INVARIANT, xlsx_cover_sheet_first, 'text-only cover sheet inserted first'),
    Transform('md_rewrap_72', 'md', INVARIANT, md_rewrap_72, 'paragraphs and list items hard-wrapped at 72 columns'),
    Transform('md_rewrap_narrow', 'md', INVARIANT, md_rewrap_narrow, 'paragraphs and list items hard-wrapped at 40 columns'),
    Transform('md_unwrap', 'md', INVARIANT, md_unwrap, 'each paragraph and list item on one line'),
    Transform('md_star_bullets', 'md', INVARIANT, md_star_bullets, "'-' bullets written as '*'"),
    Transform('md_trailing_ws', 'md', INVARIANT, md_trailing_ws, 'one trailing space on every line'),
    Transform('md_crlf', 'md', INVARIANT, md_crlf, 'CRLF line endings'),
    Transform('md_smart_quotes', 'md', SENSITIVE, md_smart_quotes, 'typographic quotes and apostrophes'),
    Transform('html_reindent', 'html', INVARIANT, html_reindent, 'indentation whitespace changed'),
    Transform('html_break_blocks', 'html', INVARIANT, html_break_blocks, 'line break before every block tag'),
    Transform('html_rewrap_text', 'html', INVARIANT, html_rewrap_text, 'long text runs broken across lines'),
    Transform('html_escape_quotes', 'html', INVARIANT, html_escape_quotes, "' and \" in text as &#39; / &quot;"),
    Transform('html_entities_nonascii', 'html', INVARIANT, html_entities_nonascii, 'non-ASCII text as &#NNN; references'),
    Transform('html_upper_tags', 'html', INVARIANT, html_upper_tags, 'tag names upper-cased'),
    Transform('ics_crlf', 'ics', INVARIANT, ics_crlf, 'CRLF line endings (RFC 5545)'),
    Transform('ics_fold', 'ics', INVARIANT, ics_fold, 'long lines folded at 40 octets (RFC 5545 section 3.1)'),
    Transform('move_to_subdir', 'workspace', SENSITIVE, None, 'deliverables moved into output/'),
    Transform('rename_case', 'workspace', SENSITIVE, None, 'deliverable file names upper-cased'),
]}

# ---------------------------------------------------------------- classification

_ORDER_ASK = re.compile(
    r'\b(sort(ed|ing)?|order(ed)? by|in (\w+ ){0,2}order(?! to)|rank(ed|ing)?|alphabetical(ly)?|chronological(ly)?|'
    r'ascending|descending|(largest|biggest|highest|oldest|newest|earliest|latest|most \w+) first|top \d+)\b', re.I)
_IMPORT_ASK = re.compile(r'\b(import(ing|er|s)?|template|upload(ing)?)\b', re.I)
_TABLE_CHECKS = {'csv_columns', 'csv_row_count', 'csv_set_equal', 'csv_values_match', 'forecast_error', 'not_fooled'}
_FIRST_SHEET = re.compile(r'\.active\b|worksheets\[0\]|sheetnames\[0\]|sheet_name=0|read_excel\((?![^)]*sheet_name)')

RULES = {
    'csv_reverse_rows': 'sensitive if the ask or a check name on the file uses ordering language, or a custom/plan module reads the file',
    'csv_reverse_columns': 'sensitive if csv_columns exact:true on the file or the ask mentions import/template/upload',
    'csv_bom': 'sensitive if csv_columns exact:true or the ask mentions import/template/upload; else invariant',
    'csv_header_case': 'sensitive if csv_columns exact:true or the ask mentions import/template/upload',
    'xlsx_sheet_order': 'sensitive if a check reads the workbook as one table without naming a sheet',
    'xlsx_cover_sheet_first': 'sensitive if a check reads the workbook as one table without naming a sheet',
}


def _path_matches(pattern: str, relpath: str) -> bool:
    pat, rel = str(pattern).lower(), relpath.lower()
    return fnmatch.fnmatch(rel, pat) or fnmatch.fnmatch(os.path.basename(rel), pat) or \
        ('/' not in pat and fnmatch.fnmatch(os.path.basename(rel), pat))


def checks_on(task: dict, relpath: str) -> list[dict]:
    """Built-in checks that read this deliverable (by path glob)."""
    out = []
    for c in task.get('checks') or []:
        paths = [c.get('path')] + [x.get('path') for x in c.get('forbidden_text') or []] + \
                [(c.get('flag') or {}).get('path')]
        if any(p and _path_matches(p, relpath) for p in paths):
            out.append(c)
    return out


def custom_reads(custom_sources: dict[str, str], relpath: str) -> list[str]:
    """Names of task-local check modules whose source mentions the deliverable (by basename or stem)."""
    base = os.path.basename(relpath)
    stem = os.path.splitext(base)[0]
    ext = os.path.splitext(base)[1]
    hits = []
    for mod, src in custom_sources.items():
        if base in src or re.search(r'["\'/]' + re.escape(stem) + r'(\.|["\'*])', src) or \
                (os.path.dirname(relpath) and os.path.dirname(relpath) + '/' in src and ext in src):
            hits.append(mod)
    return hits


def classify(tname: str, task: dict, relpath: str, custom_sources: dict[str, str] | None = None) -> tuple[str, str]:
    """Class of transform `tname` applied to deliverable `relpath` of `task` (the parsed task.yaml),
    with a one-line reason. `custom_sources` maps task-local module names to their source text."""
    t = TRANSFORMS[tname]
    custom_sources = custom_sources or {}
    if t.default_class == SENSITIVE and tname != 'csv_bom':
        return SENSITIVE, 'contract-sensitive by definition'
    ask = str(task.get('ask') or '')
    on = checks_on(task, relpath)
    exact = any(c.get('type') == 'csv_columns' and c.get('exact') for c in on)
    importish = bool(_IMPORT_ASK.search(ask))
    if tname == 'csv_reverse_rows':
        m = _ORDER_ASK.search(ask)
        if m:
            return SENSITIVE, f'ask uses ordering language ({m.group(0)!r})'
        for c in on:
            if c.get('type') == 'csv_columns':
                continue  # "template columns in order" is about column order, not row order
            m = _ORDER_ASK.search(str(c.get('name') or ''))
            if m:
                return SENSITIVE, f'check {c.get("name")!r} names an order'
        mods = custom_reads(custom_sources, relpath)
        if mods:
            return SENSITIVE, f'custom module {mods[0]} reads the file; order not provably free'
        return INVARIANT, 'no check or ask constrains row order'
    if tname in ('csv_reverse_columns', 'csv_header_case', 'csv_bom'):
        if exact:
            return SENSITIVE, 'csv_columns exact:true (fixed template)'
        if importish:
            return SENSITIVE, 'ask is an import/template/upload'
        return INVARIANT, 'no exact template; columns are matched by name'
    if tname in ('xlsx_sheet_order', 'xlsx_cover_sheet_first'):
        table = [c for c in on if c.get('type') in _TABLE_CHECKS and not c.get('sheet')]
        if table:
            return SENSITIVE, f'{table[0].get("type")} reads the workbook as one table (first sheet)'
        for mod in custom_reads(custom_sources, relpath):
            if _FIRST_SHEET.search(custom_sources[mod]):
                return SENSITIVE, f'custom module {mod} reads the first/active sheet'
        return INVARIANT, 'checks scan every sheet or name one'
    return t.default_class, 'default class'

# ---------------------------------------------------------------- running


def load_task(task_id: str) -> tuple[str, dict, dict[str, str]]:
    import yaml
    td = os.path.join(DESK, task_id)
    with open(os.path.join(td, 'task.yaml')) as f:
        task = yaml.safe_load(f) or {}
    sources = {}
    for c in task.get('checks') or []:
        if c.get('type') in ('custom', 'plan_feasible'):
            mod = c.get('module', 'check.py' if c['type'] == 'custom' else 'plan_check.py')
            p = os.path.join(td, mod)
            if os.path.isfile(p):
                with open(p, encoding='utf-8', errors='replace') as f:
                    sources[mod] = f.read()
    return td, task, sources


def deliverables(sol_dir: str) -> list[str]:
    out = []
    for base, dirs, files in os.walk(sol_dir):
        dirs[:] = [d for d in dirs if d not in ('__pycache__',)]
        for f in files:
            out.append(os.path.relpath(os.path.join(base, f), sol_dir))
    return sorted(out)


def apply_transform(tname: str, src_dir: str, dst_dir: str, files: list[str]) -> list[str]:
    """Copy src_dir to dst_dir with the transform applied to every matching deliverable.
    Returns the relpaths it changed (empty: not applicable or a no-op)."""
    t = TRANSFORMS[tname]
    shutil.copytree(src_dir, dst_dir)
    changed = []
    if t.kind == 'workspace':
        for rel in files:
            if tname == 'move_to_subdir':
                new = os.path.join('output', rel)
            else:
                d, b = os.path.split(rel)
                s, e = os.path.splitext(b)
                new = os.path.join(d, s.upper() + e)
            if new == rel:
                continue
            os.makedirs(os.path.dirname(os.path.join(dst_dir, new)), exist_ok=True)
            os.replace(os.path.join(dst_dir, rel), os.path.join(dst_dir, new))
            changed.append(rel)
        return changed
    for rel in files:
        if file_kind(rel) != t.kind:
            continue
        p = os.path.join(dst_dir, rel)
        if t.kind == 'xlsx':
            tmp = p + '.mt.xlsx'
            if t.fn(p, tmp):
                os.replace(tmp, p); changed.append(rel)
        else:
            with open(p, 'rb') as f:
                before = f.read()
            after = t.fn(before)
            if after != before:
                with open(p, 'wb') as f:
                    f.write(after)
                changed.append(rel)
    return changed


def _install_recalc_cache(grade_mod) -> None:
    """Content-keyed memo over grade.recalculated_workbook: a workbook the transform left byte-identical
    is recalculated once per run instead of once per transform. Offline only; grade.py is not modified."""
    if getattr(grade_mod, '_metamorphic_cache', False):
        return
    orig = grade_mod.recalculated_workbook
    memo: dict[str, str] = {}

    def cached(path: str) -> str:
        try:
            with open(path, 'rb') as f:
                h = hashlib.sha256(f.read()).hexdigest()
        except OSError:
            return orig(path)
        hit = memo.get(h)
        if hit and os.path.exists(hit):
            return hit
        out = orig(path)
        if out != path:  # never memoise the fallback (the input itself lives in a temp workspace)
            memo[h] = out
        return out
    grade_mod.recalculated_workbook = cached
    grade_mod._metamorphic_cache = True


def _grade(td: str, ws: str) -> dict:
    sys.path.insert(0, os.path.join(ROOT, 'bench'))
    import grade as grade_mod
    _install_recalc_cache(grade_mod)
    return grade_mod.grade(td, ws)


def _failed(g: dict) -> list[dict]:
    return [{'name': c['name'], 'type': c['type'], 'detail': c['detail']} for c in g['checks']
            if c['required'] and not c['passed']]


def run_task(task_id: str, tnames: list[str], log=None) -> dict:
    td, task, sources = load_task(task_id)
    sol = os.path.join(td, 'reference_solution')
    rec: dict = {'task': task_id, 'baseline': None, 'results': {}}
    if not os.path.isdir(sol):
        rec['baseline'] = 'missing'
        return rec
    files = deliverables(sol)
    rec['deliverables'] = files
    work = tempfile.mkdtemp(prefix='metamorphic-')
    try:
        base_ws = os.path.join(work, 'baseline')
        shutil.copytree(sol, base_ws)
        t0 = time.time()
        g = _grade(td, base_ws)
        rec['baseline'] = 'pass' if g['passed'] else 'FAIL'
        if not g['passed']:
            rec['baseline_failed'] = _failed(g)
            return rec
        if log: log(f'  {task_id}: baseline pass ({time.time() - t0:.1f}s)')
        for tn in tnames:
            t = TRANSFORMS[tn]
            targets = files if t.kind == 'workspace' else [f for f in files if file_kind(f) == t.kind]
            if not targets:
                continue  # transform does not apply to this task's deliverable types
            ws = os.path.join(work, tn)
            changed = apply_transform(tn, sol, ws, files)
            if not changed:
                rec['results'][tn] = {'status': 'noop'}
                continue
            classes = [classify(tn, task, rel, sources) for rel in changed]
            cls = SENSITIVE if any(c == SENSITIVE for c, _ in classes) else INVARIANT
            reason = next((r for c, r in classes if c == cls), classes[0][1])
            t0 = time.time()
            g = _grade(td, ws)
            r = {'status': 'pass' if g['passed'] else 'fail', 'class': cls, 'reason': reason, 'files': changed}
            if g.get('grader_errors'):
                r['status'] = 'error'
            if not g['passed']:
                r['failed'] = _failed(g)
            rec['results'][tn] = r
            if log: log(f'  {task_id} x {tn}: {r["status"]} [{cls}] ({time.time() - t0:.1f}s)')
            shutil.rmtree(ws, ignore_errors=True)
        return rec
    finally:
        shutil.rmtree(work, ignore_errors=True)


def candidates(records: list[dict]) -> list[dict]:
    """Invariant transform + fail (or grader error): candidate false negatives."""
    out = []
    for rec in records:
        for tn, r in rec['results'].items():
            if r.get('class') == INVARIANT and r['status'] in ('fail', 'error'):
                for f in r.get('failed', []):
                    out.append({'task': rec['task'], 'transform': tn, 'files': r['files'], 'check': f['name'],
                                'check_type': f['type'], 'detail': f['detail']})
    return out


def sensitive_fails(records: list[dict]) -> list[dict]:
    out = []
    for rec in records:
        for tn, r in rec['results'].items():
            if r.get('class') == SENSITIVE and r['status'] in ('fail', 'error'):
                out.append({'task': rec['task'], 'transform': tn, 'reason': r['reason'],
                            'checks': [f['name'] for f in r.get('failed', [])]})
    return out


LEGEND = ('. pass   F FAIL on invariant transform (candidate false negative)   f fail on contract-sensitive transform\n'
          's pass on contract-sensitive transform   E grader error   = no-op   (blank) not applicable')


def matrix(records: list[dict], tnames: list[str]) -> str:
    used = [t for t in tnames if any(t in r['results'] for r in records)]
    idx = {t: i + 1 for i, t in enumerate(used)}
    w = max([len(r['task']) for r in records] + [4])
    lines = ['transforms: ' + '  '.join(f'{i}={t}' for t, i in idx.items()), '']
    lines.append(f'{"task":{w}}  base  ' + ' '.join(f'{idx[t]:>2}' for t in used))
    for rec in records:
        cells = []
        for t in used:
            r = rec['results'].get(t)
            if r is None: c = ''
            elif r['status'] == 'noop': c = '='
            elif r['status'] == 'error': c = 'E'
            elif r['status'] == 'pass': c = '.' if r['class'] == INVARIANT else 's'
            else: c = 'F' if r['class'] == INVARIANT else 'f'
            cells.append(f'{c:>2}')
        lines.append(f'{rec["task"]:{w}}  {str(rec["baseline"]):4}  ' + ' '.join(cells))
    lines += ['', LEGEND]
    return '\n'.join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('tasks', nargs='*')
    ap.add_argument('--transforms', help='comma-separated transform names (default: all)')
    ap.add_argument('--json', help='write the full records and candidate list to this path')
    ap.add_argument('--all', action='store_true', help='every desk task (slow: LibreOffice per workbook)')
    ap.add_argument('--list-transforms', action='store_true')
    ap.add_argument('-q', '--quiet', action='store_true')
    a = ap.parse_args(argv)
    if a.list_transforms:
        for t in TRANSFORMS.values():
            print(f'{t.name:24} {t.kind:9} {t.default_class:18} {t.doc}' +
                  (f'\n{"":53}rule: {RULES[t.name]}' if t.name in RULES else ''))
        return 0
    tnames = list(TRANSFORMS)
    if a.transforms:
        tnames = [x.strip() for x in a.transforms.split(',') if x.strip()]
        bad = [x for x in tnames if x not in TRANSFORMS]
        if bad:
            ap.error(f'unknown transforms {bad}; see --list-transforms')
    if a.all:
        ids = sorted(d for d in os.listdir(DESK) if os.path.isfile(os.path.join(DESK, d, 'task.yaml')))
    else:
        ids = a.tasks or DEFAULT_SAMPLE
    log = None if a.quiet else (lambda s: print(s, file=sys.stderr, flush=True))
    records = []
    for tid in ids:
        if not os.path.isfile(os.path.join(DESK, tid, 'task.yaml')):
            print(f'unknown task {tid}', file=sys.stderr); continue
        records.append(run_task(tid, tnames, log))
    cands = candidates(records)
    sens = sensitive_fails(records)
    print(matrix(records, tnames))
    print()
    for rec in records:
        if rec['baseline'] != 'pass':
            print(f'BASELINE {rec["baseline"]}: {rec["task"]} ' +
                  '; '.join(f'{f["name"]}: {f["detail"][:120]}' for f in rec.get('baseline_failed', [])))
    print(f'Candidate false negatives (invariant transform, fail): {len(cands)}')
    for c in cands:
        print(f'  {c["task"]} x {c["transform"]} [{", ".join(c["files"])}]\n'
              f'      check {c["check"]!r} ({c["check_type"]}): {c["detail"][:300]}')
    print(f'Contract-sensitive fails (information): {len(sens)}')
    for s in sens:
        print(f'  {s["task"]} x {s["transform"]}: {", ".join(s["checks"])}  ({s["reason"]})')
    if a.json:
        with open(a.json, 'w') as f:
            json.dump({'records': records, 'candidate_false_negatives': cands, 'contract_sensitive_fails': sens,
                       'transforms': {t.name: {'kind': t.kind, 'class': t.default_class, 'doc': t.doc,
                                               'rule': RULES.get(t.name)} for t in TRANSFORMS.values()}},
                      f, indent=2)
    if any(r['baseline'] == 'FAIL' for r in records):
        return 2
    return 1 if cands else 0


if __name__ == '__main__':
    sys.exit(main())
