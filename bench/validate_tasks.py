#!/usr/bin/env python3
"""Sanity-check every desk task: grader must PASS on reference_solution and FAIL on the untouched workspace.
Usage: validate_tasks.py [task_id ...]"""
import os, sys, json, hashlib, shutil, subprocess
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from grade import grade
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
root = os.path.join(ROOT, 'tasks', 'desk')
strict = '--strict' in sys.argv
collisions_only = '--collisions' in sys.argv
ids = [a for a in sys.argv[1:] if not a.startswith('--')] or sorted(d for d in os.listdir(root) if os.path.isfile(os.path.join(root, d, 'task.yaml')))

def tree_hash(d: str) -> str:
    h = hashlib.sha256()
    for base, _, files in sorted(os.walk(d)):
        for f in sorted(files):
            p = os.path.join(base, f); h.update(os.path.relpath(p, d).encode()); h.update(open(p, 'rb').read())
    return h.hexdigest()


def loose_pins(td: str) -> list[str]:
    """Workbook figures computed from exact source data tie to the cent: tolerance at most 1.00.
    A figure that involves currency conversion, proration, or a documented estimate declares
    `rounding: <reason>` on the check and may then use rel_tol up to 0.001."""
    import yaml as _yaml
    spec_all = _yaml.safe_load(open(os.path.join(td, 'task.yaml'))) or {}
    out = []
    for c in spec_all.get('checks') or []:
        if c.get('type') != 'xlsx_value_present': continue
        exp = abs(float(c['expected'])); rel = float(c.get('rel_tol', 0.01)); tol = max(exp * rel, 0.01)
        if c.get('rounding'):
            if rel > 0.001: out.append(f"{c.get('name')!r}: rel_tol {rel} exceeds 0.001 even with rounding: {c['rounding']}")
        elif tol > 1.0:
            out.append(f"{c.get('name')!r}: tolerance {tol:.2f} on {exp} exceeds 1.00; tighten rel_tol, or set rounding: <reason> if conversion or proration makes cents ambiguous")
    return out

def raw_collisions(td: str) -> list[str]:
    """A pinned workbook figure must not already sit in the raw inputs on a row or column carrying its label:
    an agent that pastes the export into its workbook would pass that check without doing the work."""
    import csv as _csv, glob as _glob, yaml as _yaml
    from grade import cell_num as to_num
    spec_all = _yaml.safe_load(open(os.path.join(td, 'task.yaml'))) or {}
    pins = [c for c in spec_all.get('checks') or [] if c.get('type') == 'xlsx_value_present' and not c.get('raw_value_ok')]
    if not pins: return []
    grids = []
    for p in sorted(_glob.glob(os.path.join(td, 'workspace', '**', '*'), recursive=True)):
        low = p.lower()
        try:
            if low.endswith(('.csv', '.tsv', '.txt')) and low.endswith(('.csv', '.tsv')):
                text = open(p, 'rb').read().decode('utf-8', errors='replace').replace('\r\n', '\n')
                delim = '\t' if low.endswith('.tsv') else ','
                grids.append((os.path.basename(p), [row for row in _csv.reader(text.splitlines(), delimiter=delim)]))
            elif low.endswith(('.xlsx', '.xlsm')):
                from openpyxl import load_workbook
                wb = load_workbook(p, data_only=True, read_only=True)
                for sh in wb.worksheets:
                    grids.append((f'{os.path.basename(p)}!{sh.title}', [[c for c in row] for row in sh.iter_rows(values_only=True)]))
        except Exception:
            continue
    hits = []
    for spec in pins:
        expected = float(spec['expected']); rel = float(spec.get('rel_tol', 0.01)); tol = max(abs(expected) * rel, 0.01)
        near = str(spec.get('near_text', '')).lower()
        for name, grid in grids:
            if not grid: continue
            ncols = max(len(r) for r in grid)
            col_text = [' '.join(str(r[j]).lower() for r in grid if j < len(r) and r[j] not in (None, '')) for j in range(ncols)]
            found = None
            for i, r in enumerate(grid):
                row_text = ' '.join(str(v).lower() for v in r if v not in (None, ''))
                for j, v in enumerate(r):
                    if v in (None, '') or isinstance(v, bool): continue
                    x = to_num(v)
                    if x is None or abs(x - expected) > tol: continue
                    if not near or near in row_text or near in col_text[j]:
                        found = f'{name} row {i + 1} col {j + 1} = {v!r}'; break
                if found: break
            if found:
                hits.append(f'{spec.get("name")!r} (expected {expected}, near {near!r}) already in raw input: {found}')
                break
    return hits

bad = 0
for t in ids:
    td = os.path.join(root, t)
    if collisions_only:
        hs = raw_collisions(td)
        print(f"{'BAD' if hs else 'ok ':3} {t}" + ''.join(f'\n      raw-input collision: {h[:220]}' for h in hs)); bad += bool(hs); continue
    if strict:
        # Regenerate in a scratch sibling (same depth, so ../../lib imports resolve), never in place:
        # seed variants and hand-edited tasks must not be overwritten by a validation run.
        gen = os.path.join(td, 'gen.py')
        if not os.path.exists(gen):
            print(f'BAD {t:22} no gen.py'); bad += 1; continue
        scratch = os.path.join(root, f'_validate_{t}')
        hashes, failed = [], False
        for _ in range(2):
            shutil.rmtree(scratch, ignore_errors=True)
            shutil.copytree(td, scratch, ignore=shutil.ignore_patterns('__pycache__'))
            seed_args = []
            try:
                import yaml as _y
                seed = (_y.safe_load(open(os.path.join(td, 'task.yaml'))) or {}).get('generator_seed')
                if seed is not None: seed_args = ['--seed', str(seed)]
            except Exception: pass
            r = subprocess.run([sys.executable, os.path.join(scratch, 'gen.py'), *seed_args], capture_output=True, text=True, cwd=scratch)
            if r.returncode != 0:
                print(f'BAD {t:22} gen.py failed: {r.stderr.strip()[-200:]}'); failed = True; break
            hashes.append(tuple(tree_hash(os.path.join(scratch, d)) for d in ('workspace', 'reference', 'reference_solution')))
        committed = tuple(tree_hash(os.path.join(td, d)) for d in ('workspace', 'reference', 'reference_solution'))
        shutil.rmtree(scratch, ignore_errors=True)
        if failed: bad += 1; continue
        if hashes[0] != hashes[1]:
            print(f'BAD {t:22} gen.py is not deterministic (two runs differ)'); bad += 1; continue
        if hashes[0] != committed:
            print(f'    {t:22} note: committed files differ from gen.py output (stale or hand-edited; set generator_seed in task.yaml for seed variants)')
    sol = os.path.join(td, 'reference_solution')
    ws = os.path.join(td, 'workspace')
    g_sol = grade(td, sol) if os.path.isdir(sol) else None
    g_ws = grade(td, ws)
    ok_sol = bool(g_sol and g_sol['passed'])
    ok_ws = not g_ws['passed']
    status = 'ok' if (ok_sol and ok_ws) else 'BAD'
    n_checks = len(g_ws["checks"]); n_traps = 0
    try:
        import yaml; n_traps = len(yaml.safe_load(open(os.path.join(td, 'task.yaml'))).get('traps') or [])
    except Exception: pass
    LEGACY = {'customer-dedupe', 'expense-categorize', 'monthly-report', 'monthly-report-v9', 'payments-match', 'payments-match-v2', 'quote-comparison'}
    if strict and t not in LEGACY and (n_checks < 4 or n_traps < 3 or not os.path.isdir(sol)):
        status = 'BAD'
    collisions = raw_collisions(td) if strict else []
    if collisions and t not in LEGACY:
        status = 'BAD'
    loose = loose_pins(td) if strict and t not in LEGACY else []
    if loose:
        status = 'BAD'
    print(f'{status:3} {t:22} solution={"pass" if ok_sol else "FAIL" if os.path.isdir(sol) else "MISSING"} workspace={"fail" if ok_ws else "PASSES(!)"} checks={n_checks} traps={n_traps}')
    bad += status == 'BAD'
    for h in collisions: print(f'      raw-input collision: {h[:220]}')
    for h in loose: print(f'      loose tolerance: {h[:240]}')
    if g_sol and not g_sol['passed']:
        for c in g_sol['checks']:
            if not c['passed']: print(f'      solution failed: {c["name"]}: {c["detail"][:120]}')
sys.exit(1 if bad else 0)
