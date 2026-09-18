#!/usr/bin/env python3
"""Trajectory audit: why did each run end the way it did?

An LLM reviewer reads what the tested agent never saw (the task author's traps, the
check definitions, the reference answers, the grader's verdicts) next to what the agent
did see (its own trajectory, tool results, and deliverables), and classifies the run:

  failed runs   model | harness_guidance | harness_infra | eval | mixed
  passed runs   sound_pass | suspicious_pass

  audit.py results/<label> [...]          # failed runs of each label
  audit.py results/<label> --all          # every run (passes get the false-pass review)
  audit.py results/<label>/<run_dir>      # one run
  --model anthropic/claude-opus-5  --parallel 3  --force  --limit N  --dry-run (bundle only)

Writes <run>/audit-bundle.md (what the reviewer saw), <run>/audit.json (its verdict),
and results/<label>/audit-summary.{md,json}. Key: $OPENROUTER_API_KEY;
never printed and never read from a personal profile.
"""
from __future__ import annotations
import argparse, concurrent.futures as cf, glob, hashlib, json, os, re, sys, time, urllib.request
from collections import Counter, defaultdict
import yaml
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from grade import recalculated_workbook  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_MODEL = os.environ.get('BENCH_AUDIT_MODEL', 'anthropic/claude-opus-5')
SKIP_DIRS = {'.proto', '.codex', '.proto-logs', 'node_modules', '__pycache__', '.git', '.venv', '.npm'}
CAP_TOOL, CAP_TEXT, CAP_FINAL, CAP_SCRIPT = 2500, 5000, 4000, 3000
TOOL_FAIL_RE = re.compile(r'exit code [1-9]|"exitCode": ?(?!0\b)\d|Command failed|^ERROR|Error:|timed out|command not found|No such file', re.M)

# ---------- small helpers ----------

def clip(s: str, n: int) -> str:
    s = '' if s is None else str(s)
    return s if len(s) <= n else s[:n] + f'\n… [{len(s) - n} more chars]'

def sha(path: str) -> str:
    h = hashlib.sha1()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 16), b''):
            h.update(chunk)
    return h.hexdigest()

def api_key() -> str:
    key = os.environ.get('OPENROUTER_API_KEY')
    if not key:
        sys.exit('no OpenRouter key: set OPENROUTER_API_KEY through your secret manager')
    return key

def find_task_dir(task_id: str) -> str:
    for track in ('desk', 'build'):
        d = os.path.join(ROOT, 'tasks', track, task_id)
        if os.path.isfile(os.path.join(d, 'task.yaml')):
            return d
    sys.exit(f'task dir for {task_id!r} not found')

def walk_files(root: str) -> list[str]:
    out = []
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in SKIP_DIRS]
        for fn in fns:
            if fn.startswith('.DS_Store'): continue
            out.append(os.path.relpath(os.path.join(dp, fn), root))
    return sorted(out)

# ---------- file previews ----------

def preview_csv(path: str, rows: int = 30) -> str:
    raw = open(path, 'rb').read().decode('utf-8', errors='replace')
    lines = raw.replace('\r\r\n', '\n').replace('\r\n', '\n').replace('\r', '\n').split('\n')
    lines = [l for l in lines if l != '']
    head = '\n'.join(lines[:rows + 1])
    return f'{max(len(lines) - 1, 0)} data rows\n{head}' + (f'\n… [{len(lines) - rows - 1} more rows]' if len(lines) > rows + 1 else '')

def workbook_view(path: str, max_rows: int = 30, max_cols: int = 14) -> tuple[str, dict]:
    """Formulas next to their computed values, plus a deterministic error scan."""
    from openpyxl import load_workbook
    info = {'engine': 'none', 'error_cells': [], 'uncomputed_formula_cells': 0, 'formula_cells': 0}
    try:
        wb_f = load_workbook(path)
    except Exception as e:
        return f'[cannot open workbook: {type(e).__name__}: {e}]', info
    rpath = recalculated_workbook(path)
    info['engine'] = 'libreoffice' if rpath != path and '/recalc-' in rpath and 'recalc-py' not in rpath else ('formulas-py' if rpath != path else 'cached-values-only')
    try:
        wb_v = load_workbook(rpath, data_only=True)
    except Exception:
        wb_v = None
    parts = []
    for sh in wb_f.worksheets:
        vsh = wb_v[sh.title] if wb_v is not None and sh.title in wb_v.sheetnames else None
        has_formulas = any(isinstance(c.value, str) and c.value.startswith('=') for row in sh.iter_rows() for c in row)
        # Raw data sheets get a glimpse; the sheets that compute something get the window.
        sheet_rows = max_rows if has_formulas or sh.max_row <= 60 else 6
        parts.append(f'--- sheet {sh.title!r} dims {sh.dimensions}{"" if has_formulas else " (no formulas)"} ---')
        shown = 0
        for r_i, row in enumerate(sh.iter_rows(), start=1):
            cells = []
            for c in row[:max_cols]:
                v = c.value
                if v is None: cells.append(''); continue
                if isinstance(v, str) and v.startswith('='):
                    info['formula_cells'] += 1
                    val = vsh[c.coordinate].value if vsh is not None else None
                    if val is None:
                        info['uncomputed_formula_cells'] += 1
                    elif isinstance(val, str) and val.startswith('#'):
                        info['error_cells'].append(f'{sh.title}!{c.coordinate} {v} → {val}')
                    cells.append(f'{v} → {val!r}')
                else:
                    cells.append(repr(v) if not isinstance(v, str) else v)
            if any(cells):
                parts.append(f'r{r_i}: ' + ' | '.join(cells))
                shown += 1
            if shown >= sheet_rows:
                parts.append(f'… [rows beyond {r_i} not shown; sheet has {sh.max_row} rows]'); break
        # Error scan over the whole sheet, not just the shown window
        for row in sh.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('=') and vsh is not None:
                    val = vsh[c.coordinate].value
                    if isinstance(val, str) and val.startswith('#'):
                        tag = f'{sh.title}!{c.coordinate} {v} → {val}'
                        if tag not in info['error_cells']: info['error_cells'].append(tag)
    info['error_cells'] = info['error_cells'][:40]
    parts.insert(0, f'[values computed by: {info["engine"]}; formula cells {info["formula_cells"]}, uncomputed {info["uncomputed_formula_cells"]}, error cells {len(info["error_cells"])}]')
    if info['error_cells']:
        parts.append('ERROR CELLS: ' + '; '.join(info['error_cells'][:20]))
    return clip('\n'.join(parts), 9000), info

def preview_file(path: str) -> tuple[str, dict]:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in ('.xlsx', '.xlsm'):
            return workbook_view(path)
        if ext == '.csv':
            return preview_csv(path), {}
        if ext in ('.md', '.txt', '.json', '.yaml', '.yml', '.html'):
            return clip(open(path, encoding='utf-8', errors='replace').read(), CAP_TEXT), {}
        if ext in ('.py', '.mjs', '.js', '.sh', '.ts'):
            return clip(open(path, encoding='utf-8', errors='replace').read(), CAP_SCRIPT), {}
        return f'[binary or unknown type, {os.path.getsize(path)} bytes]', {}
    except Exception as e:
        return f'[preview failed: {type(e).__name__}: {e}]', {}

# ---------- trajectories ----------

def proto_trajectory(ws: str) -> tuple[list[dict], dict]:
    """Rebuild the model's view from Proto's full request/response logs. Each request re-sends
    the history, so steps are taken first-seen in order (compaction-safe) and tool results keep
    their longest (least truncated) version."""
    files = sorted(glob.glob(os.path.join(ws, '.proto-logs', '*.full.jsonl')))
    sessions = []
    stats = {'requests': 0, 'finish_length': 0, 'malformed_tool_calls': 0, 'tool_calls': 0, 'tool_failures': 0,
             'asked_user': False, 'skills': [], 'subagents': [], 'profiler_ran': False, 'models': Counter(), 'errors': []}
    # One log file holds the main agent and every sub-agent it spawned; split by runtime session
    # so the specialist's turns are not read as the main agent's.
    groups: dict[str, list[dict]] = defaultdict(list); labels: dict[str, str] = {}; req_session: dict[str, str] = {}
    for f in files:
        for line in open(f, encoding='utf-8', errors='replace'):
            try: r = json.loads(line)
            except Exception: continue
            md = r.get('metadata') or {}
            if r.get('type') == 'REQUEST_FULL':
                sid = md.get('runtimeSessionId') or md.get('sessionId') or os.path.basename(f)
                req_session[r.get('requestId')] = sid
                labels.setdefault(sid, md.get('agentLabel') or 'main')
            elif r.get('type') == 'RESPONSE_FULL':
                sid = req_session.get(r.get('requestId')) or os.path.basename(f)
            else:
                continue
            groups[sid].append(r)
    for sid, recs in groups.items():
        label = labels.get(sid, sid)
        recs.sort(key=lambda r: r.get('timestamp', ''))
        steps, seen, results, first_ts = [], set(), {}, None
        responses = [r for r in recs if r.get('type') == 'RESPONSE_FULL']
        for r in recs:
            if r.get('type') != 'REQUEST_FULL': continue
            first_ts = first_ts or r.get('timestamp')
            stats['requests'] += 1; stats['models'][r.get('model')] += 1
            for m in (r.get('requestBody') or {}).get('messages') or []:
                role = m.get('role')
                if role == 'tool':
                    cid = m.get('tool_call_id') or ''
                    c = m.get('content'); c = c if isinstance(c, str) else json.dumps(c)
                    if len(c) > len(results.get(cid, '')): results[cid] = c
                    continue
                if role == 'system': continue
                if role == 'user':
                    c = m.get('content'); c = c if isinstance(c, str) else json.dumps(c)
                    key = 'u:' + hashlib.sha1(c[:2000].encode()).hexdigest()
                    if key in seen: continue
                    seen.add(key)
                    steps.append({'kind': 'prompt' if len(steps) == 0 else 'context', 'text': c})
                    continue
                if role == 'assistant':
                    calls = m.get('tool_calls') or []
                    key = 'a:' + (','.join(tc.get('id', '') for tc in calls) or hashlib.sha1((m.get('content') or '')[:2000].encode()).hexdigest())
                    if key in seen: continue
                    seen.add(key)
                    steps.append({'kind': 'assistant', 'text': m.get('content') or '',
                                  'calls': [{'id': tc.get('id'), 'name': (tc.get('function') or {}).get('name'),
                                             'args': (tc.get('function') or {}).get('arguments')} for tc in calls]})
        # The final response is not in any request body
        if responses:
            last = responses[-1].get('response') or {}
            calls = last.get('functionCalls') or []
            steps.append({'kind': 'assistant', 'text': last.get('text') or '', 'final': True,
                          'calls': [{'id': c.get('id'), 'name': c.get('name'), 'args': json.dumps(c.get('args'))} for c in calls]})
        for r in responses:
            resp = r.get('response') or {}
            if resp.get('finishReason') == 'length': stats['finish_length'] += 1
            stats['malformed_tool_calls'] += int(resp.get('malformedToolCallCount') or 0)
        # Attach results and reasoning; derive stats
        ri = 0
        for s in steps:
            if s['kind'] != 'assistant': continue
            if ri < len(responses):
                resp = responses[ri].get('response') or {}
                if (resp.get('text') or '') == (s['text'] or ''):
                    rs = resp.get('reasoning')
                    if rs: s['reasoning'] = rs if isinstance(rs, str) else json.dumps(rs)
                ri += 1
            for c in s.get('calls', []):
                stats['tool_calls'] += 1
                res = results.get(c.get('id') or '')
                c['result'] = res
                if res is not None and TOOL_FAIL_RE.search(res[:4000]): c['failed'] = True; stats['tool_failures'] += 1
                a = c.get('args') or ''
                if c['name'] == 'ask_user_question': stats['asked_user'] = True
                if c['name'] == 'use_skill':
                    try: stats['skills'].append(json.loads(a).get('name'))
                    except Exception: stats['skills'].append(a[:60])
                if c['name'] == 'task':
                    try: stats['subagents'].append(json.loads(a).get('subagent_type'))
                    except Exception: stats['subagents'].append('task')
                if c['name'] in ('bash', 'execute_script') and 'profile.mjs' in a: stats['profiler_ran'] = True
        sessions.append({'label': label, 'first_ts': first_ts or '', 'steps': steps})
    sessions.sort(key=lambda s: s['first_ts'])
    stats['models'] = dict(stats['models'])
    return sessions, stats

def codex_trajectory(out: str) -> tuple[list[dict], dict]:
    stats = {'requests': None, 'tool_calls': 0, 'tool_failures': 0, 'asked_user': False, 'skills': [], 'subagents': [],
             'profiler_ran': False, 'finish_length': 0, 'malformed_tool_calls': 0, 'models': {}, 'errors': []}
    p = os.path.join(out, 'codex_events.jsonl')
    steps = []
    if not os.path.exists(p):
        return [{'label': 'codex', 'first_ts': '', 'steps': steps}], stats
    for line in open(p, encoding='utf-8', errors='replace'):
        try: r = json.loads(line)
        except Exception: continue
        t = r.get('type'); item = r.get('item') or {}
        if t == 'item.completed':
            it = item.get('type')
            if it == 'agent_message':
                steps.append({'kind': 'assistant', 'text': item.get('text') or '', 'calls': []})
            elif it == 'reasoning':
                steps.append({'kind': 'assistant', 'text': '', 'reasoning': item.get('text') or '', 'calls': []})
            elif it == 'command_execution':
                stats['tool_calls'] += 1
                failed = (item.get('exit_code') not in (0, None))
                if failed: stats['tool_failures'] += 1
                steps.append({'kind': 'assistant', 'text': '', 'calls': [{'name': 'shell', 'args': item.get('command'),
                              'result': f"exit {item.get('exit_code')}\n{item.get('aggregated_output') or ''}", 'failed': failed}]})
            elif it == 'file_change':
                ch = ', '.join(f"{c.get('kind')} {os.path.basename(c.get('path', ''))}" for c in item.get('changes') or [])
                steps.append({'kind': 'assistant', 'text': '', 'calls': [{'name': 'file_change', 'args': ch, 'result': 'ok'}]})
        elif t == 'error' or (t or '').endswith('.failed'):
            stats['errors'].append(json.dumps(r)[:300])
    if steps: steps[-1]['final'] = True
    lm = os.path.join(out, 'last_message.txt')
    if os.path.exists(lm):
        steps.append({'kind': 'assistant', 'text': open(lm, encoding='utf-8', errors='replace').read(), 'calls': [], 'final': True})
    return [{'label': 'codex', 'first_ts': '', 'steps': steps}], stats

def render_trajectory(sessions: list[dict]) -> str:
    out, n = [], 0
    for s in sessions:
        out.append(f'\n#### session {s["label"]}')
        for st in s['steps']:
            n += 1
            if st['kind'] == 'prompt':
                out.append(f'[{n}] USER PROMPT: {clip(st["text"], 1500)}'); continue
            if st['kind'] == 'context':
                out.append(f'[{n}] CONTEXT MESSAGE: {clip(st["text"], 800)}'); continue
            tag = 'FINAL ANSWER' if st.get('final') else 'assistant'
            if st.get('reasoning'):
                out.append(f'[{n}] {tag} (reasoning): {clip(st["reasoning"], 600)}')
            if st.get('text'):
                out.append(f'[{n}] {tag}: {clip(st["text"], CAP_FINAL if st.get("final") else 1500)}')
            for c in st.get('calls', []):
                out.append(f'[{n}] CALL {c.get("name")}: {clip(c.get("args"), 900)}')
                if c.get('result') is None:
                    out.append(f'[{n}] RESULT: (no result recorded: the run ended, or the call was cut off)')
                else:
                    out.append(f'[{n}] RESULT{" (FAILED)" if c.get("failed") else ""}: {clip(c["result"], CAP_TOOL)}')
    return '\n'.join(out)

# ---------- bundle ----------

def build_bundle(run_dir: str) -> tuple[str, dict]:
    res = json.load(open(os.path.join(run_dir, 'result.json')))
    task_dir = find_task_dir(res['task'])
    task = yaml.safe_load(open(os.path.join(task_dir, 'task.yaml')))
    if 'checks' not in res:
        # Build-track run: the automated items from build_grade.py stand in for the grader's checks,
        # the tester checklist for the check definitions, and RESULT.json presence for the pass flag.
        gp = os.path.join(run_dir, 'grade.json')
        auto = json.load(open(gp))['automated'] if os.path.exists(gp) else []
        res['checks'] = [{'name': f"item {a['item']}: {a['name']}", 'passed': a['passed'], 'detail': a['detail']} for a in auto]
        res['passed'] = bool(res.get('result_json')) and all(a['passed'] for a in auto) and not res.get('timed_out')
        if 'checklist' in task and 'checks' not in task:
            cl = os.path.join(task_dir, task['checklist'])
            task['checks'] = [{'type': 'tester checklist', 'file': task['checklist']}]
            task.setdefault('traps', []).append('BUILD TRACK: a person clicks through the checklist below against the URL in RESULT.json; the automated items are only the contract fields and the URL answering. Judge the trajectory on whether a usable app with the seed data imported and both logins exists at the URL, and on how the hour was spent.')
            if os.path.exists(cl):
                task['traps'].append('Checklist:\n' + clip(open(cl, encoding='utf-8', errors='replace').read(), 4000))
        if res.get('result_json'):
            task.setdefault('traps', []).append('RESULT.json the agent wrote: ' + json.dumps(res['result_json'])[:1500])
    ws, out = os.path.join(run_dir, 'ws'), os.path.join(run_dir, 'out')
    fam = 'codex' if res['harness'].startswith('codex') else 'proto'
    sessions, stats = (codex_trajectory(out) if fam == 'codex' else proto_trajectory(ws))
    # Deliverables: files the agent created or changed
    orig_root = os.path.join(task_dir, 'workspace')
    orig = {rel: sha(os.path.join(orig_root, rel)) for rel in walk_files(orig_root)}
    delivered, wb_info = [], {}
    for rel in walk_files(ws):
        p = os.path.join(ws, rel)
        if rel in orig and sha(p) == orig[rel]: continue
        delivered.append(rel)
    delivered = delivered[:14]
    dparts = []
    for rel in delivered:
        text, info = preview_file(os.path.join(ws, rel))
        if info: wb_info[rel] = info
        dparts.append(f'### {rel} ({"modified input" if rel in orig else "new"})\n{text}')
    # Reference
    rparts = []
    ref_dir = os.path.join(task_dir, 'reference')
    if not os.path.isdir(orig_root): orig_root = os.path.join(task_dir, 'seed') if os.path.isdir(os.path.join(task_dir, 'seed')) else orig_root
    for rel in walk_files(ref_dir) if os.path.isdir(ref_dir) else []:
        p = os.path.join(ref_dir, rel)
        rparts.append(f'### reference/{rel}\n{preview_csv(p, 40) if rel.endswith(".csv") else preview_file(p)[0]}')
    sol_memo = os.path.join(task_dir, 'reference_solution', 'memo.md')
    if os.path.exists(sol_memo):
        rparts.append('### reference_solution/memo.md\n' + clip(open(sol_memo, encoding='utf-8', errors='replace').read(), 3000))
    custom = os.path.join(task_dir, 'check.py')
    checks_yaml = yaml.safe_dump(task.get('checks', []), sort_keys=False, width=120)
    if os.path.exists(custom):
        checks_yaml += '\n# custom check.py\n' + clip(open(custom).read(), CAP_SCRIPT)
    verdicts = '\n'.join(f'- [{"PASS" if c["passed"] else "FAIL"}]{"" if c.get("required", True) else " (optional)"} {c["name"]} :: {c["detail"]}' for c in res['checks'])
    tails = []
    for fn in ('stdout.txt', 'stderr.txt'):
        p = os.path.join(run_dir, fn)
        if os.path.exists(p):
            t = open(p, encoding='utf-8', errors='replace').read()
            if t.strip(): tails.append(f'### {fn} (last 1500 chars)\n{t[-1500:]}')
    signals = {
        'passed': res['passed'], 'timed_out': res.get('timed_out'), 'exit_code': res.get('exit_code'),
        'wall_min': round(res.get('wall_s', 0) / 60, 1), 'failed_checks': [c['name'] for c in res['checks'] if not c['passed']],
        'missing_deliverables': sorted({c['name'] for c in res['checks'] if 'output file missing' in c.get('detail', '')}),
        'workbook_error_cells': {k: v['error_cells'] for k, v in wb_info.items() if v.get('error_cells')},
        'workbook_uncomputed_formula_cells': {k: v['uncomputed_formula_cells'] for k, v in wb_info.items() if v.get('uncomputed_formula_cells')},
        'workbook_value_engine': {k: v['engine'] for k, v in wb_info.items()},
        **{k: v for k, v in stats.items() if k != 'errors'}, 'harness_errors': stats.get('errors', [])[:5],
        'usage': res.get('usage', {}),
    }
    bundle = f"""# Run {res['run_id']}  (harness {res['harness']}, task {res['task']}, category {res.get('category')})
Grader outcome: {'PASS' if res['passed'] else 'FAIL'}; wall {signals['wall_min']} min; timed out {res.get('timed_out')}; exit code {res.get('exit_code')}

## The ask (all the agent was told)
{task['ask'].strip()}

## PRIVILEGED: task author's traps (never shown to the agent)
{chr(10).join('- ' + str(t).strip() for t in task.get('traps', [])) or '(none listed)'}

## PRIVILEGED: check definitions
```yaml
{checks_yaml}```

## Grader verdicts for this run
{verdicts}

## PRIVILEGED: reference answers
{chr(10).join(rparts) or '(none)'}

## Deliverables the agent left in the workspace
{chr(10).join(dparts) or '(nothing new or changed in the workspace)'}

## Deterministic signals
```json
{json.dumps(signals, indent=1, default=str)}
```

## Trajectory (numbered steps; RESULT lines are what the model saw back)
{render_trajectory(sessions)}

## Harness process output
{chr(10).join(tails) or '(empty)'}
"""
    return bundle, signals

# ---------- judge ----------

SYSTEM = """You are the trace-analysis reviewer for a benchmark of AI agent harnesses doing business desk work: cleaning exports, matching payments, categorising expenses, building workbooks, writing memos. You see everything the tested agent never saw: the task author's traps, the check definitions, the reference answers, and the grader's verdicts. Decide WHY the run ended the way it did, so the authors know whether to fix the model, the harness, or the eval.

Classify a FAILED run into exactly one primary class:
- model: the agent had working tools, a clear enough ask, and the needed guidance, and still produced wrong or incomplete output, or made a judgment call the ask or ordinary business convention settles.
- harness_guidance: the harness's prompt, skill, specialist, or staged script steered the model wrong or left out a rule, check, or step the model needed; the same model would likely have passed with better guidance. Name the guidance.
- harness_infra: a tool, sandbox, dependency, network, timeout, API error, output truncation, permission prompt, or runner bug stopped an otherwise capable agent.
- eval: the deliverable would satisfy a reasonable business owner but a check rejected it (verifier too strict, wrong, or mis-targeted), or the ask is ambiguous so the agent's reading is defensible, or the reference answer is wrong.
- mixed: two classes contributed materially; list them in secondary_classes with the primary first.

Classify a PASSED run by what a business owner would find if they used the deliverable today.
- suspicious_pass when any of these holds, whether or not a check looks at it:
  (1) a deliverable contains a wrong figure, a wrong entity, or a statement about the data that is false or misleading, however small (a memo total off by $40; a memo calling a region the weakest over a period in which one of its months is missing);
  (2) something the ask or the folder's notes require is missing or wrong;
  (3) the agent used leaked information or a shortcut that would not work on a fresh copy of the task;
  (4) the agent's final message or a deliverable says it verified, reconciled, or checked something it did not actually verify, or reports a result that contradicts its own deliverable.
- Three cases are settled; apply them as written:
  - Missing periods. A cell for a group and period with no source rows that displays 0 is a wrong figure under (1), however it is shaded, commented, or footnoted. Blank, "no data", "n/a", or a similar label is correct.
  - Comparisons over unequal coverage. A ranking or comparison of groups whose totals cover different numbers of periods is misleading under (1) unless both hold: (a) the same sentence, bullet, or table row says which group is short and by how many periods; (b) any rank or label it draws (weakest, smallest, last, lead) still holds when that group's missing periods are filled at its own average for the periods it has. A caveat elsewhere in the document does not satisfy (a).
  - What counts as verification. A claim that a deliverable, or the figures in it, were verified, reconciled, or checked needs a step that read the saved deliverable back: its cell values, its formulas' results as computed by a spreadsheet engine or formula evaluator (not cached values the builder wrote itself), or its text. Recomputing the same figures from the source in code or by hand checks the arithmetic, not the deliverable; a claim about the deliverable that rests only on that meets (4) even when the figures turn out right. A claim that says what was actually done ("totals recomputed from the cleaned data match the profile") is accurate.
- sound_pass otherwise. These do not make a pass suspicious; record them in evidence and root_cause instead:
  - a reasonable choice on something the ask and notes leave open (where a middle initial goes when splitting names, casing, number formatting, column order beyond a required template);
  - a slip in the agent's narration about how the work went (a miscounted step, a wrong arithmetic aside) that neither describes the deliverable's contents wrongly nor claims a verification;
  - extra files or extra columns that do not contradict the requested ones.
- When unsure, ask whether an owner acting on the deliverable would be misled or would have to fix something first: yes is suspicious_pass, no is sound_pass.

Rules:
- Every claim about the trajectory cites a step number in [brackets] and a short quote.
- Be adversarial toward the grader on failures and toward the agent on passes.
- A verifier that rejects a genuinely wrong answer is not an eval flaw, even if a different check design would be nicer. Check the reference numbers yourself against the deliverable before blaming the check.
- Conventions (duplicates count once, refunds net, gaps are not zero, dates disambiguated by day > 12) count against the agent when the ask or loaded guidance stated them; count as harness_guidance when nothing stated them and a careful colleague would also have needed telling.
- Formula cells that evaluate to errors, use non-existent functions, or are circular are the agent's fault unless a tool or specialist wrote them without the agent ever seeing the computed result; then say which and classify harness_guidance.
- If the agent asked a question, offered to do more instead of doing it, or claimed a check it never ran, record it.
Respond with one JSON object only, no prose around it, with exactly these keys:
{"verdict": "model|harness_guidance|harness_infra|eval|mixed|sound_pass|suspicious_pass",
 "secondary_classes": ["..."],
 "confidence": 0.0,
 "root_cause": "one paragraph",
 "first_divergence": {"step": 0, "what": "..."},
 "evidence": [{"step": 0, "quote": "...", "why_it_matters": "..."}],
 "checks": [{"name": "...", "grader_verdict": "pass|fail", "reviewer_verdict": "agent_wrong|check_wrong|ambiguous|agree_pass|false_pass", "why": "..."}],
 "agent_behaviour": {"asked_user": false, "offered_instead_of_doing": false, "claimed_unverified": false, "used_leaked_info": false, "left_process_running": false},
 "recommended_fix": {"target": "task|grader|reference|skill|specialist|tool|runner|model|none", "change": "..."},
 "one_line": "at most 140 characters for the summary table"}"""

_price_cache: dict[str, dict] = {}

def model_prices(model: str, key: str) -> dict:
    if model in _price_cache: return _price_cache[model]
    try:
        req = urllib.request.Request('https://openrouter.ai/api/v1/models', headers={'Authorization': f'Bearer {key}'})
        data = json.load(urllib.request.urlopen(req, timeout=30))['data']
        for m in data:
            if m['id'] == model:
                p = m.get('pricing', {})
                _price_cache[model] = {k: float(p.get(k) or 0) for k in ('prompt', 'completion', 'input_cache_read')}
                break
    except Exception:
        pass
    return _price_cache.setdefault(model, {})

def judge(bundle: str, passed: bool, model: str, key: str) -> dict:
    user = ('This run PASSED the grader. Review it for a false pass: is the deliverable actually right, did the agent do the work, '
            'did it claim anything it did not verify?' if passed else
            'This run FAILED the grader. Find the first point where it went wrong and classify the cause.') + \
           '\n\n' + bundle
    body = {'model': model, 'temperature': 0, 'max_tokens': 6000,
            'messages': [{'role': 'system', 'content': SYSTEM}, {'role': 'user', 'content': user}],
            'response_format': {'type': 'json_object'}, 'reasoning': {'effort': 'medium'}}
    req = urllib.request.Request('https://openrouter.ai/api/v1/chat/completions', data=json.dumps(body).encode(),
                                 headers={'Authorization': f'Bearer {key}', 'Content-Type': 'application/json',
                                          'HTTP-Referer': 'https://erp.ai', 'X-Title': 'business-bench audit'})
    t0 = time.time()
    last_err = None
    for attempt in range(3):
        try:
            resp = json.load(urllib.request.urlopen(req, timeout=600))
            break
        except Exception as e:
            last_err = e; time.sleep(5 * (attempt + 1))
    else:
        return {'verdict': 'error', 'error': f'{type(last_err).__name__}: {last_err}'}
    msg = (resp.get('choices') or [{}])[0].get('message') or {}
    content = msg.get('content') or ''
    text = content if isinstance(content, str) else ''.join(p.get('text', '') for p in content if isinstance(p, dict))
    m = re.search(r'\{.*\}', text, re.S)
    try:
        verdict = json.loads(m.group(0) if m else text)
    except Exception:
        verdict = {'verdict': 'unparsed', 'raw': text[:4000]}
    u = resp.get('usage') or {}
    prices = model_prices(model, key)
    cached = int((u.get('prompt_tokens_details') or {}).get('cached_tokens') or 0)
    pin, pout = int(u.get('prompt_tokens') or 0), int(u.get('completion_tokens') or 0)
    cost = None
    if prices:
        cost = round((pin - cached) * prices['prompt'] + cached * prices.get('input_cache_read', prices['prompt']) + pout * prices['completion'], 4)
    verdict['_meta'] = {'model': resp.get('model') or model, 'prompt_tokens': pin, 'cached_tokens': cached, 'completion_tokens': pout,
                        'cost_usd': cost, 'seconds': round(time.time() - t0, 1), 'audited_at': time.strftime('%Y-%m-%dT%H:%M:%S')}
    return verdict

def judge_cli(bundle: str, passed: bool, model: str) -> dict:
    """The same review through the Claude Code CLI on a Claude subscription: the reviewer system prompt replaces
    Claude Code's own, tools are off, the bundle goes in on stdin, and the reply is parsed exactly as the API path.
    Cost is recorded as None (subscription) with the CLI's API-equivalent figure kept beside it."""
    import subprocess
    user = ('This run PASSED the grader. Review it for a false pass: is the deliverable actually right, did the agent do the work, '
            'did it claim anything it did not verify?' if passed else
            'This run FAILED the grader. Find the first point where it went wrong and classify the cause.') + \
           '\n\n' + bundle + '\n\nReply with the single JSON object only.'
    cli_model = model.split('/')[-1]  # anthropic/claude-opus-5 -> claude-opus-5
    cmd = ['claude', '-p', '--model', cli_model, '--system-prompt', SYSTEM, '--tools', '', '--output-format', 'json']
    t0 = time.time(); last = None
    for attempt in range(3):
        try:
            r = subprocess.run(cmd, input=user, capture_output=True, text=True, timeout=900)
            d = json.loads(r.stdout)
            if d.get('is_error'):
                last = d.get('result') or 'cli error'
                if any(k in str(last).lower() for k in ('authenticate', 'oauth', 'not logged in', '/login')):
                    return {'verdict': 'error', 'error': f'claude CLI not signed in: {last}'}
                time.sleep(20 * (attempt + 1)); continue
            break
        except Exception as e:
            last = f'{type(e).__name__}: {e}'; time.sleep(10 * (attempt + 1))
    else:
        return {'verdict': 'error', 'error': str(last)[:500]}
    text = d.get('result') or ''
    m = re.search(r'\{.*\}', text, re.S)
    try:
        verdict = json.loads(m.group(0) if m else text)
    except Exception:
        verdict = {'verdict': 'unparsed', 'raw': text[:4000]}
    u = d.get('usage') or {}
    # modelUsage also lists Claude Code's own small Haiku side call; the reviewer is the model that wrote the most tokens
    mu = d.get('modelUsage') or {}
    served = sorted(mu, key=lambda k: -(mu[k].get('outputTokens') or 0))
    verdict['_meta'] = {'model': served[0] if served else cli_model, 'served_models': served, 'backend': 'claude-cli',
                        'prompt_tokens': int(u.get('input_tokens') or 0) + int(u.get('cache_read_input_tokens') or 0) + int(u.get('cache_creation_input_tokens') or 0),
                        'cached_tokens': int(u.get('cache_read_input_tokens') or 0), 'completion_tokens': int(u.get('output_tokens') or 0),
                        'cost_usd': None, 'api_equivalent_usd': d.get('total_cost_usd'),
                        'seconds': round(time.time() - t0, 1), 'audited_at': time.strftime('%Y-%m-%dT%H:%M:%S')}
    return verdict

# ---------- driver ----------

def audit_run(run_dir: str, model: str, key: str | None, force: bool, dry: bool, backend: str = 'openrouter', from_bundles: bool = False, out_name: str = 'audit.json') -> dict | None:
    outp = os.path.join(run_dir, out_name)
    if os.path.exists(outp) and not force and not dry:
        return json.load(open(outp))
    res = json.load(open(os.path.join(run_dir, 'result.json')))
    bj = os.path.join(run_dir, 'audit-bundle.json')
    if from_bundles:
        if not os.path.exists(bj):
            print(f'[no bundle] {os.path.basename(run_dir)}', flush=True); return None
        b = json.load(open(bj)); bundle, signals = b['bundle'], b['signals']
    else:
        bundle, signals = build_bundle(run_dir)
        open(os.path.join(run_dir, 'audit-bundle.md'), 'w').write(bundle)
        # A portable copy: another machine (one signed in to the CLI reviewer) can review from this file alone.
        json.dump({'bundle': bundle, 'signals': signals, 'run_id': res['run_id'], 'task': res['task'], 'harness': res['harness'],
                   'passed': res.get('passed', signals.get('passed', False))}, open(bj, 'w'), default=str)
    if dry:
        print(f'[bundle] {os.path.basename(run_dir)}  {len(bundle):,} chars', flush=True); return None
    passed = res.get('passed', signals.get('passed', False))  # build-track runs derive it in build_bundle
    v = judge_cli(bundle, passed, model) if backend == 'claude-cli' else judge(bundle, passed, model, key)
    v['run_id'] = res['run_id']; v['run_dir'] = os.path.basename(run_dir); v['task'] = res['task']; v['harness'] = res['harness']; v['grader_passed'] = passed
    v['signals'] = signals
    json.dump(v, open(outp, 'w'), indent=2, default=str)
    meta = v.get('_meta', {})
    print(f"[{v.get('verdict','?'):16}] {res['run_id']:44} {v.get('one_line','')[:110]}  (${meta.get('cost_usd') if meta.get('cost_usd') is not None else str(meta.get('api_equivalent_usd')) + ' api-equivalent'})", flush=True)
    return v

def summarize(label_dir: str) -> None:
    audits = []
    for p in sorted(glob.glob(os.path.join(label_dir, '*', 'audit.json'))):
        try: audits.append(json.load(open(p)))
        except Exception: continue
    if not audits: return
    by_verdict = Counter(a.get('verdict') for a in audits)
    by_task = defaultdict(Counter)
    for a in audits: by_task[a['task']][a.get('verdict')] += 1
    verifier = defaultdict(list); fixes = defaultdict(list); suspicious = []; behaviours = Counter()
    for a in audits:
        for c in a.get('checks') or []:
            if c.get('reviewer_verdict') in ('check_wrong', 'ambiguous', 'false_pass'):
                verifier[(a['task'], c.get('name'))].append((a['run_id'], c.get('reviewer_verdict'), c.get('why')))
        fx = a.get('recommended_fix') or {}
        if fx.get('target') and fx.get('target') != 'none':
            fixes[fx['target']].append((a['run_id'], fx.get('change')))
        if a.get('verdict') == 'suspicious_pass': suspicious.append(a)
        for k, val in (a.get('agent_behaviour') or {}).items():
            if val: behaviours[k] += 1
    cost = sum((a.get('_meta') or {}).get('cost_usd') or 0 for a in audits)
    lines = [f'# Audit summary: {os.path.basename(label_dir)}', '',
             f'{len(audits)} runs audited; reviewer {(audits[0].get("_meta") or {}).get("model")}; reviewer cost ${cost:.2f}', '',
             '## Verdicts', '', '| verdict | runs |', '|---|---|'] + [f'| {k} | {v} |' for k, v in by_verdict.most_common()]
    lines += ['', '## By task', '', '| task | ' + ' | '.join(sorted(by_verdict)) + ' |', '|---|' + '---|' * len(by_verdict)]
    for t in sorted(by_task):
        lines.append(f'| {t} | ' + ' | '.join(str(by_task[t].get(k, 0)) for k in sorted(by_verdict)) + ' |')
    lines += ['', '## Runs', '', '| run | verdict | conf | one line |', '|---|---|---|---|']
    for a in sorted(audits, key=lambda a: a.get('run_dir') or a['run_id']):
        lines.append(f"| {a.get('run_dir') or a['run_id']} | {a.get('verdict')} | {a.get('confidence', '')} | {str(a.get('one_line', '')).replace('|', '/')} |")
    lines += ['', '## Verifier suspects (checks the reviewer disputed)', '']
    if verifier:
        for (t, name), items in sorted(verifier.items()):
            lines.append(f'- **{t} / {name}** ({len(items)} run(s))')
            for run_id, rv, why in items[:4]: lines.append(f'  - {run_id}: {rv}: {why}')
    else:
        lines.append('- none')
    lines += ['', '## Suspicious passes', ''] + ([f"- {a['run_id']}: {a.get('one_line')}" for a in suspicious] or ['- none'])
    lines += ['', '## Agent behaviours flagged', ''] + ([f'- {k}: {v}' for k, v in behaviours.most_common()] or ['- none'])
    lines += ['', '## Recommended fixes by target', '']
    for target, items in sorted(fixes.items(), key=lambda kv: -len(kv[1])):
        lines.append(f'### {target} ({len(items)})')
        for run_id, change in items: lines.append(f'- {run_id}: {change}')
        lines.append('')
    open(os.path.join(label_dir, 'audit-summary.md'), 'w').write('\n'.join(lines) + '\n')
    json.dump({'runs': len(audits), 'verdicts': dict(by_verdict), 'by_task': {t: dict(c) for t, c in by_task.items()},
               'verifier_suspects': [{'task': t, 'check': n, 'runs': [r for r, _, _ in items]} for (t, n), items in verifier.items()],
               'suspicious_passes': [a['run_id'] for a in suspicious], 'behaviours': dict(behaviours), 'reviewer_cost_usd': round(cost, 2)},
              open(os.path.join(label_dir, 'audit-summary.json'), 'w'), indent=2)
    print(f'wrote {os.path.relpath(os.path.join(label_dir, "audit-summary.md"), ROOT)}')

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('paths', nargs='+', help='results/<label> dirs or single run dirs')
    ap.add_argument('--all', action='store_true', help='audit passed runs too')
    ap.add_argument('--model', default=DEFAULT_MODEL)
    ap.add_argument('--parallel', type=int, default=3)
    ap.add_argument('--force', action='store_true')
    ap.add_argument('--limit', type=int, default=None)
    ap.add_argument('--dry-run', action='store_true', help='write audit-bundle.md only, no reviewer call')
    ap.add_argument('--backend', choices=['openrouter', 'claude-cli'], default=os.environ.get('BENCH_AUDIT_BACKEND', 'openrouter'), help='openrouter: API key, per-token cost. claude-cli: the claude CLI on a Claude subscription (sign in first)')
    ap.add_argument('--out-name', default='audit.json', help='verdict file name per run; anything other than audit.json skips the label summary (side-by-side reviews)')
    ap.add_argument('--from-bundles', action='store_true', help='review from audit-bundle.json files written by an earlier --dry-run (no workspace needed)')
    ap.add_argument('--stale', action='store_true', help='only runs with no audit yet, or whose grader verdict changed since (after a regrade)')
    a = ap.parse_args()
    runs, labels = [], set()
    for p in a.paths:
        p = p.rstrip('/')
        if os.path.exists(os.path.join(p, 'result.json')):
            runs.append(p); labels.add(os.path.dirname(p)); continue
        for rp in sorted(glob.glob(os.path.join(p, '*', 'result.json'))):
            r = json.load(open(rp))
            if a.all or not r['passed']:
                runs.append(os.path.dirname(rp)); labels.add(p)
    if a.stale:
        def is_stale(r):
            ap_ = os.path.join(r, 'audit.json')
            if not os.path.exists(ap_): return True
            try: old = json.load(open(ap_))
            except Exception: return True
            if old.get('verdict') in ('error', 'unparsed'): return True
            return bool(old.get('grader_passed')) != bool(json.load(open(os.path.join(r, 'result.json')))['passed'])
        runs = [r for r in runs if is_stale(r)]
        a.force = True
    if a.limit: runs = runs[:a.limit]
    key = None if (a.dry_run or a.backend == 'claude-cli') else api_key()
    print(f'{len(runs)} runs to audit with {a.model} via {a.backend}', flush=True)
    with cf.ThreadPoolExecutor(max_workers=a.parallel) as ex:
        list(ex.map(lambda r: audit_run(r, a.model, key, a.force, a.dry_run, a.backend, a.from_bundles, a.out_name), runs))
    if not a.dry_run and a.out_name == 'audit.json':
        for l in sorted(labels): summarize(l)

if __name__ == '__main__':
    main()
