"""bench/metamorphic.py: transforms as pure functions and the per-task classification rules.
No LibreOffice and no grading here; the tool's end-to-end run is `python3 bench/metamorphic.py`."""
import csv, io, os, re, sys, tempfile, unittest
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'bench'))
import metamorphic as mt  # noqa: E402

CSV = b'id,name,amount\n1,"Smith, Ann",10.50\n2,Bob,3\n3,Cy,7\n'


def rows(data: bytes):
    text = data.decode('utf-8-sig')
    return list(csv.reader(io.StringIO(text, newline='')))


def records(data: bytes):
    r = rows(data)
    return [dict(zip(r[0], x)) for x in r[1:] if x]


class VisibleText(HTMLParser):
    """Rendered text approximation: text outside script/style with whitespace collapsed."""
    def __init__(self):
        super().__init__(convert_charrefs=True); self.parts = []; self.skip = 0
    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'): self.skip += 1
    def handle_endtag(self, tag):
        if tag in ('script', 'style'): self.skip -= 1
    def handle_data(self, d):
        if not self.skip: self.parts.append(d)


def visible(data: bytes) -> str:
    p = VisibleText(); p.feed(data.decode()); return ' '.join(' '.join(p.parts).split())


def tags(data: bytes):
    out = []
    class P(HTMLParser):
        def handle_starttag(self, tag, attrs): out.append(tag)
    P().feed(data.decode()); return out


class CsvTransforms(unittest.TestCase):
    def test_crlf_keeps_records(self):
        out = mt.csv_crlf(CSV)
        self.assertIn(b'\r\n', out); self.assertNotIn(b'\n', out.replace(b'\r\n', b''))
        self.assertEqual(records(out), records(CSV))

    def test_quote_all(self):
        out = mt.csv_quote_all(CSV)
        self.assertTrue(out.startswith(b'"id","name","amount"'))
        self.assertEqual(records(out), records(CSV))

    def test_reverse_rows_same_multiset_header_first(self):
        out = mt.csv_reverse_rows(CSV)
        self.assertEqual(rows(out)[0], ['id', 'name', 'amount'])
        self.assertEqual([r['id'] for r in records(out)], ['3', '2', '1'])
        self.assertEqual(sorted(map(str, records(out))), sorted(map(str, records(CSV))))

    def test_reverse_columns_same_records(self):
        out = mt.csv_reverse_columns(CSV)
        self.assertEqual(rows(out)[0], ['amount', 'name', 'id'])
        self.assertEqual(records(out), records(CSV))

    def test_bom_idempotent(self):
        out = mt.csv_bom(CSV)
        self.assertTrue(out.startswith(b'\xef\xbb\xbf'))
        self.assertEqual(mt.csv_bom(out), out)
        self.assertEqual(records(out), records(CSV))

    def test_bom_survives_other_transforms(self):
        out = mt.csv_quote_all(mt.csv_bom(CSV))
        self.assertTrue(out.startswith(b'\xef\xbb\xbf"id"'))

    def test_no_final_newline_and_header_case(self):
        self.assertFalse(mt.csv_no_final_newline(CSV).endswith(b'\n'))
        self.assertEqual(rows(mt.csv_header_case(b'external_id,name\n1,a\n'))[0], ['External Id', 'Name'])

    def test_embedded_newline_preserved(self):
        data = b'id,note\n1,"line one\nline two"\n2,x\n'
        for fn in (mt.csv_crlf, mt.csv_quote_all, mt.csv_reverse_rows, mt.csv_reverse_columns):
            key = lambda recs: sorted(str(sorted(r.items())) for r in recs)
            self.assertEqual(key(records(fn(data))), key(records(data)), fn.__name__)

    def test_crlf_input_line_endings_kept(self):
        data = CSV.replace(b'\n', b'\r\n')
        self.assertIn(b'\r\n', mt.csv_reverse_rows(data))

    def test_pure(self):
        before = bytes(CSV)
        for name, t in mt.TRANSFORMS.items():
            if t.kind == 'csv':
                self.assertEqual(t.fn(CSV), t.fn(CSV), name)
        self.assertEqual(CSV, before)


MD = """# Weekly memo

Revenue for the quarter was $380,537, up 10.3% from $344,988 in Q1. May revenue is restated after finance found a duplicate invoice.

- First bullet that is long enough to need wrapping when the width is narrow, with Umber Ceramics Studio named.
- Second bullet
  continued on the next line.
1. Numbered item with a number 2026 inside that is long enough to wrap somewhere near the middle.

| a | b |
|---|---|
| 1 | 2 |

```
code   stays   put
```
---
Thanks,
Marcus
"""


def md_words(text: str):
    """Word stream with Markdown list markers normalised, for comparing wrapped and unwrapped text."""
    return re.sub(r'(^|\n)\s*[*-]\s', r'\1- ', text).split()


class MarkdownTransforms(unittest.TestCase):
    def test_rewrap_narrow_width_and_words(self):
        out = mt.md_rewrap_narrow(MD.encode()).decode()
        prose = [ln for ln in out.split('\n') if not ln.startswith(('|', '```', 'code'))]
        self.assertTrue(all(len(ln) <= 40 or len(ln.split()) == 1 or ln.startswith('#') for ln in prose), out)
        self.assertEqual(md_words(out), md_words(MD))

    def test_rewrap_keeps_structure(self):
        out = mt.md_rewrap_narrow(MD.encode()).decode()
        self.assertIn('| a | b |\n|---|---|\n| 1 | 2 |', out)
        self.assertIn('```\ncode   stays   put\n```', out)
        self.assertIn('# Weekly memo\n', out)
        self.assertIn('\n---\n', out)
        # every list item still starts a line; no continuation line starts with list syntax
        self.assertEqual(len(re.findall(r'^- ', out, re.M)), 2)
        self.assertEqual(len(re.findall(r'^\d+\. ', out, re.M)), 1)
        for ln in out.split('\n'):
            if ln.startswith('  '):
                self.assertFalse(re.match(r'\s+([-*+]|\d+[.)])\s', ln), ln)

    def test_unwrap_joins_paragraph_lines(self):
        out = mt.md_unwrap(MD.encode()).decode()
        self.assertIn('- Second bullet continued on the next line.', out)
        self.assertEqual(md_words(out), md_words(MD))

    def test_rewrap_72(self):
        out = mt.md_rewrap_72(MD.encode()).decode()
        self.assertEqual(md_words(out), md_words(MD))

    def test_star_bullets(self):
        out = mt.md_star_bullets(MD.encode()).decode()
        self.assertIn('* First bullet', out)
        self.assertIn('\n---\n', out)  # thematic break untouched
        self.assertNotIn('\n- ', out)

    def test_trailing_ws_single_space(self):
        out = mt.md_trailing_ws(MD.encode()).decode()
        self.assertIn('# Weekly memo \n', out)
        self.assertNotIn('  \n', out.replace('code   stays   put ', ''))

    def test_crlf_roundtrip(self):
        out = mt.md_crlf(MD.encode())
        self.assertEqual(out.replace(b'\r\n', b'\n'), MD.encode())

    def test_smart_quotes(self):
        out = mt.md_smart_quotes(b'It\'s "done" and \'quoted\'').decode()
        self.assertEqual(out, 'It’s “done” and ‘quoted’')

    def test_crlf_input_rewrap_keeps_crlf(self):
        out = mt.md_rewrap_narrow(MD.replace('\n', '\r\n').encode())
        self.assertNotIn(b'\n', out.replace(b'\r\n', b''))


HTML = b"""<!doctype html>
<html><head><title>Menu</title><style>p > a { color: red; }</style>
<script>if (a<b && c>d) { x = "don't"; }</script></head>
<body>
  <h1>Fall menu</h1>
  <ul><li>Pork and chive gyoza <span>$9.50</span></li><li>Chef's special: a very long description of the dish that goes on and on caf\xc3\xa9</li></ul>
  <pre>  keep
     this</pre>
  <p>He said "hello".</p>
</body></html>
"""


class HtmlTransforms(unittest.TestCase):
    def check_render_same(self, fn):
        out = fn(HTML)
        self.assertEqual(visible(out), visible(HTML), fn.__name__)
        self.assertEqual([t.lower() for t in tags(out)], tags(HTML), fn.__name__)
        return out

    def test_all_render_invariant(self):
        for name, t in mt.TRANSFORMS.items():
            if t.kind == 'html':
                self.check_render_same(t.fn)

    def test_raw_elements_untouched(self):
        for name, t in mt.TRANSFORMS.items():
            if t.kind == 'html':
                out = t.fn(HTML)
                self.assertIn(b'if (a<b && c>d) { x = "don\'t"; }', out, name)
                self.assertIn(b'  keep\n     this', out, name)
                self.assertIn(b'p > a { color: red; }', out, name)

    def test_specific_effects(self):
        self.assertIn(b'Chef&#39;s', mt.html_escape_quotes(HTML))
        self.assertIn(b'&quot;hello&quot;', mt.html_escape_quotes(HTML))
        self.assertIn(b'caf&#233;', mt.html_entities_nonascii(HTML))
        self.assertIn(b'<LI>', mt.html_upper_tags(HTML))
        self.assertIn(b'</SCRIPT>', mt.html_upper_tags(HTML))
        self.assertIn(b'\n<li>', mt.html_break_blocks(HTML))
        wrapped = mt.html_rewrap_text(HTML)
        self.assertIn(b'\n', re.search(rb'Chef.*?caf', wrapped, re.S).group(0))
        self.assertIn(b'\n    <h1>', mt.html_reindent(HTML))


ICS = ('BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\nSUMMARY:Puppy class — week one at the north field with Sam\n'
       'DESCRIPTION:' + 'x' * 120 + '\nEND:VEVENT\nEND:VCALENDAR\n').encode()


def ics_unfold(text: str) -> str:
    return re.sub(r'\r?\n[ \t]', '', text)


class IcsTransforms(unittest.TestCase):
    def test_fold_unfolds_to_original(self):
        out = mt.ics_fold(ICS)
        lines = out.decode().split('\n')
        self.assertTrue(all(len(ln.encode()) <= 40 for ln in lines))
        self.assertEqual(ics_unfold(out.decode()), ICS.decode())

    def test_fold_respects_utf8_boundaries(self):
        mt.ics_fold(ICS).decode('utf-8')  # raises if a multi-byte character was split

    def test_crlf(self):
        self.assertEqual(mt.ics_crlf(ICS).replace(b'\r\n', b'\n'), ICS)


class XlsxTransforms(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.d = tempfile.mkdtemp()
        wb = Workbook(); sh = wb.active; sh.title = 'Summary'
        sh.append(['Job', 'Revenue']); sh.append(['Harborview', 100]); sh.append(['Total', '=SUM(B2:B2)'])
        wb.create_sheet('Detail').append(['x', 1])
        self.src = os.path.join(self.d, 'in.xlsx'); wb.save(self.src)
        one = Workbook(); one.active.append(['a', 1]); self.one = os.path.join(self.d, 'one.xlsx'); one.save(self.one)

    def cells(self, path):
        from openpyxl import load_workbook
        wb = load_workbook(path)
        return {sh.title: [[c.value for c in r] for r in sh.iter_rows()] for sh in wb.worksheets}, wb.sheetnames

    def test_resave_same_cells(self):
        dst = os.path.join(self.d, 'o.xlsx')
        self.assertTrue(mt.xlsx_resave(self.src, dst))
        self.assertEqual(self.cells(dst), self.cells(self.src))

    def test_extra_sheet_appended_text_only(self):
        dst = os.path.join(self.d, 'o.xlsx'); mt.xlsx_extra_sheet(self.src, dst)
        cells, names = self.cells(dst)
        self.assertEqual(names, ['Summary', 'Detail', 'Notes'])
        self.assertTrue(all(isinstance(v, str) and not v.startswith('=') for r in cells['Notes'] for v in r if v is not None))
        self.assertEqual(cells['Summary'], self.cells(self.src)[0]['Summary'])

    def test_sheet_order(self):
        dst = os.path.join(self.d, 'o.xlsx')
        self.assertTrue(mt.xlsx_sheet_order(self.src, dst))
        cells, names = self.cells(dst)
        self.assertEqual(names, ['Detail', 'Summary'])
        self.assertEqual(cells['Summary'][2][1], '=SUM(B2:B2)')
        self.assertFalse(mt.xlsx_sheet_order(self.one, os.path.join(self.d, 'n.xlsx')))
        self.assertFalse(os.path.exists(os.path.join(self.d, 'n.xlsx')))

    def test_cover_first(self):
        dst = os.path.join(self.d, 'o.xlsx'); mt.xlsx_cover_sheet_first(self.src, dst)
        from openpyxl import load_workbook
        wb = load_workbook(dst)
        self.assertEqual(wb.sheetnames, ['Cover', 'Summary', 'Detail'])
        self.assertEqual(wb.active.title, 'Cover')

    def test_source_untouched(self):
        before = Path(self.src).read_bytes()
        for t in mt.TRANSFORMS.values():
            if t.kind == 'xlsx':
                t.fn(self.src, os.path.join(self.d, t.name + '.xlsx'))
        self.assertEqual(Path(self.src).read_bytes(), before)


def task(ask='Clean this up.', checks=()):
    return {'id': 't', 'ask': ask, 'checks': list(checks)}


class Classification(unittest.TestCase):
    def test_every_transform_has_a_class(self):
        for t in mt.TRANSFORMS.values():
            self.assertIn(t.default_class, (mt.INVARIANT, mt.SENSITIVE))
            self.assertIn(t.kind, ('csv', 'xlsx', 'md', 'html', 'ics', 'workspace'))

    def test_column_reorder_exact_template_is_sensitive(self):
        t = task(checks=[{'type': 'csv_columns', 'path': 'out.csv', 'columns': ['a', 'b'], 'exact': True}])
        self.assertEqual(mt.classify('csv_reverse_columns', t, 'out.csv')[0], mt.SENSITIVE)
        t = task(checks=[{'type': 'csv_columns', 'path': 'out.csv', 'columns': ['a', 'b']}])
        self.assertEqual(mt.classify('csv_reverse_columns', t, 'out.csv')[0], mt.INVARIANT)

    def test_exact_template_on_another_file_does_not_count(self):
        t = task(checks=[{'type': 'csv_columns', 'path': 'other.csv', 'columns': ['a'], 'exact': True}])
        self.assertEqual(mt.classify('csv_reverse_columns', t, 'out.csv')[0], mt.INVARIANT)

    def test_import_ask_makes_bom_and_columns_sensitive(self):
        t = task(ask='Build the Zendesk import file from the template.')
        for n in ('csv_bom', 'csv_reverse_columns', 'csv_header_case'):
            self.assertEqual(mt.classify(n, t, 'users.csv')[0], mt.SENSITIVE, n)
        self.assertEqual(mt.classify('csv_bom', task(), 'report.csv')[0], mt.INVARIANT)

    def test_row_order(self):
        self.assertEqual(mt.classify('csv_reverse_rows', task(), 'x.csv')[0], mt.INVARIANT)
        self.assertEqual(mt.classify('csv_reverse_rows', task('List them oldest first.'), 'x.csv')[0], mt.SENSITIVE)
        self.assertEqual(mt.classify('csv_reverse_rows', task('Sorted by vendor please.'), 'x.csv')[0], mt.SENSITIVE)
        self.assertEqual(mt.classify('csv_reverse_rows', task('Open purchase orders in order to pay.'), 'x.csv')[0],
                         mt.INVARIANT)
        named = task(checks=[{'type': 'custom', 'name': 'rows sorted by student name', 'path': 'x.csv'}])
        self.assertEqual(mt.classify('csv_reverse_rows', named, 'x.csv')[0], mt.SENSITIVE)
        cols = task(checks=[{'type': 'csv_columns', 'name': 'template columns in order', 'path': 'x.csv', 'columns': ['a']}])
        self.assertEqual(mt.classify('csv_reverse_rows', cols, 'x.csv')[0], mt.INVARIANT)
        src = {'check.py': 'p = find(ws, "grid.csv")'}
        self.assertEqual(mt.classify('csv_reverse_rows', task(), 'grid.csv', src)[0], mt.SENSITIVE)
        self.assertEqual(mt.classify('csv_reverse_rows', task(), 'other.csv', src)[0], mt.INVARIANT)

    def test_sheet_order(self):
        table = task(checks=[{'type': 'csv_values_match', 'path': 'prices.xlsx', 'ref': 'p.csv', 'key': 'k', 'columns': ['v']}])
        self.assertEqual(mt.classify('xlsx_sheet_order', table, 'prices.xlsx')[0], mt.SENSITIVE)
        named = task(checks=[{'type': 'csv_values_match', 'path': 'prices.xlsx', 'sheet': 'Prices', 'ref': 'p.csv',
                              'key': 'k', 'columns': ['v']}])
        self.assertEqual(mt.classify('xlsx_sheet_order', named, 'prices.xlsx')[0], mt.INVARIANT)
        scan = task(checks=[{'type': 'xlsx_value_present', 'path': 'm.xlsx', 'expected': 1}])
        self.assertEqual(mt.classify('xlsx_cover_sheet_first', scan, 'm.xlsx')[0], mt.INVARIANT)
        src = {'check.py': 'wb = load_workbook(find(ws, "m.xlsx")).active'}
        self.assertEqual(mt.classify('xlsx_cover_sheet_first', scan, 'm.xlsx', src)[0], mt.SENSITIVE)

    def test_fixed_classes(self):
        for n in ('md_rewrap_narrow', 'md_unwrap', 'html_escape_quotes', 'xlsx_resave', 'ics_fold', 'csv_crlf'):
            self.assertEqual(mt.classify(n, task(), 'x')[0], mt.INVARIANT, n)
        for n in ('md_smart_quotes', 'move_to_subdir', 'rename_case'):
            self.assertEqual(mt.classify(n, task(), 'x')[0], mt.SENSITIVE, n)

    def test_checks_on_globs_and_subdirs(self):
        t = task(checks=[{'type': 'file_exists', 'path': 'notices/*.md'}, {'type': 'file_exists', 'path': 'memo.md'},
                         {'type': 'not_fooled', 'path': 'a.csv', 'forbidden_text': [{'path': 'memo.md', 'phrases': ['x']}]}])
        self.assertEqual(len(mt.checks_on(t, 'notices/Aspen_House.md')), 1)
        self.assertEqual(len(mt.checks_on(t, 'memo.md')), 2)


class Apply(unittest.TestCase):
    def test_apply_and_workspace_moves(self):
        d = tempfile.mkdtemp(); src = os.path.join(d, 'src'); os.makedirs(os.path.join(src, 'sub'))
        Path(src, 'a.csv').write_bytes(CSV)
        Path(src, 'sub', 'memo.md').write_text('- x\n')
        files = mt.deliverables(src)
        self.assertEqual(files, ['a.csv', 'sub/memo.md'])
        self.assertEqual(mt.apply_transform('csv_bom', src, os.path.join(d, 'o1'), files), ['a.csv'])
        self.assertEqual(mt.apply_transform('md_crlf', src, os.path.join(d, 'o2'), files), ['sub/memo.md'])
        self.assertEqual(mt.apply_transform('md_unwrap', src, os.path.join(d, 'o3'), files), [])  # no-op
        mt.apply_transform('move_to_subdir', src, os.path.join(d, 'o4'), files)
        self.assertTrue(os.path.isfile(os.path.join(d, 'o4', 'output', 'sub', 'memo.md')))
        mt.apply_transform('rename_case', src, os.path.join(d, 'o5'), files)
        self.assertTrue(os.path.isfile(os.path.join(d, 'o5', 'A.csv')))
        self.assertEqual(Path(src, 'a.csv').read_bytes(), CSV)

    def test_candidates_and_matrix(self):
        recs = [{'task': 't1', 'baseline': 'pass', 'results': {
            'csv_crlf': {'status': 'fail', 'class': mt.INVARIANT, 'files': ['a.csv'], 'reason': '',
                         'failed': [{'name': 'c', 'type': 'csv_columns', 'detail': 'd'}]},
            'csv_bom': {'status': 'fail', 'class': mt.SENSITIVE, 'files': ['a.csv'], 'reason': 'r',
                        'failed': [{'name': 'c', 'type': 'csv_columns', 'detail': 'd'}]},
            'csv_quote_all': {'status': 'pass', 'class': mt.INVARIANT, 'files': ['a.csv'], 'reason': ''}}}]
        self.assertEqual([c['transform'] for c in mt.candidates(recs)], ['csv_crlf'])
        self.assertEqual([c['transform'] for c in mt.sensitive_fails(recs)], ['csv_bom'])
        m = mt.matrix(recs, ['csv_crlf', 'csv_quote_all', 'csv_bom', 'md_crlf'])
        self.assertRegex(m, r't1\s+pass\s+F\s+\.\s+f')
        self.assertNotIn('md_crlf', m)


class Isolation(unittest.TestCase):
    def test_runner_and_grader_do_not_import_it(self):
        for f in ('grade.py', 'run.py', 'validate_tasks.py', 'regrade.py'):
            p = ROOT / 'bench' / f
            if p.exists():
                self.assertNotIn('metamorphic', p.read_text(), f)

    def test_import_does_not_load_grader(self):
        import subprocess
        code = 'import sys; sys.path.insert(0, "bench"); import metamorphic; print("grade" in sys.modules, "pandas" in sys.modules)'
        out = subprocess.run([sys.executable, '-c', code], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        self.assertEqual(out, 'False False')


if __name__ == '__main__':
    unittest.main()
