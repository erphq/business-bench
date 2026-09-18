#!/usr/bin/env python3
"""stripe-customer-import: a software company's legacy billing export turned into a Stripe customer import file.

    python gen.py [--seed N]

Business: Kestrel Analytics bills customers in five currencies from an old billing tool and is moving billing to
Stripe. The finance lead wrote down how the old export maps onto Stripe's customer import columns.

Traps (each caught by a check, see task.yaml):
  * account numbers are six digits with leading zeros, stored as numbers with a 000000 display format, so a plain
    read gives 417 instead of 000417                                            (checks: one customer per legacy account; plan metadata)
  * "$" means US dollars for US customers but Canadian or Australian dollars for customers there, one Canadian
    customer is explicitly billed in US$, and a blank currency takes the country's default  (check: currency)
  * suite, unit, flat, level and floor sit inside the street cell after a comma or on a second line (and in front
    of the street for UK and Australian addresses); they belong in address_line2 (check: address lines)
  * states and provinces are sometimes spelled out; Stripe wants the abbreviation for US, CA and AU and nothing for
    other countries, and US ZIP codes starting with 0 lost their leading zero in the export (check: city, state and postal code)
  * plan names are typed loosely and "Pro" is the old name of Business; metadata[plan] takes the plan code from
    the price list                                                              (check: plan metadata)
  * cancelled accounts and the internal QA accounts are in the export and stay out (checks: one customer per legacy account; row count)
"""
from __future__ import annotations
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "lib"))
from bizgen import *  # noqa: E402,F403

TEMPLATE = ["email", "name", "description", "phone", "currency", "address_line1", "address_line2", "address_city", "address_state",
            "address_postal_code", "address_country", "metadata[legacy_account]", "metadata[plan]"]
US_NAMES = {"OR": "Oregon", "TX": "Texas", "CO": "Colorado", "WI": "Wisconsin", "NC": "North Carolina", "ID": "Idaho", "VT": "Vermont",
            "NM": "New Mexico", "MI": "Michigan", "GA": "Georgia", "WA": "Washington", "AZ": "Arizona", "VA": "Virginia", "CA": "California",
            "OH": "Ohio", "NE": "Nebraska", "RI": "Rhode Island"}
ZERO_ZIPS = [("Burlington", "VT", "05401"), ("Providence", "RI", "02903"), ("Boston", "MA", "02110"), ("Portland", "ME", "04101")]
US_NAMES.update({"MA": "Massachusetts", "ME": "Maine"})
CA_CITIES = [("Toronto", "ON", "Ontario", "M5V 2T6"), ("Vancouver", "BC", "British Columbia", "V6B 1A1"),
             ("Calgary", "AB", "Alberta", "T2P 1J9"), ("Montreal", "QC", "Quebec", "H3B 2Y5")]
AU_CITIES = [("Sydney", "NSW", "New South Wales", "2000"), ("Melbourne", "VIC", "Victoria", "3000"), ("Brisbane", "QLD", "Queensland", "4000")]
EU = {"GB": [("London", "Greater London", "EC2A 3AR", ["Leonard Street", "Old Street", "Tabernacle Street"]),
             ("Manchester", "", "M1 1AE", ["Portland Street", "Oldham Street"]), ("Leeds", "West Yorkshire", "LS1 4AP", ["Park Row", "Wellington Street"])],
      "DE": [("Berlin", "Berlin", "10115", ["Invalidenstraße", "Chausseestraße"]), ("München", "Bayern", "80331", ["Sendlinger Straße", "Kaufingerstraße"])],
      "IE": [("Dublin", "Co. Dublin", "D02 X285", ["Grand Canal Street", "Baggot Street"])],
      "NL": [("Amsterdam", "Noord-Holland", "1012 AB", ["Damrak", "Herengracht"])]}
DEFAULT_CCY = {"US": "usd", "CA": "cad", "AU": "aud", "GB": "gbp", "DE": "eur", "IE": "eur", "NL": "eur"}
CCY_TEXT = {"usd": ["USD", "US$", "$"], "cad": ["$", "CAD", "C$"], "aud": ["$", "A$", "AUD"], "gbp": ["£", "GBP"], "eur": ["€", "EUR", "Euro"]}
PLANS = {"starter_monthly": ["Starter (monthly)", "Starter/mo", "starter monthly"], "starter_annual": ["Starter (annual)", "Starter - Annual"],
         "team_monthly": ["Team (monthly)", "team monthly", "Team/mo"], "team_annual": ["Team (annual)", "Team - Annual", "team yearly"],
         "business_monthly": ["Business (monthly)", "Pro Monthly", "PRO - monthly"], "business_annual": ["Business (annual)", "Pro (annual)", "Pro yearly"]}
CUSTOMER_SUFFIX = ["", " Inc", " LLC", " Ltd", " Pty Ltd", " GmbH", " BV"]
EXTRA_NAMES = ["Northwind Grocers", "Bluebird Clinics", "Harrow & Pike", "Coastline Credit Union", "Tidewater Schools", "Lumen Architects",
               "Brassica Farms", "Fjord Travel", "Greywell Hotels", "Orbit Dental Labs", "Paloma Foods", "Quayside Marinas", "Rook Legal",
               "Sable Fitness", "Tern Insurance", "Upland Realty", "Verdant Energy", "Wharf Street Books", "Yarrow Health", "Zinnia Events"]


def build(seed: int) -> dict:
    r = rng(seed)
    names = [c[0] for c in COMPANIES if c[0] != "Kestrel Analytics"] + EXTRA_NAMES
    r.shuffle(names)
    plan_countries = ["US"] * 21 + ["CA"] * 6 + ["GB"] * 4 + ["AU"] * 3 + ["DE"] * 2 + ["IE"] + ["NL"]
    r.shuffle(plan_countries)
    accts = r.sample(range(100, 99999), 60)
    customers = []
    for k, cc in enumerate(plan_countries):
        f, l = person(r)
        nm = names[k]
        suffix = {"US": [" Inc", " LLC", ""], "CA": [" Inc", " Ltd", ""], "GB": [" Ltd", ""], "AU": [" Pty Ltd"], "DE": [" GmbH"],
                  "IE": [" Ltd"], "NL": [" BV"]}[cc]
        nm = nm + r.choice(suffix)
        dom = re_domain(names[k])
        c = {"name": nm, "contact": f"{f} {l}", "email": f"{f.lower()}.{l.lower()}@{dom}", "country": cc, "tags": set(),
             "acct": f"{accts[k]:06d}", "plan": r.choice(list(PLANS)), "status": "Active", "line2": ""}
        c["currency"] = DEFAULT_CCY[cc]
        if cc == "US":
            city, st, z = r.choice(CITIES)
            c.update(city=city, state=st, postal=z, line1=f"{r.randint(12, 9899)} {r.choice(STREETS)}", region_text=st)
            c["phone"] = phone_variant(phone_digits(r), 0)
        elif cc == "CA":
            city, st, full, z = r.choice(CA_CITIES)
            c.update(city=city, state=st, postal=z, line1=f"{r.randint(12, 2899)} {r.choice(STREETS)}", region_text=st)
            c["phone"] = phone_variant(f"{r.choice(['416', '604', '403', '514'])}555{r.randint(1000, 9999)}", 0)
            c["full_region"] = full
        elif cc == "AU":
            city, st, full, z = r.choice(AU_CITIES)
            c.update(city=city, state=st, postal=z, line1=f"{r.randint(1, 480)} {r.choice(['George St', 'Pitt St', 'Collins St', 'Queen St'])}",
                     region_text=st, full_region=full)
            c["phone"] = f"+61 2 5550 {r.randint(1000, 9999)}"
        else:
            city, region, z, streets = r.choice(EU[cc])
            num = r.randint(1, 180)
            line1 = f"{num} {r.choice(streets)}" if cc in ("GB", "IE") else f"{r.choice(streets)} {num}"
            c.update(city=city, state="", postal=z, line1=line1, region_text=region)
            c["phone"] = {"GB": f"+44 20 7946 0{r.randint(100, 999)}", "DE": f"+49 30 5550{r.randint(1000, 9999)}",
                          "IE": f"+353 1 555 {r.randint(1000, 9999)}", "NL": f"+31 20 555 {r.randint(1000, 9999)}"}[cc]
        customers.append(c)
    by = lambda cc: [c for c in customers if c["country"] == cc]
    # zero-leading ZIPs on three US customers
    for c, (city, st, z) in zip(by("US")[:3], ZERO_ZIPS):
        c.update(city=city, state=st, postal=z, region_text=st); c["tags"].add("geo")
    for c in by("US"):
        if c["postal"].startswith("0"):
            c["tags"].add("geo")
    # spelled-out regions
    for c in by("US")[3:7] + by("CA")[:3] + by("AU")[:2]:
        c["region_text"] = c.get("full_region") or US_NAMES[c["state"]]; c["tags"].add("geo")
    for c in by("GB") + by("DE") + by("IE") + by("NL"):
        if c["region_text"]:
            c["tags"].add("geo")
    # line2 variants
    us_units = [("Suite", 100, 900), ("Ste", 100, 450), ("Unit", 1, 40), ("Floor", 2, 12)]
    for c in r.sample(by("US")[3:] + by("CA"), 7):
        word, lo, hi = r.choice(us_units)
        c["line2"] = f"{word} {r.randint(lo, hi)}"; c["line2_style"] = r.choice(["comma", "newline"]); c["tags"].add("addr")
    for c in by("GB")[:2]:
        c["line2"] = r.choice(["Flat", "Floor"]) + f" {r.randint(1, 9)}"; c["line2_style"] = "front"; c["tags"].add("addr")
    for c in by("AU")[:2]:
        c["line2"] = f"Level {r.randint(2, 30)}"; c["line2_style"] = "front"; c["tags"].add("addr")
    # currency text
    for c in customers:
        cc, ccy = c["country"], c["currency"]
        opts = CCY_TEXT[ccy]
        if cc == "US":
            c["ccy_text"] = r.choice(opts + [""])
        elif cc in ("CA", "AU"):
            c["ccy_text"] = r.choice(["$", "$", opts[1], opts[2]])
        else:
            c["ccy_text"] = r.choice(opts + [""])
        if c["ccy_text"] not in (ccy.upper(),):
            c["tags"].add("ccy")
    ca = by("CA")[-1]
    ca["currency"] = "usd"; ca["ccy_text"] = "US$"; ca["tags"].add("ccy")
    for c in by("CA")[:2] + by("AU")[:1]:
        c["ccy_text"] = "$"; c["tags"].add("ccy")
    # plans text
    for c in customers:
        c["plan_text"] = r.choice(PLANS[c["plan"]])
        if c["plan_text"] not in (PLANS[c["plan"]][0],) or c["plan"].startswith("business"):
            c["tags"].add("plan")
    for c in r.sample(customers, 4):
        c["plan"] = r.choice(["business_monthly", "business_annual"])
        c["plan_text"] = r.choice(PLANS[c["plan"]][1:]); c["tags"].add("plan")
    # excluded rows
    extra = []
    for k in range(3):
        f, l = person(r)
        nm = names[len(plan_countries) + k]
        extra.append({"name": nm + " Inc", "contact": f"{f} {l}", "email": f"{f.lower()}@{re_domain(nm)}", "country": "US", "tags": {"excluded"},
                      "acct": f"{accts[len(plan_countries) + k]:06d}", "plan": r.choice(list(PLANS)), "status": r.choice(["Cancelled", "Canceled"]),
                      "line2": "", "currency": "usd", "ccy_text": "USD", "city": "Denver", "state": "CO", "postal": "80202",
                      "line1": f"{r.randint(12, 9899)} {r.choice(STREETS)}", "region_text": "CO", "phone": phone_variant(phone_digits(r), 0)})
    for k, label in enumerate(["QA test account", "TEST - do not bill"]):
        extra.append({"name": label, "contact": "QA Team", "email": f"qa+{k + 1}@kestrelanalytics.io", "country": "US", "tags": {"excluded"},
                      "acct": f"{accts[len(plan_countries) + 3 + k]:06d}", "plan": "team_monthly", "status": "Active", "line2": "",
                      "currency": "usd", "ccy_text": "USD", "city": "Boulder", "state": "CO", "postal": "80302", "line1": "1 Test Way",
                      "region_text": "CO", "phone": ""})
    for c in customers + extra:
        if c["acct"].startswith("0"):
            c["tags"].add("zeros")
        c["plan_text"] = c.get("plan_text") or PLANS[c["plan"]][0]
    rows = customers + extra
    r.shuffle(rows)
    return {"customers": customers, "rows": rows}


def re_domain(name: str) -> str:
    import re
    return re.sub(r"[^a-z0-9]", "", name.lower()) + r_tld(name)


def r_tld(name: str) -> str:
    return ".com" if len(name) % 3 else ".co"


def street_cell(c) -> str:
    st = c.get("line2_style")
    if not c["line2"]:
        return c["line1"]
    if st == "comma":
        return f"{c['line1']}, {c['line2']}"
    if st == "newline":
        return f"{c['line1']}\n{c['line2']}"
    return f"{c['line2']}, {c['line1']}"


def emit(seed: int) -> None:
    d = build(seed)
    customers, rows = d["customers"], d["rows"]
    ws, ref, sol = task_dirs(HERE)
    r = rng(seed + 3)
    xrows = []
    for c in rows:
        postal = int(c["postal"]) if c["country"] == "US" else c["postal"]
        xrows.append([int(c["acct"]), c["name"], c["contact"], c["email"], c["phone"], street_cell(c), c["city"], c["region_text"],
                      postal, c["country"], c["ccy_text"], c["plan_text"], c["status"], f"**** {r.randint(1000, 9999)}",
                      r.choice(["", "", "", "PO required on invoices", "annual prepay", "migrated from invoice billing 2024"])])
    path = os.path.join(ws, "billing_customers_export.xlsx")
    write_xlsx(path, {"Customers": {
        "preamble": [["LedgerBell billing - customer export", "", "", "", "", "", "", "", "", "", "", "", "", "", "generated 2026-09-09 14:02 UTC"]],
        "header": ["Account No", "Customer", "Billing Contact", "Billing Email", "Phone", "Street", "City", "State/Region", "Postal Code",
                   "Country", "Currency", "Plan", "Status", "Card", "Internal Notes"],
        "rows": xrows, "widths": {"B": 30, "C": 20, "D": 36, "F": 34, "L": 20, "O": 32}}}, creator="LedgerBell")
    # show account numbers with their leading zeros, the way the old system displays them
    from openpyxl import load_workbook
    wb = load_workbook(path)
    sh = wb["Customers"]
    for row in sh.iter_rows(min_row=3, max_col=1):
        for cell in row:
            if isinstance(cell.value, int):
                cell.number_format = "000000"
    for row in sh.iter_rows(min_row=3, min_col=9, max_col=9):
        for cell in row:
            if isinstance(cell.value, int):
                cell.number_format = "0"
    wb.save(path)
    freeze_zip(path)
    write_csv(os.path.join(ws, "stripe_customer_import_template.csv"), TEMPLATE,
              [["billing@example.com", "Example Co", "Billing contact: Sam Example", "+1 303 555 0100", "usd", "100 Main St", "Suite 200",
                "Denver", "CO", "80202", "US", "001234", "team_monthly"]])
    write_text(os.path.join(ws, "stripe_migration_notes.md"), (
        "# Moving billing customers to Stripe\n\n"
        "Notes from Fatima (finance). The import file follows `stripe_customer_import_template.csv` exactly: same columns, same\n"
        "order, one row per customer we still bill. Leave out cancelled accounts and our own QA/test accounts.\n\n"
        "## Columns\n\n"
        "- **email, name, phone**: billing email, the customer (company) name, billing phone.\n"
        "- **description**: `Billing contact: <contact name>`.\n"
        "- **currency**: the three-letter ISO code, lowercase, of the currency we bill them in. LedgerBell shows a bare `$` for\n"
        "  US, Canadian and Australian dollars alike, so a `$` means the customer's own country's dollar. `US$` always means US\n"
        "  dollars. If the currency is blank, they are billed in their country's currency (US usd, CA cad, AU aud, GB gbp,\n"
        "  Germany / Ireland / Netherlands eur).\n"
        "- **address_line1 / address_line2**: LedgerBell has one Street box. Suite, Ste, Unit, Flat, Floor and Level go in\n"
        "  address_line2 exactly as written; the building number and street go in address_line1. (UK and Australian\n"
        "  addresses put the flat or level in front of the street.)\n"
        "- **address_city, address_postal_code**: as the customer's address has them. US ZIP codes are five digits.\n"
        "- **address_state**: the postal abbreviation for the US, Canada and Australia (CO, ON, NSW). Leave it blank for\n"
        "  every other country; Stripe does not need a county or Land.\n"
        "- **address_country**: two-letter ISO code.\n"
        "- **metadata[legacy_account]**: the LedgerBell account number, always six digits with its leading zeros\n"
        "  (`001234`, not `1234`). We use it to match payments during the cutover.\n"
        "- **metadata[plan]**: the plan code from our price list below.\n\n"
        "## Price list codes\n\n"
        "| Plan | Monthly | Annual |\n|---|---|---|\n"
        "| Starter | starter_monthly | starter_annual |\n"
        "| Team | team_monthly | team_annual |\n"
        "| Business | business_monthly | business_annual |\n\n"
        "Business was called Pro until 2025 and some old accounts still say Pro.\n\n"
        "Do not import card details; Stripe collects those from customers directly.\n"))
    out = sorted(customers, key=lambda c: c["acct"])
    header = TEMPLATE
    rrows = [[c["email"], c["name"], f"Billing contact: {c['contact']}", c["phone"], c["currency"], c["line1"], c["line2"], c["city"],
              c["state"], c["postal"], c["country"], c["acct"], c["plan"]] for c in out]
    write_csv(os.path.join(ref, "stripe_customers.csv"), header, rrows)
    write_csv(os.path.join(sol, "stripe_customers.csv"), header, rrows)

    def keys(tag):
        return sorted(c["acct"] for c in out if tag in c["tags"])
    write_json(os.path.join(ref, "notes.json"), {t: keys(t) for t in ("zeros", "ccy", "addr", "geo", "plan")} |
               {"excluded": sorted(c["acct"] for c in rows if "excluded" in c["tags"])})
    write_task_yaml(HERE, {
        "id": "stripe-customer-import", "track": "desk", "category": "reformatting",
        "title": "Build the Stripe customer import from the old billing export",
        "ask": ("We're moving billing to Stripe. Please turn the customer export from our old billing system into Stripe's customer "
                "import file using their template; Fatima's migration notes explain the columns. Save it as stripe_customers.csv.\n"),
        "followup": None, "timeout_s": 1200,
        "traps": [
            "account numbers are six digits with leading zeros stored as numbers under a 000000 display format; a plain read gives 417 instead of 000417, which breaks the key and the metadata (checks: one customer per legacy account; plan metadata)",
            "a bare $ means the customer's own country's dollar (cad for Canada, aud for Australia), one Canadian customer is billed in US$, and blank currency takes the country default; symbols and words must become lowercase ISO codes (check: currency)",
            "suites, units and floors sit in the Street cell after a comma or on a second line, and UK flats and Australian levels sit in front of the street; they belong in address_line2 with the number and street in address_line1 (check: address lines)",
            "regions are sometimes spelled out (Ontario, New South Wales, Massachusetts), non-US/CA/AU regions such as Greater London or Bayern must be blank, and New England ZIP codes (at least three rows) lost their leading zero in the export (check: city, state and postal code)",
            "plan names are typed loosely (Team/mo, team yearly, PRO - monthly) and Pro is the old name of Business; metadata[plan] takes the price-list code (check: plan metadata)",
            "three cancelled accounts and two internal QA accounts (both marked Active) are in the export and stay out; the export also has a title row above the header (checks: one customer per legacy account; row count)",
        ],
        "checks": [
            {"type": "csv_columns", "name": "Stripe template columns, exact order", "path": "stripe_customers.csv", "columns": header, "exact": True},
            {"type": "csv_set_equal", "name": "one customer per legacy account", "path": "stripe_customers.csv",
             "column": "metadata[legacy_account]", "ref": "stripe_customers.csv", "normalize": ["strip"]},
            {"type": "csv_row_count", "name": "row count", "path": "stripe_customers.csv", "equals_ref": "stripe_customers.csv"},
            {"type": "csv_values_match", "name": "currency", "path": "stripe_customers.csv", "ref": "stripe_customers.csv",
             "key": "email", "columns": ["currency"], "min_accuracy": 1.0, "must_match_keys": emails(out, "ccy")},
            {"type": "csv_values_match", "name": "address lines", "path": "stripe_customers.csv", "ref": "stripe_customers.csv",
             "key": "email", "columns": ["address_line1", "address_line2"], "normalize": ["alnum"], "min_accuracy": 1.0,
             "must_match_keys": emails(out, "addr")},
            {"type": "csv_values_match", "name": "city, state and postal code", "path": "stripe_customers.csv", "ref": "stripe_customers.csv",
             "key": "email", "columns": ["address_city", "address_state", "address_postal_code", "address_country"], "normalize": ["alnum"],
             "min_accuracy": 1.0, "must_match_keys": emails(out, "geo")},
            {"type": "csv_values_match", "name": "plan metadata", "path": "stripe_customers.csv", "ref": "stripe_customers.csv",
             "key": "email", "columns": ["metadata[legacy_account]", "metadata[plan]"], "normalize": ["strip", "lower"], "min_accuracy": 1.0,
             "must_match_keys": emails(out, "plan") + [e for e in emails(out, "zeros") if e not in emails(out, "plan")]},
        ],
    })
    print(f"seed={seed}: {len(rows)} export rows, {len(customers)} customers; " + ", ".join(f"{t}={len(keys(t))}" for t in ("zeros", "ccy", "addr", "geo", "plan")))


def emails(out, tag):
    return sorted(c["email"] for c in out if tag in c["tags"])


if __name__ == "__main__":
    emit(argparse_seed())
