"""Every unit is on the roll with the right occupant, monthly rent and renewal flag.

A unit's rows are found by its id ("A-101", "A101") or by the unit number on a row (or under a section
heading, or on a sheet) that names its building. For each unit:
  * vacant units: some row says vacant and does not carry the tenant who moved out;
  * occupied units: some row carries the current tenant's last name and the monthly rent (lease rent plus
    the month-to-month fee where it applies);
  * renewal flag: flagged exactly when the lease ends October 1 - December 31. A row counts as flagged when a
    cell says "expiring" (not "expired"), when a yes/true/x sits under a column headed expiring/renewal/ending,
    or when the row sits on a sheet or under a heading about expiring leases or renewals.
Expected values come from reference/rent_roll.csv, so nothing here is pinned to a seed.
"""
import csv
import glob
import importlib.util
import os
import re

FLAG_TEXT = re.compile(r"\bexpir(ing|es soon|y soon)\b|\brenewals? (due|needed|window)\b|\bup for renewal\b|\bdue for renewal\b|\bto renew\b", re.I)
FLAG_HEADER = re.compile(r"expir|renew|ending|ends by|end by|window|flag", re.I)
TRUTHY = {"yes", "y", "true", "x", "✓", "✔", "flag", "flagged", "renew", "1"}
SECTION = re.compile(r"expir|renew", re.I)


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*\(?-?[$€£]?\s?(\d[\d,]*(?:\.\d+)?)\s*(/\s*mo(nth)?|per month|monthly)?\s*\)?\s*", str(v), re.I)
    return float(m.group(1).replace(",", "")) if m else None


def _alnum(v):
    return re.sub(r"[^a-z0-9]", "", str(v).lower())


def check(ws, ref):
    name = "every unit with tenant, rent and renewal flag"
    hits = sorted(glob.glob(os.path.join(ws, "rent_roll.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "rent_roll.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "rent_roll.xlsx not found"}]
    try:
        with open(os.path.join(ref, "rent_roll.csv"), encoding="utf-8") as f:
            truth = list(csv.DictReader(f))
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]

    ids = {_alnum(t["unit"]): t for t in truth}
    buildings = {t["building"].split()[0].lower() for t in truth}

    # collect (unit_key, row_text, numbers, flagged) for every row that names a unit
    found = {k: [] for k in ids}
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows(values_only=True)]
        sheet_building = next((b for b in buildings if b in sh.title.lower()), None)
        sheet_section = bool(SECTION.search(sh.title))
        section_building, section_flag, header = None, False, []
        for row in grid:
            cells = ["" if v is None else str(v).strip() for v in row]
            nonempty = [c for c in cells if c]
            text = " ".join(nonempty).lower()
            keys = [ids_key for ids_key in (_alnum(c) for c in nonempty) if ids_key in ids]
            if not keys:
                b = next((b for b in buildings if b in text), None)
                nums = {re.sub(r"^(#|apt\.?\s*|unit\s*)", "", c.lower()).split(".")[0] for c in nonempty}
                bld = b or section_building or sheet_building
                if bld:
                    keys = [k for k, t in ids.items() if t["building"].split()[0].lower() == bld and t["unit_number"] in nums]
                    if len(keys) > 1:
                        keys = []
            if not keys:
                if not nonempty:
                    section_flag = False
                    continue
                if len(nonempty) <= 2:
                    b = next((b for b in buildings if b in text), None)
                    if b:
                        section_building = b
                    if SECTION.search(text):
                        section_flag = True
                    elif b is None and not any(_num(c) is not None for c in nonempty):
                        section_flag = False
                elif sum(1 for c in nonempty if _num(c) is None) >= 3:
                    header = cells
                continue
            flagged = sheet_section or section_flag or any(FLAG_TEXT.search(c) for c in cells)
            if not flagged:
                for j, c in enumerate(cells):
                    if c.lower() in TRUTHY and j < len(header) and FLAG_HEADER.search(header[j] or ""):
                        flagged = True
                        break
            nums = [n for n in (_num(v) for v in row) if n is not None]
            for k in set(keys):
                found[k].append((text, nums, flagged))

    missing, wrong_occ, wrong_flag = [], [], []
    for k, t in ids.items():
        rows = found[k]
        if not rows:
            missing.append(t["unit"])
            continue
        if t["status"] == "vacant":
            former = t["former_tenant_last"].lower()
            if not any("vacant" in text and not (former and re.search(r"\b" + re.escape(former) + r"\b", text)) for text, _, _ in rows):
                wrong_occ.append(f"{t['unit']} should be vacant")
        else:
            last, rent = t["tenant_last"].lower(), float(t["monthly_rent"])
            if not any(re.search(r"\b" + re.escape(last) + r"\b", text) and any(abs(n - rent) <= 0.5 for n in nums) for text, nums, _ in rows):
                wrong_occ.append(f"{t['unit']} should be {t['tenant_last']} at {rent:,.0f}")
        want = t["expiring"] == "yes"
        got = any(fl for _, _, fl in rows)
        if want != got:
            wrong_flag.append(f"{t['unit']} {'not flagged' if want else 'flagged but should not be'}")
    ok = not missing and not wrong_occ and not wrong_flag
    detail = (f"units missing {missing}; " if missing else "") + (f"occupancy/rent {wrong_occ[:6]}; " if wrong_occ else "") + \
             (f"renewal flags {wrong_flag[:6]}" if wrong_flag else "")
    return [{"name": name, "passed": ok, "detail": detail or f"{len(ids)} units with the right tenant, rent and renewal flag"}]
