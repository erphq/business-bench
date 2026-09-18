"""card-expense-coding: the review list.

Required: every charge the policy cannot code (blank account) and every personal charge, by txn_id. For each trip day
whose meals ran over the per diem, at least one of that day's meal charges (the policy puts the overage on the largest
meal; a list that names each meal that day is also accepted). Nothing else: no ordinary business charge and no declined
authorisation. Ids come from reference/owes.json.
"""
import csv
import glob
import io
import json
import os
import re


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _ids(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        vals = [list(r) for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        vals = list(csv.reader(io.StringIO(text)))
    if not vals:
        return None
    hdr = [re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_") for h in vals[0]]
    if "txn_id" not in hdr:
        return None
    j = hdr.index("txn_id")
    return {str(r[j]).strip().upper() for r in vals[1:] if j < len(r) and r[j] is not None and str(r[j]).strip()}


def check(ws, ref):
    name = "review list"
    p = _find(ws, "review.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "review.csv not found"}]
    got = _ids(p)
    if got is None:
        return [{"name": name, "passed": False, "detail": "review.csv has no txn_id column"}]
    owes = json.load(open(os.path.join(ref, "owes.json")))
    required = {i.upper() for i in owes["review_required"]}
    over_days = [[i.upper() for i in ids] for ids in owes["meal_days_over"]]
    allowed = required | {i for ids in over_days for i in ids}
    missing = sorted(required - got)
    days_missing = [ids for ids in over_days if not (set(ids) & got)]
    extra = sorted(got - allowed)
    ok = not missing and not days_missing and not extra
    return [{"name": name, "passed": ok,
             "detail": f"{len(got)} listed; missing={missing[:8]}; over-per-diem days not listed={len(days_missing)}; should not be listed={extra[:8]}"}]
