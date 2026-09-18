#!/usr/bin/env python3
"""shared-cost-allocation: a cabinet shop's August shared costs to its five departments by headcount and floor space.

    python gen.py [--seed N] [--naive DIR]

Business: Hollow Creek Woodworks builds and installs custom cabinetry. The controller allocates the month's shared
overhead to departments: building costs by square feet, people costs by headcount. The GL export holds the month's
lines, HR's roster and the floor plan sit in one workbook, and the controller's memo carries the rules, including
where the rounding residual goes.

Traps (each caught by a check, see task.yaml):
  * building accounts go by square feet and people accounts by headcount          (checks: every department total)
  * the common area and the sublet space are not part of the square-foot base    (checks: Mill Shop total; Finishing total)
  * headcount at month end: part-timers are a half, a leaver, a September starter
    and a 1099 contractor are out                                               (checks: Install total; Showroom total)
  * lines already coded to a department are direct costs, not shared            (checks: Finishing total; Showroom total)
  * a utility refund nets against its account                                    (check: grand total)
  * each department's piece rounds to the cent and the residual goes to the
    department with the largest base for that account                           (check: every account splits to the cent)
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from fractions import Fraction

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "lib"))
from bizgen import *  # noqa: E402,F403
from openpyxl.utils import get_column_letter  # noqa: E402

DEPTS = [("MILL", "Mill Shop"), ("FIN", "Finishing"), ("INST", "Install"), ("SHOW", "Showroom"), ("ADM", "Admin")]
SQFT_ACCTS = [("6100", "Rent"), ("6120", "Property insurance"), ("6210", "Electric"), ("6220", "Gas"), ("6230", "Water & sewer"),
              ("6300", "Janitorial"), ("6310", "Building repairs")]
HC_ACCTS = [("6400", "IT support"), ("6410", "Software subscriptions"), ("6500", "Payroll service"), ("6510", "HR & benefits platform"),
            ("6600", "Phones & internet"), ("6700", "Office supplies")]
VENDORS = {"6100": ["Creekside Industrial Properties"], "6120": ["Everline Insurance"], "6210": ["Valley Electric Co-op"],
           "6220": ["Northern Gas Utility"], "6230": ["City of Ashford Utilities"], "6300": ["Brightway Cleaning"],
           "6310": ["Hollis Door & Dock", "Apex Fire Protection", "Summit HVAC Service"], "6400": ["TechBridge IT"],
           "6410": ["Microsoft 365", "QuickBooks Online", "SketchUp Pro", "Cabinet Vision", "Adobe"], "6500": ["Gusto"],
           "6510": ["Justworks"], "6600": ["Lumen Business", "Verizon Wireless"], "6700": ["Staples", "Amazon Business", "Uline"]}


def cents(x: Fraction) -> int:
    return int((Decimal(x.numerator) / Decimal(x.denominator)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def build(seed: int) -> dict:
    r = rng(seed)
    # ---- roster
    names = people(r, 40)
    roster = []
    plan = {"MILL": (8, 1), "FIN": (4, 2), "INST": (9, 0), "SHOW": (2, 2), "ADM": (3, 1)}
    k = 0
    for code, (ft, pt) in plan.items():
        for i in range(ft + pt):
            f, l = names[k]; k += 1
            roster.append({"name": f"{f} {l}", "dept": code, "status": "Full-time" if i < ft else "Part-time",
                           "hired": day_in(r, date(2015, 1, 1), date(2026, 6, 30)), "term": None})
    f, l = names[k]; k += 1
    roster.append({"name": f"{f} {l}", "dept": "INST", "status": "Full-time", "hired": date(2023, 4, 3), "term": date(2026, 8, r.randint(7, 21))})
    f, l = names[k]; k += 1
    roster.append({"name": f"{f} {l}", "dept": "SHOW", "status": "Full-time", "hired": date(2026, 9, r.randint(1, 14)), "term": None})
    f, l = names[k]; k += 1
    roster.append({"name": f"{f} {l}", "dept": "MILL", "status": "1099 contractor", "hired": date(2026, 5, 11), "term": None})
    f, l = names[k]; k += 1
    roster.append({"name": f"{f} {l}", "dept": "FIN", "status": "Part-time", "hired": date(2024, 2, 5), "term": date(2026, 7, r.randint(10, 30))})
    asof = date(2026, 8, 31)
    hc = {c: Fraction(0) for c, _ in DEPTS}
    for p in roster:
        if p["status"] == "1099 contractor" or p["hired"] > asof or (p["term"] and p["term"] <= asof):
            continue
        hc[p["dept"]] += Fraction(1) if p["status"] == "Full-time" else Fraction(1, 2)

    # ---- floor plan
    areas = [("Mill room and lumber rack", "MILL", r.randint(58, 70) * 100), ("Assembly bay", "MILL", r.randint(14, 20) * 100),
             ("Spray booth and drying room", "FIN", r.randint(22, 30) * 100), ("Install staging and van bay", "INST", r.randint(16, 22) * 100),
             ("Showroom and design studio", "SHOW", r.randint(15, 19) * 100), ("Offices", "ADM", r.randint(8, 11) * 100),
             ("Break room and restrooms (common)", "", r.randint(6, 8) * 100), ("Suite B - sublet to Rivera Upholstery", "", r.randint(11, 14) * 100)]
    sq = {c: sum(a[2] for a in areas if a[1] == c) for c, _ in DEPTS}

    # ---- GL lines for August
    lines = []

    def line(acct, amt, vendor=None, dept="", memo="", day=None):
        lines.append({"date": day or date(2026, 8, r.randint(1, 28)), "acct": acct, "vendor": vendor or r.choice(VENDORS[acct]),
                      "amt": amt, "dept": dept, "memo": memo})

    line("6100", r.randint(138, 152) * 10000, "Creekside Industrial Properties", memo="August rent - 1450 Mill Race Rd", day=date(2026, 8, 1))
    line("6120", r.randint(98000, 132000), memo="Property policy monthly installment")
    elec = r.randint(290000, 380000)
    line("6210", elec, memo="Service 07/15-08/14")
    line("6210", -r.randint(15000, 26000), memo="Billing adjustment credit - meter misread in June")
    line("6220", r.randint(38000, 61000), memo="Gas service July")
    line("6230", r.randint(21000, 34000), memo="Water & sewer bi-monthly")
    line("6300", r.randint(130000, 152000), "Brightway Cleaning", memo="Nightly cleaning - August")
    line("6310", r.randint(42000, 88000), "Hollis Door & Dock", memo="Overhead door spring replacement")
    line("6310", r.randint(31000, 56000), "Apex Fire Protection", memo="Sprinkler inspection")
    line("6310", r.randint(64000, 98000), "Summit HVAC Service", dept="FIN", memo="Spray booth exhaust fan motor")
    line("6400", r.randint(120000, 135000), "TechBridge IT", memo="Managed IT - August")
    for v, lo, hi in (("Microsoft 365", 38000, 46000), ("QuickBooks Online", 20000, 24000), ("Adobe", 8000, 12000)):
        line("6410", r.randint(lo, hi), v, memo="Monthly subscription")
    line("6410", r.randint(29000, 36000), "Cabinet Vision", dept="SHOW", memo="Design seats - showroom")
    line("6500", r.randint(26000, 33000), "Gusto", memo="Payroll processing")
    line("6510", r.randint(36000, 47000), "Justworks", memo="Benefits platform")
    line("6600", r.randint(31000, 38000), "Lumen Business", memo="Fiber internet")
    line("6600", r.randint(19000, 27000), "Verizon Wireless", memo="Crew phones")
    for v in ("Staples", "Amazon Business", "Uline"):
        line("6700", r.randint(4000, 29000), v, memo="Supplies")
    lines.sort(key=lambda x: (x["date"], x["acct"]))

    # ---- truth
    splits, net = {}, {}
    for acct, label in SQFT_ACCTS + HC_ACCTS:
        base = sq if (acct, label) in SQFT_ACCTS else hc
        amt = sum(x["amt"] for x in lines if x["acct"] == acct and not x["dept"])
        net[acct] = amt
        total = sum(base.values())
        exact = {c: Fraction(amt) * Fraction(base[c]) / Fraction(total) for c, _ in DEPTS}
        pieces = {c: cents(v) for c, v in exact.items()}
        largest = max(base, key=lambda c: base[c])
        pieces[largest] += amt - sum(pieces.values())
        splits[acct] = {"label": label, "basis": "sqft" if base is sq else "headcount", "amt": amt, "pieces": pieces, "exact": exact,
                        "largest": largest}
    dept_tot = {c: sum(s["pieces"][c] for s in splits.values()) for c, _ in DEPTS}
    return {"roster": roster, "areas": areas, "lines": lines, "hc": hc, "sq": sq, "splits": splits, "net": net, "dept_tot": dept_tot,
            "grand": sum(net.values())}


def naive_split(d: dict) -> dict:
    """Everything by headcount off the raw roster (everyone counts as one), direct-coded lines included, pieces rounded independently."""
    hc = {c: 0 for c, _ in DEPTS}
    for p in d["roster"]:
        hc[p["dept"]] += 1
    out = {}
    for acct, label in SQFT_ACCTS + HC_ACCTS:
        amt = sum(x["amt"] for x in d["lines"] if x["acct"] == acct)
        out[acct] = {"label": label, "amt": amt, "pieces": {c: cents(Fraction(amt * hc[c], sum(hc.values()))) for c, _ in DEPTS}}
    return out


def acceptable(d: dict) -> bool:
    for base in (d["sq"], d["hc"]):
        vals = sorted(base.values())
        if vals[-1] == vals[-2]:
            return False
    if max(d["hc"], key=lambda c: d["hc"][c]) == max(d["sq"], key=lambda c: d["sq"][c]):
        return False
    residuals = 0
    for s in d["splits"].values():
        for v in s["exact"].values():
            if (v * 1) .denominator != 1 and (v - int(v)) == Fraction(1, 2):
                return False  # an exact half cent would make half-up rounding engine-dependent
        residuals += s["amt"] != sum(cents(v) for v in s["exact"].values())
    if residuals < 4:
        return False
    nv = naive_split(d)
    for c, _ in DEPTS:
        if abs(sum(s["pieces"][c] for s in nv.values()) - d["dept_tot"][c]) < 500:
            return False
    return True


# --------------------------------------------------------------------------- deliverable

def workbook(d: dict, splits: dict, hc: dict, sq: dict) -> dict:
    nd = len(DEPTS)
    bases = [[name, float(hc[c]), sq[c]] for c, name in DEPTS]
    bases.append(["Total", f"=SUM(B2:B{nd + 1})", f"=SUM(C2:C{nd + 1})"])
    tot_row = nd + 2
    rows = []
    for i, (acct, s) in enumerate(splits.items(), start=2):
        basis = "Headcount" if s.get("basis", "headcount") == "headcount" else "Square feet"
        col = "B" if basis == "Headcount" else "C"
        line = [f"{acct} {s['label']}", basis, s["amt"] / 100]
        largest = s.get("largest")
        dept_cols = [get_column_letter(4 + j) for j in range(nd)]
        for j, (c, _) in enumerate(DEPTS):
            if c == largest:
                others = "+".join(f"{dept_cols[k]}{i}" for k in range(nd) if k != j)
                line.append(f"=C{i}-({others})")
            else:
                line.append(f"=ROUND(C{i}*Bases!${col}${2 + j}/Bases!${col}${tot_row},2)")
        line.append(f"=C{i}-SUM({dept_cols[0]}{i}:{dept_cols[-1]}{i})")
        rows.append(line)
    n = len(splits) + 1
    rows.append(["Total", "", f"=SUM(C2:C{n})"] + [f"=SUM({get_column_letter(4 + j)}2:{get_column_letter(4 + j)}{n})" for j in range(nd)]
                + [f"=SUM({get_column_letter(4 + nd)}2:{get_column_letter(4 + nd)}{n})"])
    return {
        "Allocation": {"header": ["Account", "Basis", "Shared cost"] + [name for _, name in DEPTS] + ["Unallocated"], "rows": rows,
                       "number_formats": {get_column_letter(j): "#,##0.00" for j in range(3, 5 + nd)}, "widths": {"A": 30, "B": 12}},
        "Bases": {"header": ["Department", "Headcount at 2026-08-31", "Square feet"], "rows": bases, "widths": {"A": 14, "B": 24, "C": 12}},
    }


def naive_workbook(d: dict) -> dict:
    """Static independently-rounded pieces laid out like the reference (no residual step)."""
    nv = naive_split(d)
    nd = len(DEPTS)
    rows = []
    for i, (acct, s) in enumerate(nv.items(), start=2):
        rows.append([f"{acct} {s['label']}", "Headcount", s["amt"] / 100] + [s["pieces"][c] / 100 for c, _ in DEPTS])
    n = len(nv) + 1
    rows.append(["Total", "", f"=SUM(C2:C{n})"] + [f"=SUM({get_column_letter(4 + j)}2:{get_column_letter(4 + j)}{n})" for j in range(nd)])
    return {"Allocation": {"header": ["Account", "Basis", "Shared cost"] + [name for _, name in DEPTS], "rows": rows}}


def emit(seed: int, naive_dir: str | None) -> None:
    d = build(seed)
    if naive_dir:
        os.makedirs(naive_dir, exist_ok=True)
        write_xlsx(os.path.join(naive_dir, "allocation.xlsx"), naive_workbook(d), creator="naive")
        return
    ws, ref, sol = task_dirs(HERE)
    dept_name = dict(DEPTS)
    acct_name = dict(SQFT_ACCTS + HC_ACCTS)

    # ---- workspace: GL export
    write_csv(os.path.join(ws, "gl_detail_overhead_2026-08.csv"), ["Date", "Account", "Account name", "Vendor", "Memo", "Dept", "Amount"],
              [[x["date"].strftime("%m/%d/%Y"), x["acct"], acct_name[x["acct"]], x["vendor"], x["memo"], x["dept"], money_str(x["amt"] / 100, 7)]
               for x in d["lines"]], preamble=["Hollow Creek Woodworks LLC", "General ledger detail - overhead accounts 6100-6799",
                                                "08/01/2026 - 08/31/2026"], bom=True)

    # ---- workspace: roster and floor plan
    roster_rows = [[p["name"], dept_name[p["dept"]], p["status"], p["hired"], p["term"]] for p in sorted(d["roster"], key=lambda p: (p["dept"], p["name"]))]
    write_xlsx(os.path.join(ws, "headcount_and_floorplan.xlsx"), {
        "Roster": {"merged_title": "Employee roster (HR export 09/02/2026)", "header": ["Name", "Department", "Status", "Hire date", "Termination date"],
                   "rows": roster_rows, "widths": {"A": 22, "B": 14, "C": 16, "D": 12, "E": 16}},
        "Floor plan": {"merged_title": "1450 Mill Race Rd - space by use", "header": ["Area", "Department", "Square feet"],
                       "rows": [[a[0], dept_name.get(a[1], ""), a[2]] for a in d["areas"]], "widths": {"A": 38, "B": 14, "C": 12}},
    }, creator="Hollow Creek")

    # ---- workspace: last month's allocation, done by the previous bookkeeper (a template and a distractor)
    r = rng(seed + 33)
    jul_rows = []
    for acct, label in SQFT_ACCTS + HC_ACCTS:
        amt = d["net"][acct] / 100 * r.uniform(0.9, 1.08)
        shares = [r.uniform(0.1, 0.4) for _ in DEPTS]
        tot = sum(shares)
        jul_rows.append([f"{acct} {label}", round(amt, 2)] + [round(amt * s / tot, 2) for s in shares])
    write_xlsx(os.path.join(ws, "overhead_allocation_2026-07.xlsx"), {"July": {
        "merged_title": "Overhead allocation - July 2026 (prepared by T. Nguyen)",
        "header": ["Account", "Amount"] + [name for _, name in DEPTS], "rows": jul_rows, "widths": {"A": 30}}}, creator="T. Nguyen")

    # ---- workspace: controller memo
    write_text(os.path.join(ws, "memo_overhead_allocation.txt"),
               "To: Bookkeeping\nFrom: Carla Mendes, Controller\nRe: Monthly overhead allocation - how we do it now\n\n"
               "Tomas's July sheet split everything by his own estimates. From August on we allocate on real bases:\n\n"
               "1. Building costs go by square feet: rent, property insurance, electric, gas, water & sewer, janitorial and building\n"
               "   repairs. People costs go by headcount: IT support, software subscriptions, payroll service, the HR & benefits\n"
               "   platform, phones & internet and office supplies.\n"
               "2. Square feet: each department's own space from the floor plan. The common break room and restrooms and the space\n"
               "   we sublet are not a department and stay out of the base.\n"
               "3. Headcount: employees on the payroll on the last day of the month. Part-time counts as one half. Contractors are not\n"
               "   employees. Somebody who left during the month, or who has not started yet, does not count.\n"
               "4. A GL line that already has a department code is a direct cost of that department. It is not overhead and does\n"
               "   not go into the allocation at all. Credits and refunds net against their account.\n"
               "5. Rounding: split each account on its own. Work out each department's share, round it to the cent, and put\n"
               "   whatever difference is left (a cent or two) on the department with the biggest base for that account, so every\n"
               "   account allocates to the penny.\n\n"
               "I want accounts down the side, the five departments across, and department totals, with the formulas left in.\n")

    # ---- reference
    body = [[acct, s["label"], s["basis"], f"{s['amt'] / 100:.2f}"] + [f"{s['pieces'][c] / 100:.2f}" for c, _ in DEPTS] for acct, s in d["splits"].items()]
    write_csv(os.path.join(ref, "allocation.csv"), ["account", "label", "basis", "amount"] + [name for _, name in DEPTS], body)
    write_json(os.path.join(ref, "notes.json"), {
        "splits": {f"{acct} {s['label']}": {dept_name[c]: round(s["pieces"][c] / 100, 2) for c, _ in DEPTS} for acct, s in d["splits"].items()},
        "department_totals": {dept_name[c]: round(v / 100, 2) for c, v in d["dept_tot"].items()}, "grand_total": round(d["grand"] / 100, 2),
        "headcount": {dept_name[c]: float(v) for c, v in d["hc"].items()}, "square_feet": {dept_name[c]: v for c, v in d["sq"].items()}})

    # ---- reference solution
    write_xlsx(os.path.join(sol, "allocation.xlsx"), workbook(d, d["splits"], d["hc"], d["sq"]), creator="Hollow Creek")

    def pin(c):
        return {"type": "xlsx_value_present", "name": f"{dept_name[c]} total", "path": "allocation.xlsx", "expected": round(d["dept_tot"][c] / 100, 2),
                "rel_tol": 0.000001, "near_text": dept_name[c].lower()}

    leaver = next(p for p in d["roster"] if p["term"] and p["term"].month == 8)
    starter = next(p for p in d["roster"] if p["hired"] > date(2026, 8, 31))
    contractor = next(p for p in d["roster"] if p["status"] == "1099 contractor")
    write_task_yaml(HERE, {
        "id": "shared-cost-allocation", "track": "desk", "category": "bookkeeping",
        "title": "Allocate August overhead to departments",
        "ask": "Please do the August overhead allocation to the departments. Carla's memo explains the new way she wants it done. Save it as allocation.xlsx.\n",
        "followup": None, "timeout_s": 1200,
        "traps": [
            "building accounts split by square feet and people accounts by headcount; one base for everything, or July's hand estimates, "
            "moves every department (checks: Mill Shop total; Finishing total; Install total; Showroom total; Admin total)",
            "the floor plan lists the common break room and restrooms and the sublet suite with no department; they stay out of the "
            "square-foot base rather than being spread or kept as a sixth column (checks: Mill Shop total; Finishing total)",
            f"headcount is a month-end count with part-timers at a half: {leaver['name']} left in August, {starter['name']} starts in "
            f"September, {contractor['name']} is a 1099 contractor and one part-timer left in July, all still on the roster export "
            "(checks: Install total; Showroom total)",
            "the Summit HVAC spray-booth repair is coded FIN and the Cabinet Vision seats are coded SHOW; both are direct costs and "
            "stay out of the allocation (checks: Finishing total; Showroom total; grand total allocated)",
            "the electric account carries a negative billing-adjustment credit that nets against the bill, and amounts are text with "
            "thousands separators under a three-line preamble (check: grand total allocated)",
            "each department's piece rounds to the cent and the leftover cent or two goes to the department with the largest base for "
            "that account (Mill Shop for square feet, Install for headcount); rounding every piece independently leaves several "
            "accounts a cent off (check: every account splits to the cent)",
        ],
        "checks": [
            {"type": "xlsx_has_formulas", "name": "live formulas", "path": "allocation.xlsx", "min_count": 10},
            {"type": "xlsx_no_errors", "name": "no formula errors", "path": "allocation.xlsx"},
            {"type": "custom", "name": "every account splits to the cent", "module": "check.py"},
            {"type": "xlsx_value_present", "name": "grand total allocated", "path": "allocation.xlsx", "expected": round(d["grand"] / 100, 2),
             "rel_tol": 0.000001, "near_text": "total"},
        ] + [pin(c) for c, _ in DEPTS],
    })
    print(f"seed={seed} hc={ {k: float(v) for k, v in d['hc'].items()} } sq={d['sq']} grand={d['grand'] / 100:.2f}")
    print({dept_name[c]: v / 100 for c, v in d["dept_tot"].items()})


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--naive", default=None)
    a_ = ap.parse_args()
    for attempt in range(400):
        if acceptable(build(a_.seed * 1000 + attempt)):
            break
    else:
        raise SystemExit("no acceptable draw")
    emit(a_.seed * 1000 + attempt, a_.naive)
