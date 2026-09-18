"""Policy section per flagged line.

Every line the reference flags must appear in violations.csv with its policy section written somewhere on that row
(a policy_section column, or the reason text: "3.1", "Section 3.1", "§3.1" and "Policy 3.1 lodging cap" all pass).
Rows are matched on line_id, case-insensitively. Sections come from reference/violations.csv.
"""
import csv
import glob
import io
import os
import re


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _norm(c):
    return re.sub(r"[^a-z0-9]+", "_", str(c).strip().lower()).strip("_")


def _rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        vals = [list(r) for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
        if not vals:
            return []
        hdr = [_norm(h) for h in vals[0]]
        return [{h: ("" if v is None else str(v)) for h, v in zip(hdr, row)} for row in vals[1:]]
    text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
    rd = list(csv.reader(io.StringIO(text)))
    if not rd:
        return []
    hdr = [_norm(h) for h in rd[0]]
    return [dict(zip(hdr, row)) for row in rd[1:]]


def check(ws, ref):
    name = "policy section per flagged line"
    p = _find(ws, "violations.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "violations.csv not found"}]
    try:
        rows = _rows(p)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"unreadable: {e}"}]
    if not rows or "line_id" not in rows[0]:
        return [{"name": name, "passed": False, "detail": f"need a line_id column; have {list(rows[0]) if rows else []}"}]
    got = {}
    for row in rows:
        k = str(row.get("line_id", "")).strip().lower()
        text = " ".join(str(v) for c, v in row.items() if c != "line_id")
        got[k] = got.get(k, "") + " " + text
    want = _rows(os.path.join(ref, "violations.csv"))
    wrong = []
    for w in want:
        k = w["line_id"].strip().lower()
        sec = re.escape(w["policy_section"].strip())
        if k not in got or not re.search(r"(?<![\d.])" + sec + r"(?![\d])", got[k]):
            wrong.append(f"{w['line_id']} (want {w['policy_section']})")
    return [{"name": name, "passed": not wrong,
             "detail": f"{len(want) - len(wrong)}/{len(want)} flagged lines cite the right section" + (f"; wrong: {wrong}" if wrong else "")}]
