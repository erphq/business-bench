#!/usr/bin/env python3
"""Generator for desk/monthly-report.

Writes workspace/sales_export.csv, reference/revenue_by_region_month.csv,
reference/region_totals.csv and reference_solution/{report.xlsx, memo.md}.

Deterministic: same --seed -> byte-identical output (fixed workbook properties
and normalised zip timestamps). Other seeds re-roll order ids, amounts, dates,
which rows are duplicated/refunded and the date-format mix, but keep the traps:

  * order_date in three mixed formats (2026-03-04, 03/04/2026, 4 Mar 2026)
  * amount as "$1,234.50"; refunds as "(120.00)" (negative, parentheses)
  * 12 fully duplicated rows (same order_id, identical fields) -> count once
  * 14 refund rows (order_id = original + "-R", same month) -> netted
  * North has NO rows at all in March 2026 (data gap, not zero sales)

reference_solution/report.xlsx: Data sheet (cleaned rows) + Report sheet driven
by SUMIFS/COUNTIFS formulas + a static Snapshot sheet with the same numbers.
The Snapshot sheet exists because the shipped bench/grade.py cannot read
recalculated formula values (its *.xlsx glob misses the upper-cased file the
`formulas` engine writes); with that one-line fix the Report sheet alone passes.
Use --no-snapshot to emit the formulas-only workbook.
"""
from __future__ import annotations
import argparse, calendar, csv, os, random, re, zipfile
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
REGIONS = ["North", "South", "East", "West"]
LINES = [("Hardware", (400, 8900), 0.35), ("Subscriptions", (80, 1500), 0.40), ("Services", (250, 6000), 0.25)]
MONTHS = [(2026, m) for m in range(1, 7)]
GAP = ("North", 3)          # region/month with no rows at all
N_DUPS = 12
N_REFUNDS = 14
SENSITIVE = [("South", 2), ("West", 6)]  # cells checked in task.yaml; forced to contain a dup and a refund


def ym(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def fmt_date(d: date, style: int) -> str:
    if style == 0:
        return d.isoformat()
    if style == 1:
        return f"{d.month:02d}/{d.day:02d}/{d.year}"
    return f"{d.day} {d.strftime('%b')} {d.year}"


def fmt_amount(a: float) -> str:
    return f"({abs(a):,.2f})" if a < 0 else f"${a:,.2f}"


def normalize_zip(path: str):
    """Rewrite an xlsx zip with fixed entry timestamps so bytes are reproducible."""
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            if name == "docProps/core.xml":  # openpyxl re-stamps modified at save time
                data = re.sub(rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                              b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-07-01T00:00:00Z</dcterms:modified>', data)
            zi = zipfile.ZipInfo(name, date_time=(2026, 7, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)


def pick_line(rng: random.Random):
    r = rng.random()
    acc = 0.0
    for name, rngamt, w in LINES:
        acc += w
        if r <= acc:
            return name, rngamt
    return LINES[-1][0], LINES[-1][1]


def build(rng: random.Random):
    orders = []
    for region in REGIONS:
        for (y, m) in MONTHS:
            if (region, m) == GAP:
                continue
            for _ in range(rng.randint(35, 41)):
                day = rng.randint(1, calendar.monthrange(y, m)[1])
                line, (lo, hi) = pick_line(rng)
                orders.append(dict(date=date(y, m, day), region=region, line=line,
                                   amount=round(rng.uniform(lo, hi), 2), k=rng.random()))
    orders.sort(key=lambda o: (o["date"], o["k"]))
    for i, o in enumerate(orders):
        o["order_id"] = f"SO-{10001 + i}"
        o["style"] = rng.randrange(3)

    def choose(pool_filter, n, forced_cells, min_per_region=2):
        chosen, used = [], set()
        for (reg, mo) in forced_cells:
            cands = [o for o in orders if o["region"] == reg and o["date"].month == mo and pool_filter(o) and id(o) not in used]
            pick = rng.choice(cands); chosen.append(pick); used.add(id(pick))
        for reg in REGIONS:
            while sum(1 for c in chosen if c["region"] == reg) < min_per_region:
                cands = [o for o in orders if o["region"] == reg and pool_filter(o) and id(o) not in used]
                pick = rng.choice(cands); chosen.append(pick); used.add(id(pick))
        while len(chosen) < n:
            cands = [o for o in orders if pool_filter(o) and id(o) not in used]
            pick = rng.choice(cands); chosen.append(pick); used.add(id(pick))
        return chosen

    # refunds: own order_id (orig + "-R"), same month as the original sale
    refund_src = choose(lambda o: o["amount"] >= 2000, N_REFUNDS, SENSITIVE)
    refunds = []
    for o in refund_src:
        last = calendar.monthrange(o["date"].year, o["date"].month)[1]
        d = min(o["date"] + timedelta(days=rng.randint(2, 9)), date(o["date"].year, o["date"].month, last))
        frac = 1.0 if (o["region"], o["date"].month) in SENSITIVE else rng.choice([1.0, 1.0, 0.5])
        refunds.append(dict(date=d, region=o["region"], line=o["line"], amount=-round(o["amount"] * frac, 2),
                            k=rng.random(), order_id=o["order_id"] + "-R", style=rng.randrange(3)))
    refund_ids = {r["order_id"] for r in refunds}

    clean = sorted(orders + refunds, key=lambda o: (o["date"], o["k"]))
    rows = [dict(order_id=o["order_id"], order_date=fmt_date(o["date"], o["style"]), region=o["region"],
                 product_line=o["line"], amount=fmt_amount(o["amount"]), _d=o["date"], _a=o["amount"]) for o in clean]

    # exact duplicates of existing sale rows, dropped at random positions
    dup_src = choose(lambda o: o["amount"] >= 2000, N_DUPS, SENSITIVE)
    dup_ids = {o["order_id"] for o in dup_src}
    export = list(rows)
    for r in [r for r in rows if r["order_id"] in dup_ids]:
        export.insert(rng.randint(0, len(export)), dict(r))
    return rows, export, refund_ids, dup_ids


def acceptable(rows, export, refund_ids, dup_ids) -> bool:
    """Grader-precision guard. xlsx_value_present matches any cell on a row that
    mentions the region, so (1) no other correct cell on a sensitive region's row
    may sit within 2% of the checked cell, (2) the three naive miscomputations
    (no dedupe + parentheses ignored, no dedupe, refunds dropped) must move the
    checked cell and every region total by more than 1%, and (3) each region's
    duplicates and refunds must each be worth more than 1% of its total."""
    months = [f"{y}-{m:02d}" for (y, m) in MONTHS]
    def agg(src, val):
        out = {}
        for r in src:
            k = (r["region"], ym(r["_d"]))
            out[k] = out.get(k, 0.0) + val(r)
        return out
    ref = agg(rows, lambda r: r["_a"])
    variants = [agg(export, lambda r: abs(r["_a"])), agg(export, lambda r: r["_a"]),
                agg([r for r in rows if r["_a"] > 0], lambda r: r["_a"])]
    def total(d, reg):
        return sum(v for (r_, _), v in d.items() if r_ == reg)
    for (reg, mo) in SENSITIVE:
        key = (reg, f"2026-{mo:02d}")
        v = ref[key]
        others = [ref[(reg, m)] for m in months if (reg, m) in ref and m != key[1]] + [total(ref, reg)]
        if any(abs(o - v) <= 0.02 * v for o in others):
            return False
        for var in variants:
            if any(abs(var.get((reg, m), 0.0) - v) <= 0.01 * v for m in months):
                return False
    for reg in REGIONS:
        t = total(ref, reg)
        if any(abs(total(var, reg) - t) <= 0.01 * t for var in variants):
            return False
        dup_sum = sum(r["_a"] for r in rows if r["region"] == reg and r["order_id"] in dup_ids)
        ref_sum = sum(-r["_a"] for r in rows if r["region"] == reg and r["order_id"] in refund_ids)
        if dup_sum <= 0.01 * t or ref_sum <= 0.01 * t:
            return False
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--no-snapshot", action="store_true", help="omit the static Snapshot sheet (formulas-only workbook)")
    args = ap.parse_args()
    for attempt in range(200):
        rng = random.Random(args.seed * 1000 + attempt)
        rows, export, refund_ids, dup_ids = build(rng)
        if acceptable(rows, export, refund_ids, dup_ids):
            break
    else:
        raise SystemExit("no acceptable draw in 200 attempts")

    ws = os.path.join(HERE, "workspace"); ref = os.path.join(HERE, "reference"); sol = os.path.join(HERE, "reference_solution")
    for d in (ws, ref, sol):
        os.makedirs(d, exist_ok=True)
        for f in os.listdir(d):
            os.remove(os.path.join(d, f))

    cols = ["order_id", "order_date", "region", "product_line", "amount"]
    with open(os.path.join(ws, "sales_export.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(cols)
        for r in export:
            w.writerow([r[c] for c in cols])

    # ---- reference ----
    rev, cnt = {}, {}
    for r in rows:
        key = (r["region"], ym(r["_d"]))
        rev[key] = round(rev.get(key, 0.0) + r["_a"], 2); cnt[key] = cnt.get(key, 0) + 1
    months = [f"{y}-{m:02d}" for (y, m) in MONTHS]
    with open(os.path.join(ref, "revenue_by_region_month.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["region", "month", "revenue", "rows", "note"])
        for reg in REGIONS:
            for mo in months:
                if (reg, mo) in rev:
                    w.writerow([reg, mo, f"{rev[(reg, mo)]:.2f}", cnt[(reg, mo)], ""])
                else:
                    w.writerow([reg, mo, "", 0, "no rows in export (data gap, not zero sales)"])
    totals = {reg: round(sum(v for (r_, m_), v in rev.items() if r_ == reg), 2) for reg in REGIONS}
    with open(os.path.join(ref, "region_totals.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["region", "revenue"])
        for reg in REGIONS:
            w.writerow([reg, f"{totals[reg]:.2f}"])
        w.writerow(["ALL", f"{round(sum(totals.values()), 2):.2f}"])

    # ---- reference solution ----
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.properties.created = datetime(2026, 7, 1); wb.properties.modified = datetime(2026, 7, 1)
    wb.properties.creator = "reference"; wb.properties.lastModifiedBy = "reference"
    dsh = wb.active; dsh.title = "Data"
    dsh.append(["order_id", "order_date", "month", "region", "product_line", "amount"])
    for r in rows:
        dsh.append([r["order_id"], r["_d"].isoformat(), ym(r["_d"]), r["region"], r["product_line"], r["_a"]])
    n = dsh.max_row
    rp = wb.create_sheet("Report")
    rp.append(["Region"] + months + ["Total"])
    for c in rp[1]:
        c.font = Font(bold=True)
    for i, reg in enumerate(REGIONS, start=2):
        line = [reg]
        for j in range(len(months)):
            L = chr(ord("B") + j)
            line.append(f'=IF(COUNTIFS(Data!$D$2:$D${n},$A{i},Data!$C$2:$C${n},{L}$1)=0,"no data",'
                        f'SUMIFS(Data!$F$2:$F${n},Data!$D$2:$D${n},$A{i},Data!$C$2:$C${n},{L}$1))')
        line.append(f"=SUM(B{i}:G{i})")
        rp.append(line)
    rp.append(["Total"] + [f"=SUM({chr(ord('B') + j)}2:{chr(ord('B') + j)}5)" for j in range(len(months) + 1)])
    rp.append([])
    rp.append(["Note: North has no rows for 2026-03 in the export - a data gap, not zero sales. "
               "Duplicate rows removed and refunds netted on the Data sheet."])
    rp.column_dimensions["A"].width = 12
    if not args.no_snapshot:
        sn = wb.create_sheet("Snapshot")
        sn.append(["Static copy of the Report values (net revenue, USD)"])
        sn.append(["Region"] + months + ["Total"])
        for reg in REGIONS:
            line = [reg] + [rev.get((reg, mo), "no data") for mo in months] + [totals[reg]]
            sn.append(line)
        sn.append(["Total"] + [round(sum(rev.get((r_, mo), 0.0) for r_ in REGIONS), 2) for mo in months]
                  + [round(sum(totals.values()), 2)])
    p = os.path.join(sol, "report.xlsx"); wb.save(p); normalize_zip(p)

    n_dup = len(dup_ids); n_ref = len(refund_ids)
    ref_total = round(sum(r["_a"] for r in rows if r["order_id"] in refund_ids), 2)
    north5 = totals["North"]
    memo = f"""# Revenue by region, January-June 2026: three things to know

1. **North has no March data.** The export contains no rows at all for the North region in
   March 2026 - it is missing from the file, not a month of zero sales. The report shows
   "no data" for that cell and North's half-year total ({north5:,.2f}) covers five months only.
   Please have the March North rows re-exported before this goes anywhere.
2. **{n_dup} duplicate rows were removed.** {n_dup} orders appear twice in the export with the same
   order_id, date and amount. Each is counted once in the report.
3. **Refunds are netted.** {n_ref} refund rows (negative amounts shown in parentheses, ids ending in
   "-R") total {ref_total:,.2f} and are subtracted in the month they were issued. Region totals
   in the report are net of these refunds.

Half-year net revenue: North {totals['North']:,.2f} (5 months) · South {totals['South']:,.2f} ·
East {totals['East']:,.2f} · West {totals['West']:,.2f} · All regions {sum(totals.values()):,.2f}.

Report.xlsx: the Report sheet is driven by SUMIFS formulas over the cleaned Data sheet
(deduplicated, dates normalised, amounts parsed to numbers).
"""
    with open(os.path.join(sol, "memo.md"), "w") as f:
        f.write(memo)

    print(f"seed={args.seed} attempt={attempt} export_rows={len(export)} unique_rows={len(rows)} dups={n_dup} refunds={n_ref} refund_total={ref_total}")
    print("region totals:", {k: f"{v:.2f}" for k, v in totals.items()})
    print("sensitive cells:", {f"{r}/{y}-{m:02d}": f"{rev[(r, f'{y}-{m:02d}')]:.2f}" for (r, m) in SENSITIVE for y in [2026]})
    print("task.yaml expected values are pinned to seed 0; update them if you re-roll the seed.")


if __name__ == "__main__":
    main()
