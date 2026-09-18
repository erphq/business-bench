"""bank-reconciliation custom check, layout-agnostic.

1. both sides tie: at least two cells on rows labelled adjusted / corrected / reconciled equal the true adjusted balance
   (the bank side and the book side), after recalculation.
2. every outstanding check at 31 August, including the one carried from July, sits on a row with its amount.
3. the check keyed with swapped digits has its correction amount somewhere in the workbook.
4. a difference / variance / unreconciled row exists whose figures are all zero.
5. the matched list names the checks that cleared (at least 90% of their numbers appear).
Figures come from reference/rec.json, so nothing here is pinned to a seed.
"""
import glob
import json
import os
import re


def _grade_module():
    import importlib.util
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


ADJ = re.compile(r"\b(adjusted|corrected|reconciled)\b", re.I)
DIFF = re.compile(r"\b(difference|variance|unreconciled|out of balance|discrepancy)\b", re.I)


def check(ws, ref):
    out = []
    p = _find(ws, "reconciliation.xlsx")
    if not p:
        return [{"name": "reconciliation.xlsx present", "passed": False, "detail": "not found"}]
    try:
        rec = json.load(open(os.path.join(ref, "rec.json")))
    except Exception as e:
        return [{"name": "reference", "passed": False, "detail": f"unreadable: {e}"}]
    try:
        g = _grade_module()
        from openpyxl import load_workbook
        wb = load_workbook(g.recalculated_workbook(p), data_only=True)
        cell_num = g.cell_num
    except Exception as e:
        return [{"name": "workbook readable", "passed": False, "detail": f"{type(e).__name__}: {e}"}]

    rows = []
    for sh in wb.worksheets:
        for row in sh.iter_rows(values_only=True):
            vals = [v for v in row if v is not None and str(v).strip() != ""]
            if not vals:
                continue
            texts = [str(v) for v in vals if not isinstance(v, (int, float)) or isinstance(v, bool)]
            nums = [x for x in (cell_num(v) for v in vals) if x is not None]
            rows.append((" ".join(str(v) for v in vals), texts, nums))

    adj = rec["adjusted_balance"]
    hits = sum(1 for full, texts, nums in rows if any(ADJ.search(t) for t in texts) for x in nums if abs(x - adj) <= 0.011)
    out.append({"name": "adjusted bank and book balances both equal", "passed": hits >= 2,
                "detail": f"{hits} cell(s) on adjusted/corrected/reconciled rows equal {adj:,.2f} (need 2)"})

    missing = []
    for oc in rec["outstanding_checks"]:
        pat = re.compile(rf"(?<!\d){oc['num']}(?!\d)")
        if not any(pat.search(full) and any(abs(abs(x) - oc["amount"]) <= 0.011 for x in nums) for full, texts, nums in rows):
            missing.append(f"{oc['num']} ({oc['amount']:,.2f})")
    out.append({"name": "outstanding checks listed with amounts", "passed": not missing,
                "detail": "all listed" if not missing else f"not on a row with their amount: {missing}"})

    diff = rec["swap"]["difference"]
    ok = any(abs(abs(x) - diff) <= 0.011 for full, texts, nums in rows for x in nums)
    out.append({"name": f"check {rec['swap']['num']} correction shown", "passed": ok,
                "detail": f"{diff:,.2f} {'found' if ok else 'not found'}"})

    zero_rows = [full for full, texts, nums in rows if any(DIFF.search(t) for t in texts) and nums and all(abs(x) <= 0.005 for x in nums)]
    out.append({"name": "difference row is zero", "passed": bool(zero_rows),
                "detail": f"zero difference row: {zero_rows[0][:80]!r}" if zero_rows else "no difference/variance row whose figures are all zero"})

    blob = " ".join(full for full, _, _ in rows)
    nums_cleared = rec["cleared_check_numbers"]
    seen = [n for n in nums_cleared if re.search(rf"(?<!\d){n}(?!\d)", blob)]
    frac = len(seen) / len(nums_cleared) if nums_cleared else 1.0
    out.append({"name": "cleared checks listed as matched", "passed": frac >= 0.9,
                "detail": f"{len(seen)}/{len(nums_cleared)} cleared check numbers appear"})
    return out
