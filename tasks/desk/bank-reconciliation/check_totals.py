"""bank-reconciliation: the outstanding-checks total and the deposits-in-transit total.

Same rule as the grader's xlsx_value_present (a cell equal to the figure, after recalculation, on a row or in a column
mentioning the label) except that the sign is ignored: "Less: outstanding checks (7,700.04)" and a positive total
subtracted by formula are both right. Figures come from reference/rec.json.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from check import _find, _grade_module  # noqa: E402


def _present(wb, cell_num, expected, near):
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows(values_only=True)]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        col_text = [" ".join(str(r[j]).lower() for r in grid if j < len(r) and r[j] is not None) for j in range(ncols)]
        for r in grid:
            row_text = " ".join(str(v).lower() for v in r if v is not None)
            for j, v in enumerate(r):
                if v is None or isinstance(v, bool):
                    continue
                x = cell_num(v)
                if x is None or abs(abs(x) - expected) > 0.011:
                    continue
                if near in row_text or near in col_text[j]:
                    return f"{sh.title}: {v}"
    return None


def check(ws, ref):
    p = _find(ws, "reconciliation.xlsx")
    if not p:
        return [{"name": "reconciliation.xlsx present", "passed": False, "detail": "not found"}]
    rec = json.load(open(os.path.join(ref, "rec.json")))
    try:
        g = _grade_module()
        from openpyxl import load_workbook
        wb = load_workbook(g.recalculated_workbook(p), data_only=True)
    except Exception as e:
        return [{"name": "workbook readable", "passed": False, "detail": f"{type(e).__name__}: {e}"}]
    out = []
    for label, key, near in (("outstanding checks total", "outstanding_checks_total", "outstanding"),
                             ("deposits in transit total", "deposits_in_transit_total", "transit")):
        hit = _present(wb, g.cell_num, rec[key], near)
        out.append({"name": label, "passed": bool(hit), "detail": f"{rec[key]:,.2f} found at {hit}" if hit else f"no cell = ±{rec[key]:,.2f} near {near!r}"})
    return out
