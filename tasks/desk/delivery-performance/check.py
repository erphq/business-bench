"""Per-hub on-time percentage, accepted as a fraction (0.8140) or a percentage (81.4 / "81.4%").

xlsx_value_present pins one numeric form and a rate is legitimately written either way, so this check
recalculates the workbook itself and accepts both. Expected rates come from reference/notes.json, so
nothing here is pinned to a seed.
"""
import glob
import importlib.util
import json
import os
import re


def _recalculated(path):
    """Reuse the grader's own recalculation (LibreOffice, else the `formulas` engine)."""
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*\(?-?[$€£]?\s?(\d[\d,]*(?:\.\d+)?)\s?%?\)?\s*", str(v))
    return float(m.group(1).replace(",", "")) if m else None


def check(ws, ref):
    name = "on-time rates per hub"
    hits = sorted(glob.glob(os.path.join(ws, "delivery.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "delivery.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "delivery.xlsx not found"}]
    try:
        rates = json.load(open(os.path.join(ref, "notes.json")))["on_time_rates"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]

    cells = []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        col_text = [""] * ncols
        for r in grid:
            for j, c in enumerate(r):
                if c.value is not None:
                    col_text[j] += " " + str(c.value).lower()
        for r in grid:
            row_text = " ".join(str(c.value).lower() for c in r if c.value is not None)
            for j, c in enumerate(r):
                v = _num(c.value)
                if v is not None:
                    cells.append((row_text, col_text[j], v))

    bad = []
    for hub, want in sorted(rates.items()):
        key = hub.lower()
        ok = False
        for row_text, col_text, v in cells:
            if key not in row_text and key not in col_text:
                continue
            if abs(v - want) <= max(want * 0.01, 0.002) or abs(v - want * 100) <= max(want, 0.2):
                ok = True
                break
        if not ok:
            bad.append(f"{hub}: expected {want:.4f} (or {want * 100:.2f}%)")
    return [{"name": name, "passed": not bad,
             "detail": "; ".join(bad) if bad else f"{len(rates)} hub rates match"}]
