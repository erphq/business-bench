"""Every dish carries its plate cost and its margin, in whichever form the workbook states the margin.

Plate cost is one number. What the owner "keeps" is legitimately written as a dollar margin, a margin
percentage (0.684 or 68.4%) or its complement, the food cost percentage; any of them passes as long as
it sits on the dish's row (or column) next to the right plate cost. Expected figures come from
reference/notes.json, so nothing here is pinned to a seed.
"""
import glob
import importlib.util
import json
import os
import re


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*\(?-?[$€£]?\s?(\d[\d,]*(?:\.\d+)?)\s?%?\)?\s*", str(v))
    return float(m.group(1).replace(",", "")) if m else None


def _groups(wb):
    """(text, [numbers]) for every row and every column of every sheet."""
    out = []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        cols = [[] for _ in range(ncols)]
        for r in grid:
            text = " ".join(str(c.value).lower() for c in r if c.value is not None)
            nums = [n for n in (_num(c.value) for c in r) if n is not None]
            out.append((text, nums))
            for j, c in enumerate(r):
                cols[j].append(c.value)
        for vals in cols:
            text = " ".join(str(v).lower() for v in vals if v is not None)
            nums = [n for n in (_num(v) for v in vals) if n is not None]
            out.append((text, nums))
    return out


def _near(a, b, rel):
    return abs(a - b) <= max(abs(b) * rel, 0.005)


def check(ws, ref):
    name = "plate cost and margin, every dish"
    hits = sorted(glob.glob(os.path.join(ws, "menu_costs.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "menu_costs.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "menu_costs.xlsx not found"}]
    try:
        dishes = json.load(open(os.path.join(ref, "notes.json")))["dishes"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    groups = _groups(wb)
    bad = []
    for dish, t in sorted(dishes.items()):
        key = t["keyword"]
        forms = [t["margin"], t["margin_pct"], t["margin_pct"] * 100, 1 - t["margin_pct"], (1 - t["margin_pct"]) * 100]
        ok = False
        for text, nums in groups:
            if key not in text:
                continue
            if any(_near(n, t["cost"], 0.005) for n in nums) and any(_near(n, f, 0.01) for f in forms for n in nums):
                ok = True
                break
        if not ok:
            bad.append(f"{dish}: expected plate cost {t['cost']:.2f} with margin {t['margin']:.2f} ({t['margin_pct']:.1%})")
    return [{"name": name, "passed": not bad,
             "detail": "; ".join(bad) if bad else f"{len(dishes)} dishes carry the right plate cost and margin"}]
