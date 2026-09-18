"""chargeback-tracker: every case's evidence_due must be the reference date (any common date style accepted).

Rows are matched on the case id's letters and digits; truth is reference/cases.json.
"""
import csv
import glob
import io
import json
import os
import re
from datetime import date, datetime

FMTS = ["%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y", "%d %b %Y", "%d %B %Y",
        "%d-%b-%Y", "%Y/%m/%d", "%a %b %d %Y", "%a, %b %d, %Y", "%A, %B %d, %Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"]


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _norm(c):
    return re.sub(r"[^a-z0-9]+", "_", str(c or "").strip().lower()).strip("_")


def _rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        vals = [["" if v is None else str(v) for v in r] for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        vals = list(csv.reader(io.StringIO(text)))
    if not vals:
        return []
    hdr = [_norm(h) for h in vals[0]]
    return [dict(zip(hdr, row)) for row in vals[1:] if any(str(x).strip() for x in row)]


def parse_date(v):
    s = str(v or "").strip()
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s)
    for f in FMTS:
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", s)
    return date.fromisoformat(m.group(1)) if m else None


def check(ws, ref):
    name = "evidence deadlines"
    p = _find(ws, "chargebacks.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "chargebacks.csv not found"}]
    rows = _rows(p)
    if not rows or "case_id" not in rows[0] or "evidence_due" not in rows[0]:
        return [{"name": name, "passed": False, "detail": "need case_id and evidence_due columns"}]
    got = {}
    for r in rows:
        k = re.sub(r"[^a-z0-9]", "", r["case_id"].lower())
        got.setdefault(k, []).append(parse_date(r["evidence_due"]))
    bad = []
    truth = json.load(open(os.path.join(ref, "cases.json")))
    for t in truth:
        k = re.sub(r"[^a-z0-9]", "", t["case_id"].lower())
        vals = got.get(k) or []
        if not vals:
            bad.append(f"{t['case_id']}: missing")
        elif any(v is None or v.isoformat() != t["evidence_due"] for v in vals):
            bad.append(f"{t['case_id']}: {[v.isoformat() if v else None for v in vals]} != {t['evidence_due']}")
    return [{"name": name, "passed": not bad, "detail": "; ".join(bad[:6]) if bad else f"{len(truth)} deadlines right"}]
