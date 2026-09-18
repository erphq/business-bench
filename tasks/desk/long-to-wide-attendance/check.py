"""Rows in the attendance grid are in alphabetical order by Student Name (case-insensitive), as the district asks."""
import csv
import glob
import io
import os
import re


def _norm(c):
    return re.sub(r"[^a-z0-9]+", "_", str(c).strip().lower()).strip("_")


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _names(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        rows = [["" if v is None else str(v) for v in r] for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return None
    head = [_norm(h) for h in rows[0]]
    if "student_name" not in head:
        return None
    j = head.index("student_name")
    return [r[j].strip() for r in rows[1:] if len(r) > j and r[j].strip()]


def check(ws, ref):
    name = "rows sorted by student name"
    p = _find(ws, "attendance_grid.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "attendance_grid.csv not found"}]
    try:
        names = _names(p)
    except Exception as e:  # unreadable deliverable is a failed check, never a crash
        return [{"name": name, "passed": False, "detail": f"could not read: {e}"}]
    if not names or len(names) < 2:
        return [{"name": name, "passed": False, "detail": "no Student Name column or fewer than two rows"}]
    want = sorted(names, key=lambda s: s.lower())
    ok = [s.lower() for s in names] == [s.lower() for s in want]
    bad = next((i for i, (a, b) in enumerate(zip(names, want)) if a.lower() != b.lower()), None)
    return [{"name": name, "passed": ok, "detail": f"{len(names)} rows in order" if ok else
             f"row {bad + 2} is {names[bad]!r}, expected {want[bad]!r} in alphabetical order"}]
