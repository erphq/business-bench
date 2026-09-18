"""Every shared account splits across the five departments to the cent, with the rounding residual where the memo puts it.

The department totals are pinned separately; a residual placed on the wrong department can still land inside a total's
one-cent tolerance, so this check looks at each account's own split. It recalculates the workbook with the grader's
engine and passes when, for every account in reference/notes.json, some single row or column holds all five of that
account's department amounts (each within half a cent, each matched to a different cell). Accounts down the side or
across the top both pass.
"""
import glob
import json
import os
import sys


def _bench():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    sys.path.insert(0, os.path.join(root, "bench"))
    import grade  # noqa: E402
    return grade


def _lines(wb, cell_num):
    out = []
    for sh in wb.worksheets:
        grid = [[cell_num(c.value) for c in row] for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        out += [[v for v in r if v is not None] for r in grid]
        out += [[r[j] for r in grid if j < len(r) and r[j] is not None] for j in range(ncols)]
    return out


def _covers(line, want, tol):
    pool = list(line)
    for w in want:
        k = next((i for i, v in enumerate(pool) if abs(v - w) <= tol), None)
        if k is None:
            return False
        pool.pop(k)
    return True


def check(ws, ref):
    name = "every account splits to the cent"
    hits = sorted(glob.glob(os.path.join(ws, "allocation.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "allocation.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "allocation.xlsx not found"}]
    try:
        splits = json.load(open(os.path.join(ref, "notes.json")))["splits"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        grade = _bench()
        from openpyxl import load_workbook
        wb = load_workbook(grade.recalculated_workbook(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    lines = _lines(wb, grade.cell_num)
    bad = [acct for acct, pieces in splits.items() if not any(_covers(ln, [float(v) for v in pieces.values()], 0.005) for ln in lines)]
    return [{"name": name, "passed": not bad,
             "detail": f"{len(splits) - len(bad)}/{len(splits)} accounts split exactly" + (f"; not found: {bad[:5]}" if bad else "")}]
