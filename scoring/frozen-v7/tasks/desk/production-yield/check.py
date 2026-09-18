"""production-yield custom check: yield per line written as a fraction (0.9123) or a percentage (91.23 or '91.23%'),
on a row whose label cell names the line (or in a column headed by it). Batch-level rows - any row carrying a batch id
or a SKU - are skipped, because a single batch's yield can sit within tolerance of the line's."""
from __future__ import annotations
import glob
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "bench"))
LABEL = {"Bar line": "bar", "Granola line": "granola", "Bites line": "bites"}
BATCH_ROW = re.compile(r"\b(br|gr|bt)\d{4}-\d+\b|\b(bar|grn|bite)-[a-z]+-\d+\b")


def _find(ws: str, pattern: str):
    hits = sorted(glob.glob(os.path.join(ws, pattern))) or sorted(glob.glob(os.path.join(ws, "**", pattern), recursive=True))
    hits = [h for h in hits if os.path.isfile(h) and "/.proto" not in h and "/.codex" not in h]
    return hits[0] if hits else None


def _num(v):
    if isinstance(v, bool) or v is None:
        return None, False
    if isinstance(v, (int, float)):
        return float(v), False
    m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*%\s*", str(v))
    if m:
        return float(m.group(1)), True
    try:
        return float(str(v).replace(",", "")), False
    except ValueError:
        return None, False


def check(ws: str, ref: str) -> list[dict]:
    try:
        pins = json.load(open(os.path.join(ref, "notes.json")))["yield_pins"]
    except Exception as e:  # pragma: no cover
        return [{"name": "yield per line", "passed": False, "detail": f"reference unreadable: {e}"}]
    names = [f"yield {line}" for line, _ in pins]
    p = _find(ws, "yield.xlsx")
    if not p:
        return [{"name": n, "passed": False, "detail": "yield.xlsx missing"} for n in names]
    try:
        from grade import recalculated_workbook
        from openpyxl import load_workbook
        wb = load_workbook(recalculated_workbook(p), data_only=True)
    except Exception as e:
        return [{"name": n, "passed": False, "detail": f"workbook unreadable: {e}"} for n in names]
    grids = []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        cols: dict[int, str] = {}
        for row in grid:
            for j, c in enumerate(row):
                if c.value is not None:
                    cols[j] = cols.get(j, "") + " " + str(c.value).lower()
        grids.append((sh.title, grid, cols))
    out = []
    for (line, want), name in zip(pins, names):
        label = LABEL[line]
        hit = None
        for title, grid, cols in grids:
            for row in grid:
                vals = [str(c.value).lower() for c in row if c.value is not None]
                text = " ".join(vals)
                # a batch-level row (batch id or SKU on it) is not a line total, even if it names the line
                if BATCH_ROW.search(text):
                    continue
                labelled_row = bool(vals) and label in vals[0]
                for j, c in enumerate(row):
                    if not labelled_row and label not in cols.get(j, ""):
                        continue
                    v, pct = _num(c.value)
                    if v is None:
                        continue
                    if (not pct and abs(v - want) <= 0.0015) or abs(v - want * 100) <= 0.15:
                        hit = f"{title}!{c.coordinate} = {c.value}"
                        break
                if hit:
                    break
            if hit:
                break
        out.append({"name": name, "passed": bool(hit), "detail": hit or f"no cell = {want} (or {want * 100:.2f}%) near {label!r}"})
    return out
