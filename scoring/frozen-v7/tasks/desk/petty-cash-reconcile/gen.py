#!/usr/bin/env python3
"""petty-cash-reconcile: a pediatric clinic's $400 imprest petty cash box reconciled at month end, with the
replenishment check coded by account.

    python gen.py [--seed N] [--naive DIR]

Business: Oakhurst Pediatrics keeps a petty cash box at the front desk for stamps, office bits, the staff kitchen and
waiting-room supplies. Dana logs every voucher in a sheet, staples receipts into an envelope, and the box was topped
up to $400 on 14 August. At month end the office manager wants it reconciled and a replenishment request coded.

Traps (each caught by a check, see task.yaml):
  * a voucher in the log has no receipt: it goes to the missing-receipt suspense account, not an expense
                                                                                  (checks: kitchen total; flagged; replenishment)
  * one voucher was logged at 42.80 but its receipt says 24.80; the receipt is what was paid (checks: office total; short)
  * the 14 August replenishment resets the period; a receipt already claimed on it is still in the envelope
                                                                                  (checks: category totals; replenishment)
  * a receipt in the envelope was never logged; it is real spending              (checks: postage total; short)
  * the log's running balance has an addition slip, so its closing balance is wrong (check: short)
  * the count sheet gives denominations only                                      (checks: replenishment; short)
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "lib"))
from bizgen import *  # noqa: E402,F403

FUND = 400.00
REPLEN_DAY = date(2026, 8, 14)
ACCOUNTS = [("6610", "Postage & shipping", "postage"), ("6600", "Office supplies", "office"),
            ("6450", "Staff kitchen & refreshments", "kitchen"), ("6520", "Waiting room & patient supplies", "waiting")]
ACC = {k: (a, n) for a, n, k in ACCOUNTS}
# category: (log purposes, vendors with item lists)
SPEND = {  # category: vendors with (item, price, what the log calls it)
    "postage": [("USPS Oakhurst Station", [("Forever stamps (book of 20)", 14.60, "stamps"), ("Certified mail w/ return receipt", 9.35, "certified mail"),
                                           ("Priority Mail flat rate env", 10.40, "priority mail - lab kit")]),
                ("The UPS Store #4471", [("Ground shipping - 2 lb", 13.86, "shipping records"), ("Padded envelope", 2.49, "envelopes")])],
    "office": [("Staples #0442", [("HP 58A toner", 89.99, "printer toner"), ("Manila folders 100ct", 12.49, "file folders"),
                                  ("Label tape 1/2in", 17.99, "label tape"), ("Gel pens 12pk", 11.29, "pens")]),
               ("Office Depot #2231", [("Copy paper ream", 8.99, "copy paper"), ("Scotch tape 6pk", 9.49, "tape"), ("Binder clips", 4.79, "binder clips"),
                                       ("Sticky notes 12pk", 10.99, "sticky notes")])],
    "kitchen": [("Safeway #1882", [("Folgers 30oz", 13.99, "coffee for staff room"), ("Milk 1 gal", 4.29, "milk"), ("Paper towels 6pk", 11.49, "paper towels"),
                                   ("Dish soap", 3.99, "dish soap"), ("Bottled water 24pk", 5.99, "water for staff fridge")]),
                ("Costco Whse #0891", [("Kirkland coffee 3lb", 21.99, "coffee for staff room"), ("Water 40pk", 4.99, "water for staff fridge")])],
    "waiting": [("Target #1102", [("Sticker rolls 4pk", 9.99, "stickers for exam rooms"), ("Crayons 24ct", 1.99, "crayons"),
                                  ("Coloring books", 3.99, "coloring books"), ("Disinfecting wipes 3pk", 12.49, "toy bin wipes")]),
                ("Dollar Tree #3301", [("Picture books", 1.25, "books for waiting room"), ("Stickers", 1.25, "stickers"), ("Bubbles", 1.25, "bubbles")])],
}


def c(x):
    return round(x + 0.0, 2)


def make_receipt(r, cat, day):
    vendor, items = r.choice(SPEND[cat])
    lines, what = [], []
    for name, price, purpose in r.sample(items, r.randint(1, min(2, len(items)))):
        q = r.randint(1, 2) if price < 20 else 1
        lines.append((name, q, price))
        what.append(purpose)
    sub = c(sum(q * p for _, q, p in lines))
    tax = 0.0 if cat == "postage" else c(sub * 0.0825)
    return {"vendor": vendor, "date": day, "lines": lines, "subtotal": sub, "tax": tax, "total": c(sub + tax), "cat": cat,
            "purpose": " + ".join(what), "no": f"{r.randint(1000, 9999)}-{r.randint(10, 99)}"}


def build(seed: int) -> dict:
    r = rng(seed)
    cats = list(SPEND)
    # ---- before the 14 Aug replenishment (already claimed) ----
    before = []
    for i, day in enumerate(sorted(day_in(r, date(2026, 8, 3), date(2026, 8, 13), True) for _ in range(6))):
        rc = make_receipt(r, cats[i % 4], day)
        rc["voucher"] = f"PC-{220 + i}"
        rc["logged"] = rc["total"]
        before.append(rc)
    replen1 = c(sum(x["logged"] for x in before))
    # ---- since the replenishment ----
    after = []
    plan = ["postage", "office", "kitchen", "waiting", "office", "kitchen", "postage", "waiting", "office", "kitchen"]
    days = sorted(day_in(r, date(2026, 8, 17), date(2026, 8, 28), True) for _ in range(len(plan)))
    for i, (cat, day) in enumerate(zip(plan, days)):
        rc = make_receipt(r, cat, day)
        rc["voucher"] = f"PC-{226 + i}"
        rc["logged"] = rc["total"]
        rc["tag"] = ""
        after.append(rc)
    # the logged amount with two digits swapped (receipt is right)
    swap = next(x for x in after if x["cat"] == "office" and x["total"] >= 12 and int(x["total"]) // 10 % 10 != int(x["total"]) % 10)
    t = f"{swap['total']:.2f}"
    whole, frac = t.split(".")
    sw = whole[:-2] + whole[-1] + whole[-2] if len(whole) >= 2 else whole
    swap["logged"] = float(f"{sw}.{frac}")
    swap["tag"] = "log_error"
    # a voucher whose receipt never made it into the envelope
    missing = next(x for x in after if x["cat"] == "kitchen")
    missing["tag"] = "missing_receipt"
    # a receipt nobody logged
    unlogged = make_receipt(r, "postage", date(2026, 8, 27))
    while unlogged["total"] < 12:
        unlogged = make_receipt(r, "postage", date(2026, 8, 27))
    unlogged["voucher"] = ""
    unlogged["tag"] = "unlogged"
    # true cash left the box
    spent_after = c(sum(x["total"] for x in after) + unlogged["total"])
    short = c(r.choice([2.35, 3.15, 4.60, 5.05, 6.40, 7.75]))
    counted = c(FUND - spent_after - short)
    # log with an addition slip
    log = [{"date": date(2026, 8, 1), "voucher": "", "who": "", "what": "Balance brought forward", "out": None, "in": None}]
    for x in before:
        log.append({"date": x["date"], "voucher": x["voucher"], "who": x["vendor"].split(" #")[0], "what": x["purpose"], "out": x["logged"], "in": None})
    log.append({"date": REPLEN_DAY, "voucher": "", "who": "Replenishment - check 3318", "what": "top up to $400", "out": None, "in": replen1})
    for x in after:
        log.append({"date": x["date"], "voucher": x["voucher"], "who": x["vendor"].split(" #")[0], "what": x["purpose"], "out": x["logged"], "in": None})
    bal = FUND
    slip_at = r.randint(len(before) + 3, len(log) - 3)
    slip = float(r.choice([10, -10, 1, -1]))
    for i, row in enumerate(log):
        if i:
            bal = c(bal - (row["out"] or 0) + (row["in"] or 0))
            if i == slip_at:
                bal = c(bal + slip)
        row["balance"] = bal
    log_close = bal
    # count sheet denominations for the counted cash
    denoms = [(20.0, "Twenties"), (10.0, "Tens"), (5.0, "Fives"), (1.0, "Ones"), (0.25, "Quarters"), (0.10, "Dimes"), (0.05, "Nickels"), (0.01, "Pennies")]
    left = int(round(counted * 100))
    counts = []
    for v, name in denoms:
        cv = int(round(v * 100))
        n = left // cv
        if v >= 5 and n > 0:
            n = max(0, n - r.randint(0, 2))
        if v == 0.25:
            n = max(0, n - r.randint(0, 3))
        counts.append([name, v, n])
        left -= n * cv
    assert left == 0
    # category truths for the replenishment
    by_cat = {k: 0.0 for k in SPEND}
    for x in after:
        if x["tag"] != "missing_receipt":
            by_cat[x["cat"]] = c(by_cat[x["cat"]] + x["total"])
    by_cat["postage"] = c(by_cat["postage"] + unlogged["total"])
    replen = c(FUND - counted)
    assert abs(replen - c(sum(by_cat.values()) + missing["logged"] + short)) < 0.005
    claimed_again = before[1]
    return {"before": before, "after": after, "unlogged": unlogged, "missing": missing, "swap": swap, "claimed_again": claimed_again,
            "replen1": replen1, "short": short, "counted": counted, "counts": counts, "log": log, "log_close": log_close,
            "slip": slip, "by_cat": by_cat, "replen": replen}


def naive(d):
    """Log only: whole-month outflows by category at logged amounts, log closing balance as the cash."""
    cat = {k: 0.0 for k in SPEND}
    for x in d["before"] + d["after"]:
        cat[x["cat"]] = c(cat[x["cat"]] + x["logged"])
    return {"by_cat": cat, "replen": c(sum(x["logged"] for x in d["after"])), "short": c(d["counted"] - d["log_close"])}


def acceptable(d):
    if not (45 <= d["counted"] <= 170) or d["log_close"] <= 0:
        return False
    others = [x["logged"] for x in d["before"] + d["after"]] + [x["total"] for x in d["before"] + d["after"]]
    if any(abs(d["unlogged"]["total"] - q) < 1 for q in others):
        return False
    n = naive(d)
    for k in ("postage", "office", "kitchen"):
        if abs(n["by_cat"][k] - d["by_cat"][k]) < 1:
            return False
        # a since-14-August-only reading at logged amounts must also miss
        logged_after = c(sum(x["logged"] for x in d["after"] if x["cat"] == k))
        if abs(logged_after - d["by_cat"][k]) < 1:
            return False
    if abs(n["replen"] - d["replen"]) < 1:
        return False
    if abs(abs(n["short"]) - d["short"]) < 0.5:
        return False
    vals = [d["by_cat"][k] for k in SPEND] + [d["replen"], d["counted"]]
    if len({round(v, 2) for v in vals}) < len(vals):
        return False
    raw = {x["logged"] for x in d["before"] + d["after"]} | {x["total"] for x in d["before"] + d["after"] + [d["unlogged"]]} | \
          {row["balance"] for row in d["log"]} | {d["replen1"]}
    if any(abs(v - q) < 0.011 for v in [d["by_cat"]["postage"], d["by_cat"]["office"], d["by_cat"]["kitchen"], d["replen"]] for q in raw):
        return False
    if d["swap"]["logged"] - d["swap"]["total"] < 5 and d["swap"]["total"] - d["swap"]["logged"] < 5:
        return False
    return True


def workbook(d, path, naive_mode=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    rec = wb.active; rec.title = "Reconciliation"
    cnt = wb.create_sheet("Cash count")
    rcs = wb.create_sheet("Receipts since 14 Aug")
    cod = wb.create_sheet("Replenishment coding")
    cnt.append(["Denomination", "Value", "Count", "Amount"])
    for i, (name, v, n) in enumerate(d["counts"], start=2):
        cnt.append([name, v, n, f"=ROUND(B{i}*C{i},2)"])
    last = len(d["counts"]) + 1
    cnt.append(["Cash counted 31 Aug", None, None, f"=SUM(D2:D{last})"])
    count_cell = f"'Cash count'!D{last + 1}"

    rcs.append(["Voucher", "Date", "Paid to", "Account", "Logged", "Receipt total", "Note"])
    items = d["after"] + [d["unlogged"]]
    items = sorted(items, key=lambda x: (x["date"], x["voucher"]))
    for x in items:
        a, _ = ACC[x["cat"]]
        if naive_mode:
            if x["tag"] == "unlogged":
                continue
            rcs.append([x["voucher"], x["date"], x["vendor"], int(a), x["logged"], x["logged"], ""])
            continue
        if x["tag"] == "log_error":
            note = f"logged as {x['logged']:.2f}; receipt shows {x['total']:.2f}"
        elif x["tag"] == "missing_receipt":
            note = "NO RECEIPT - not expensed, held in 1360 until found"
        elif x["tag"] == "unlogged":
            note = "receipt in envelope, never logged"
        else:
            note = ""
        acct = 1360 if x["tag"] == "missing_receipt" else int(a)
        receipt_amt = None if x["tag"] == "missing_receipt" else x["total"]
        rcs.append([x["voucher"] or "(none)", x["date"], x["vendor"], acct, x["logged"] if x["tag"] != "unlogged" else None, receipt_amt, note])
    if not naive_mode:
        rcs.append([d["claimed_again"]["voucher"], d["claimed_again"]["date"], d["claimed_again"]["vendor"], None, None, None,
                    "already reimbursed on the 14 Aug replenishment - excluded"])
    n = rcs.max_row
    for row in rcs.iter_rows(min_row=2):
        if isinstance(row[1].value, date):
            row[1].number_format = "yyyy-mm-dd"

    cod.append(["Account", "Name", "Amount"])
    r_ = 2
    for a, name, k in ACCOUNTS:
        if naive_mode:
            cod.append([int(a), name, f"=SUMIF('Receipts since 14 Aug'!D2:D{n},{a},'Receipts since 14 Aug'!F2:F{n})"])
        else:
            cod.append([int(a), name, f"=SUMIF('Receipts since 14 Aug'!D2:D{n},{a},'Receipts since 14 Aug'!F2:F{n})"])
        r_ += 1
    if not naive_mode:
        cod.append([1360, "Missing receipts (suspense)", f"=SUMIF('Receipts since 14 Aug'!D2:D{n},1360,'Receipts since 14 Aug'!E2:E{n})"])
        cod.append([6990, "Cash over/short", "=-Reconciliation!B8"])
        r_ += 2
    cod.append([None, "Replenishment check total", f"=SUM(C2:C{r_ - 1})"])

    rows = [
        ["Oakhurst Pediatrics - petty cash reconciliation, 31 August 2026"],
        ["Imprest fund", FUND],
        ["Cash counted 31 Aug", f"={count_cell}"],
        ["Receipts on hand since the 14 Aug replenishment", f"=SUM('Receipts since 14 Aug'!F2:F{n})"],
        ["Vouchers with no receipt", f"=SUMIF('Receipts since 14 Aug'!D2:D{n},1360,'Receipts since 14 Aug'!E2:E{n})"],
        ["Accounted for", "=B3+B4+B5"],
        [],
        ["Cash over / (short)", "=ROUND(B6-B2,2)"],
        ["Replenishment needed to restore the fund", "=ROUND(B2-B3,2)"],
    ]
    if naive_mode:
        rows[2] = ["Cash per log 31 Aug", d["log_close"]]
        rows[8] = ["Replenishment needed to restore the fund", f"=SUM('Receipts since 14 Aug'!E2:E{n})"]
    for rr in rows:
        rec.append(rr)
    rec["A1"].font = Font(bold=True)
    for sh, widths in ((rec, {"A": 48, "B": 14}), (cnt, {"A": 22}), (rcs, {"A": 10, "B": 12, "C": 24, "D": 9, "E": 10, "F": 13, "G": 50}), (cod, {"B": 32})):
        for k, w in widths.items():
            sh.column_dimensions[k].width = w
    from datetime import datetime as _dt
    wb.properties.created = _dt(2026, 1, 15, 9, 0, 0); wb.properties.modified = _dt(2026, 1, 15, 9, 0, 0)
    wb.properties.creator = "reference"; wb.properties.lastModifiedBy = "reference"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    freeze_zip(path)


def receipts_pdf(d, path, r):
    blocks = [("title", "Petty cash receipts - August 2026"), ("small", "Envelope contents, scanned 31 Aug by front desk"), ("spacer", 6)]
    items = d["after"] + [d["unlogged"], d["claimed_again"]]
    items = [x for x in items if x is not d["missing"]]
    items = sorted(items, key=lambda x: x["date"])
    for x in items:
        blocks.append(("hr", None))
        blocks.append(("h", x["vendor"]))
        kv = [("Date", x["date"].strftime("%m/%d/%Y")), ("Receipt", x["no"])]
        blocks.append(("kv", kv))
        rows = [["Item", "Qty", "Price", "Amount"]] + [[nm, str(q), f"{p:.2f}", f"{q * p:.2f}"] for nm, q, p in x["lines"]]
        blocks.append(("table", rows, {"col_widths": [230, 40, 60, 70]}))
        tl = [("Subtotal", f"{x['subtotal']:.2f}")]
        if x["tax"]:
            tl.append(("Sales tax", f"{x['tax']:.2f}"))
        tl.append(("TOTAL", f"${x['total']:.2f}"))
        blocks.append(("kv", tl))
        if x is d["claimed_again"]:
            blocks.append(("small", f"handwritten: {x['voucher']} - claimed on 8/14 top-up"))
        elif x["voucher"]:
            blocks.append(("small", f"handwritten: {x['voucher']}"))
    write_pdf_document(path, blocks, font="Helvetica", base_size=10)


def emit(seed, naive_dir):
    d = build(seed)
    if naive_dir:
        workbook(d, os.path.join(naive_dir, "petty_cash.xlsx"), naive_mode=True)
        return
    ws, ref, sol = task_dirs(HERE)
    r = rng(seed + 9)
    lrows = []
    for row in d["log"]:
        lrows.append([row["date"], row["voucher"] or None, row["who"] or None, row["what"], row["out"], row["in"], row["balance"]])
    write_xlsx(os.path.join(ws, "petty_cash_log_2026-08.xlsx"), {"August": {
        "merged_title": "Front desk petty cash log - August 2026 (fund $400)",
        "header": ["Date", "Voucher", "Paid to / from", "What for", "Cash out", "Cash in", "Balance"],
        "rows": lrows, "widths": {"A": 12, "B": 9, "C": 26, "D": 32, "E": 10, "F": 10, "G": 10},
        "number_formats": {"E": "0.00", "F": "0.00", "G": "0.00"}}}, creator="Front desk")
    receipts_pdf(d, os.path.join(ws, "petty_cash_receipts_2026-08.pdf"), r)
    count_lines = ["PETTY CASH COUNT SHEET", "", "Date: 31 Aug 2026 4:45pm", "Counted by: Dana Tran    Witness: Marisol Ortiz", ""]
    for name, v, n in d["counts"]:
        count_lines.append(f"{name:<10} x {n}")
    count_lines += ["", "Total: ________", "", "Box also holds: 2 blank voucher pads, key to supply cabinet"]
    write_text(os.path.join(ws, "cash_count_2026-08-31.txt"), "\n".join(count_lines) + "\n")
    write_email_thread(os.path.join(ws, "email_from_marisol.txt"), [
        {"from": "Marisol Ortiz <mortiz@oakhurstpeds.org>", "to": "you", "date": "Mon, 31 Aug 2026 17:02",
         "subject": "petty cash for August",
         "body": ("Dana counted the box this afternoon and I witnessed it. Can you reconcile petty cash for me and work out the "
                  "replenishment check, coded by account so I can hand it to accounting? Put it in one workbook with the "
                  "totals as formulas.\n\n"
                  "The fund is $400. We topped it back up to $400 on 14 August, so this covers what has happened since then; the "
                  "receipts from before that went in with the first check.\n\n"
                  "Accounts: postage and shipping 6610, office supplies 6600, staff kitchen and refreshments 6450, waiting room "
                  "and patient supplies 6520. The receipt is what counts - if the log and a receipt disagree, the receipt is right. "
                  "If there is no receipt for a voucher, don't expense it: code it to 1360 missing receipts until Dana finds it. "
                  "Whatever is left over or short after all that goes to 6990 cash over/short.")}])

    write_json(os.path.join(ref, "petty.json"), {
        "fund": FUND, "counted": d["counted"], "replenishment": d["replen"], "short": d["short"], "by_account": {ACC[k][0]: v for k, v in d["by_cat"].items()},
        "missing_receipt": {"voucher": d["missing"]["voucher"], "amount": d["missing"]["logged"]},
        "log_error": {"voucher": d["swap"]["voucher"], "logged": d["swap"]["logged"], "receipt": d["swap"]["total"]},
        "unlogged_receipt": d["unlogged"]["total"], "claimed_again": d["claimed_again"]["voucher"], "log_closing_balance": d["log_close"],
        "naive": naive(d)})
    workbook(d, os.path.join(sol, "petty_cash.xlsx"))
    tol = lambda x: round(0.01 / max(abs(x), 1), 9)
    bc = d["by_cat"]
    write_task_yaml(HERE, {
        "id": "petty-cash-reconcile", "track": "desk", "category": "bookkeeping",
        "title": "Reconcile the front desk petty cash and code the top-up",
        "ask": ("Can you reconcile the front desk petty cash for August and work out the replenishment? Dana's log, the receipts, "
                "the count sheet and Marisol's email are in the folder. Save it as petty_cash.xlsx.\n"),
        "followup": None, "timeout_s": 1200,
        "traps": [
            f"voucher {d['missing']['voucher']} ({d['missing']['what'] if 'what' in d['missing'] else d['missing']['purpose']}, "
            f"{d['missing']['logged']:.2f}) is in the log but its receipt is not in the envelope; it goes to 1360 missing receipts, "
            "not to the kitchen account, and the workbook must flag it (checks: kitchen & refreshments total; "
            "over/short and missing receipt flagged)",
            f"voucher {d['swap']['voucher']} was logged at {d['swap']['logged']:.2f} but the receipt says {d['swap']['total']:.2f}; "
            "the email says the receipt is right (checks: office supplies total; over/short and missing receipt flagged)",
            "the log shows the 14 August top-up as cash in, which starts a new period; totalling the whole month's vouchers codes "
            f"spending already reimbursed, and the receipt for {d['claimed_again']['voucher']} from before the top-up is still in "
            "the envelope (checks: postage total; office supplies total; kitchen & refreshments total; replenishment check)",
            f"the envelope holds a receipt from {d['unlogged']['vendor']} for {d['unlogged']['total']:.2f} on 27 August that was never logged; the cash still "
            "left the box and it is postage (checks: postage total; over/short and missing receipt flagged)",
            f"the log's running balance slips by {abs(d['slip']):.2f} on one line, so its closing balance of {d['log_close']:.2f} is not what "
            "the box should hold; the shortage has to come from the count and the receipts (check: over/short and missing receipt flagged)",
            "the count sheet lists denominations with the total left blank; the replenishment is the fund less the cash "
            "actually counted, not the sum of the logged vouchers (check: replenishment check)",
        ],
        "checks": [
            {"type": "file_exists", "name": "petty_cash.xlsx exists", "path": "petty_cash.xlsx"},
            {"type": "xlsx_has_formulas", "name": "live formulas", "path": "petty_cash.xlsx", "min_count": 4},
            {"type": "xlsx_no_errors", "name": "no error cells", "path": "petty_cash.xlsx"},
            {"type": "xlsx_value_present", "name": "replenishment check", "path": "petty_cash.xlsx", "expected": d["replen"],
             "rel_tol": tol(d["replen"]), "near_text": "replenish"},
            {"type": "xlsx_value_present", "name": "postage total", "path": "petty_cash.xlsx", "expected": bc["postage"],
             "rel_tol": tol(bc["postage"]), "near_text": "postage"},
            {"type": "xlsx_value_present", "name": "office supplies total", "path": "petty_cash.xlsx", "expected": bc["office"],
             "rel_tol": tol(bc["office"]), "near_text": "office"},
            {"type": "xlsx_value_present", "name": "kitchen & refreshments total", "path": "petty_cash.xlsx", "expected": bc["kitchen"],
             "rel_tol": tol(bc["kitchen"]), "near_text": "kitchen"},
            {"type": "custom", "name": "over/short and missing receipt flagged", "module": "check.py"},
        ],
    })
    print(f"seed={seed}: counted {d['counted']} replen {d['replen']} short {d['short']} by_cat {d['by_cat']} log close {d['log_close']}")
    print("  naive:", naive(d))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--naive", default=None)
    a = ap.parse_args()
    for attempt in range(500):
        s = a.seed * 1000 + attempt
        try:
            ok = acceptable(build(s))
        except (StopIteration, AssertionError):
            ok = False
        if ok:
            break
    else:
        raise SystemExit("no acceptable draw")
    emit(s, a.naive)
