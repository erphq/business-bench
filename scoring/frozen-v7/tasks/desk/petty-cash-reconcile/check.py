"""petty-cash-reconcile custom check, layout-agnostic.

1. the cash shortage: a cell on a row mentioning short / over / shortage / discrepancy / unexplained / variance equals the
   true shortage, written either as a negative or as a positive "short" figure.
2. the voucher with no receipt is flagged: its voucher number sits on a row with a flag word (no receipt, missing, 1360,
   suspense, not found ...).
3. the missing-receipt amount is coded to 1360: a cell equal to that amount on a row mentioning 1360 or missing/suspense.
Figures come from reference/petty.json.
"""
import glob
import json
import os
import re

FLAG = re.compile(r"(no receipt|missing|\b1360\b|suspense|not found|unsupported|without (a )?receipt|no documentation|lost)", re.I)
SHORT = re.compile(r"(short|over|discrepanc|unexplained|variance|difference)", re.I)


def _grade_module():
    import importlib.util
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def check(ws, ref):
    hits = sorted(glob.glob(os.path.join(ws, "petty_cash.xlsx"))) or sorted(glob.glob(os.path.join(ws, "**", "petty_cash.xlsx"), recursive=True))
    if not hits:
        return [{"name": "petty_cash.xlsx present", "passed": False, "detail": "not found"}]
    pj = json.load(open(os.path.join(ref, "petty.json")))
    try:
        g = _grade_module()
        from openpyxl import load_workbook
        wb = load_workbook(g.recalculated_workbook(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": "workbook readable", "passed": False, "detail": f"{type(e).__name__}: {e}"}]
    rows = []
    for sh in wb.worksheets:
        for row in sh.iter_rows(values_only=True):
            vals = [v for v in row if v is not None and str(v).strip() != ""]
            if vals:
                texts = [str(v) for v in vals if isinstance(v, str)]
                nums = [x for x in (g.cell_num(v) for v in vals) if x is not None]
                rows.append((" ".join(str(v) for v in vals), texts, nums))
    out = []
    short = pj["short"]
    ok = any(any(SHORT.search(t) for t in texts) and any(abs(abs(x) - short) <= 0.011 for x in nums) for _, texts, nums in rows)
    out.append({"name": "cash shortage", "passed": ok, "detail": f"{short:.2f} {'found' if ok else 'not found'} on an over/short row"})
    v = pj["missing_receipt"]["voucher"]
    pat = re.compile(rf"\b{re.escape(v)}\b", re.I)
    ok2 = any(pat.search(full) and FLAG.search(full.replace(v, "")) for full, _, _ in rows)
    out.append({"name": f"{v} flagged as missing its receipt", "passed": ok2, "detail": "flagged" if ok2 else f"no row with {v} and a missing-receipt flag"})
    amt = pj["missing_receipt"]["amount"]
    ok3 = any(re.search(r"(\b1360\b|missing|suspense)", full, re.I) and any(abs(x - amt) <= 0.011 for x in nums) for full, _, nums in rows)
    out.append({"name": "missing receipt amount held in 1360", "passed": ok3, "detail": f"{amt:.2f} {'found' if ok3 else 'not found'} on a 1360/missing row"})
    return out
