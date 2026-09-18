"""payroll-journal-entry: the net amount (debit minus credit) on every account matches the reference entry.

Account numbers are read as the first four-digit number in account_number, so "2115" and "2115 FICA payable" are the
same account; several lines on one account are added together; a blank debit or credit is 0. Figures come from
reference/je.json.
"""
import csv
import glob
import io
import json
import os
import re


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _norm(h):
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def read_lines(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        vals = [["" if v is None else str(v) for v in row] for row in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        vals = list(csv.reader(io.StringIO(text)))
    if not vals:
        return None, "empty file"
    hdr = [_norm(h) for h in vals[0]]
    need = ["account_number", "debit", "credit"]
    if any(n not in hdr for n in need):
        return None, f"need columns {need}; have {hdr}"
    ia, idr, icr = (hdr.index(n) for n in need)
    lines = []
    for row in vals[1:]:
        if not any(str(x).strip() for x in row):
            continue
        acct = re.search(r"\d{4}", str(row[ia]) if ia < len(row) else "")
        lines.append((acct.group(0) if acct else str(row[ia]).strip(), money(row[idr] if idr < len(row) else ""), money(row[icr] if icr < len(row) else "")))
    return lines, ""


def money(v):
    s = str(v).strip()
    if s in ("", "-"):
        return 0.0
    neg = s.startswith("(") or s.startswith("-")
    try:
        x = float(re.sub(r"[^0-9.]", "", s))
    except ValueError:
        return float("nan")
    return -x if neg else x


def check(ws, ref):
    name = "account amounts"
    p = _find(ws, "payroll_je.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "payroll_je.csv not found"}]
    lines, err = read_lines(p)
    if lines is None:
        return [{"name": name, "passed": False, "detail": err}]
    want = json.load(open(os.path.join(ref, "je.json")))["net_by_account"]
    got = {}
    for acct, dr, cr in lines:
        got[acct] = got.get(acct, 0.0) + dr - cr
    wrong = [f"{a}: {got.get(a, 0.0):,.2f} (want {v:,.2f})" for a, v in sorted(want.items()) if not abs(got.get(a, 0.0) - v) <= 0.011]
    extra = [f"{a}: {v:,.2f}" for a, v in sorted(got.items()) if a not in want and abs(v) > 0.011]
    ok = not wrong and not extra
    return [{"name": name, "passed": ok, "detail": f"{len(want) - len(wrong)}/{len(want)} accounts right" + (f"; wrong: {wrong}" if wrong else "")
             + (f"; unexpected accounts: {extra}" if extra else "")}]
