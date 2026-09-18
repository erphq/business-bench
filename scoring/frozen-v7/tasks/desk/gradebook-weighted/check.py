"""Every enrolled student's final percentage and letter, on one row of grades.xlsx; the withdrawn student gets no letter.

A percentage is legitimately written 87.4, 0.874 or "87.4%", so this recalculates the workbook with the grader's
own engine and accepts either scale within 0.1 percentage points. The row is found by the student id or by both
first and last name, and the same row must carry the expected letter as its own cell ("B+"). Expected values come
from reference/notes.json, so nothing here is pinned to a seed.
"""
import glob
import importlib.util
import json
import os
import re

LETTERS = {"A", "A-", "B+", "B", "B-", "C+", "C", "C-", "D", "F", "D+", "D-"}


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s?%?\s*", str(v))
    return float(m.group(1)) if m else None


def check(ws, ref):
    name = "final grade per student"
    hits = sorted(glob.glob(os.path.join(ws, "grades.xlsx"))) or sorted(glob.glob(os.path.join(ws, "**", "grades.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "grades.xlsx not found"}]
    try:
        notes = json.load(open(os.path.join(ref, "notes.json")))
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    rows = []
    for sh in wb.worksheets:
        for row in sh.iter_rows(values_only=True):
            vals = [v for v in row if v is not None]
            if not vals:
                continue
            text = " ".join(str(v).lower() for v in vals)
            nums = [x for x in (_num(v) for v in vals) if x is not None]
            letters = {str(v).strip().upper() for v in vals if isinstance(v, str) and str(v).strip().upper() in LETTERS}
            rows.append((text, nums, letters))

    def mentions(text, sid, first, last):
        return sid in text or (re.search(rf"\b{re.escape(first.lower())}\b", text) and re.search(rf"\b{re.escape(last.lower())}\b", text))

    bad = []
    for sid, s in sorted(notes["students"].items()):
        want, letter = s["final"], s["letter"]
        ok = False
        seen = False
        for text, nums, letters in rows:
            if not mentions(text, sid, s["first"], s["last"]):
                continue
            seen = True
            if letter in letters and any(abs(x - want) <= 0.1 or abs(x * 100 - want) <= 0.1 for x in nums):
                ok = True
                break
        if not ok:
            bad.append(f"{s['first']} {s['last']}: expected {want:.1f} {letter}" + ("" if seen else " (no row)"))
    w = notes["withdrawn"]
    for text, nums, letters in rows:
        if mentions(text, w["id"], w["first"], w["last"]) and letters:
            bad.append(f"withdrawn {w['first']} {w['last']} has a letter {sorted(letters)}")
            break
    return [{"name": name, "passed": not bad,
             "detail": "; ".join(bad[:6]) + (f" (+{len(bad) - 6} more)" if len(bad) > 6 else "") if bad
             else f"{len(notes['students'])} students match; withdrawn student has no letter"}]
