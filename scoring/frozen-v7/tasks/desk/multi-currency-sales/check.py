"""April flagged, no borrowed rate.

The April GBP average is missing from the rate sheet and the owner's note forbids borrowing another month's
rate. After recalculation the workbook must have a row mentioning April (2026-04, Apr, April) that also carries
a flag word (missing, no rate, pending, not converted, ...), and no numeric cell anywhere may equal April's
pounds converted at March's or May's rate (reference/notes.json).
"""
import glob
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "bench"))

APRIL = re.compile(r"(2026-04|\b04/2026\b|\bapr\b|\bapril\b|2026-4\b)", re.I)
FLAG = re.compile(r"(missing|no rate|not available|unavailable|\bn/?a\b|pending|incomplete|awaiting|\btbd\b|"
                  r"not converted|unconverted|flag|to follow|chasing|not yet|blank|excluded|left out|no (gbp|pound)|"
                  r"fill (it )?in|needs? (a |the )?rate)", re.I)


def check(ws, ref):
    name = "April flagged, no borrowed rate"
    hits = sorted(glob.glob(os.path.join(ws, "sales_usd.xlsx"))) or sorted(glob.glob(os.path.join(ws, "**", "sales_usd.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "sales_usd.xlsx not found"}]
    try:
        from grade import recalculated_workbook, cell_num
        from openpyxl import load_workbook
        wb = load_workbook(recalculated_workbook(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    notes = json.load(open(os.path.join(ref, "notes.json")))
    borrowed = [float(x) for x in notes["borrowed_rate_figures"]]
    flagged, bad = None, []
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            texts = [str(c.value) for c in row if c.value is not None]
            joined = " | ".join(texts)
            if not flagged and APRIL.search(joined) and any(FLAG.search(t) for t in texts if not re.fullmatch(r"[\d.\-\s:]+", t)):
                flagged = f"{sh.title}!row {row[0].row}: {joined[:120]}"
            for c in row:
                v = cell_num(c.value)
                if v is None:
                    continue
                for b in borrowed:
                    if abs(v - b) <= 0.02:
                        bad.append(f"{sh.title}!{c.coordinate}={v}")
    ok = bool(flagged) and not bad
    detail = (f"flag: {flagged}" if flagged else "no row mentioning April carries a flag") + \
             (f"; borrowed-rate figures present: {bad[:4]}" if bad else "")
    return [{"name": name, "passed": ok, "detail": detail}]
