"""Every price in pos_menu.csv is a plain number (no $, no +, no words) equal to the reference within half a cent,
and rows the reference leaves blank (modifier groups) stay blank. Keyed on Code."""
import csv
import glob
import io
import os
import re

NUM = re.compile(r"^-?\d+(\.\d{1,2})?$")


def _norm(c):
    return re.sub(r"[^a-z0-9]+", "_", str(c).strip().lower()).strip("_")


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        grid = [["" if v is None else str(v) for v in r] for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        grid = list(csv.reader(io.StringIO(text)))
    if not grid:
        return []
    head = [_norm(h) for h in grid[0]]
    return [dict(zip(head, r)) for r in grid[1:]]


def check(ws, ref):
    name = "prices as plain numbers"
    p = _find(ws, "pos_menu.csv")
    if not p:
        return [{"name": name, "passed": False, "detail": "pos_menu.csv not found"}]
    try:
        got = {str(r.get("code", "")).strip().lower(): r for r in _rows(p)}
    except Exception as e:  # an unreadable deliverable fails the check, never crashes the grader
        return [{"name": name, "passed": False, "detail": f"could not read: {e}"}]
    want = _rows(os.path.join(ref, "pos_menu.csv"))
    bad = []
    for w in want:
        k = w["code"].strip().lower()
        g = got.get(k)
        if g is None:
            bad.append(f"{w['code']}: missing"); continue
        gv, rv = str(g.get("price", "")).strip(), str(w["price"]).strip()
        if rv == "":
            if gv != "":
                bad.append(f"{w['code']}: {gv!r} should be blank")
            continue
        if not NUM.match(gv) or abs(float(gv) - float(rv)) > 0.005:
            bad.append(f"{w['code']}: {gv!r} != {rv}")
    return [{"name": name, "passed": not bad, "detail": "; ".join(bad[:8]) if bad else f"{len(want)} prices match"}]
