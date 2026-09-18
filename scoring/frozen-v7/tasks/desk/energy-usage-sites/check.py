"""energy-usage-sites custom check: Depot Lane Storage's figures, with June not invented.

Depot Lane has no meter read after 31 May, so its June cannot be worked out. Omar's email: "If there is a month you
cannot work out, leave it blank and say so in the memo. I would rather have a hole in the table than a number somebody
invented." The ask wants electricity by building month by month and never asks for building totals, so Depot's
half-year total may be the five-month sum (23,572.9, the reference) or left blank or marked incomplete because June is
a hole. Either reading passes:

  A. a cell equal to Depot's five-month total on a row or column naming Depot Lane (the original pin, unchanged); or
  B. Depot's January to May figures each on a row or column naming Depot Lane, and no kWh-sized number where Depot's
     June sits.

Where Depot's June sits, for every cell holding one of Depot's monthly figures: the cell on the same row under a column
labelled June, and the cell in the same column on a row labelled June when the column's label names Depot Lane (months
down the side) or the row's label does (one row per site and month) without being a portfolio row ("All sites excl.
Depot Lane"). A label is the text before a row's (or above a column's) first number. A number there of at least a quarter of Depot's smallest
month is an invented June, or a total that counts one (a "Jan-Jun" column); a zero, a note or a blank is the hole Omar
asked for.
"""
from __future__ import annotations
import csv
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "bench"))

SITE, NEAR = "Depot Lane Storage", "depot lane"
JUNE = re.compile(r"\bjune?\b|\b2026-06\b|\b0?6/2026\b|\bjun[- ]?(20)?26\b", re.I)
PORTFOLIO = re.compile(r"\btotal|\ball\b|portfolio|\bexcl|\bexcept\b|\bwithout\b", re.I)
REL = 0.001


def check(ws: str, ref: str) -> list[dict]:
    name = "Depot Lane figures (five-month total or a blank total, June not invented)"
    try:
        from grade import cell_num, find_file, recalculated_workbook
        from openpyxl import load_workbook
    except Exception as e:  # pragma: no cover
        return [{"name": name, "passed": False, "detail": f"grader helpers unavailable: {e}"}]
    months = {r["month"]: r["kwh"] for r in csv.DictReader(open(os.path.join(ref, "kwh_by_site_month.csv"))) if r["site"] == SITE}
    known = [float(v) for m, v in sorted(months.items()) if str(v).strip()]
    total = next(float(r["kwh"]) for r in csv.DictReader(open(os.path.join(ref, "site_totals.csv"))) if r["site"] == SITE)
    floor = min(known) / 4
    p = find_file(ws, "energy.xlsx")
    if not p:
        return [{"name": name, "passed": False, "detail": "energy.xlsx missing"}]
    try:
        wb = load_workbook(recalculated_workbook(p), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"deliverable unreadable by a standard reader: {e}"}]

    close = lambda v, e: v is not None and abs(v - e) <= max(abs(e) * REL, 0.01)
    total_hit, month_hits, invented = None, {e: None for e in known}, []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        val = lambda i, j: grid[i][j].value if j < len(grid[i]) else None
        num = lambda i, j: None if val(i, j) is None or isinstance(val(i, j), bool) else cell_num(val(i, j))
        row_text = [" ".join(str(c.value).lower() for c in r if c.value is not None) for r in grid]
        col_text = [" ".join(str(val(i, j)).lower() for i in range(len(grid)) if val(i, j) is not None) for j in range(ncols)]

        def row_label(i):
            out = []
            for j in range(len(grid[i])):
                if num(i, j) is not None:
                    break
                if val(i, j) is not None:
                    out.append(str(val(i, j)))
            return " ".join(out).lower()

        def col_label(j):
            out = []
            for i in range(len(grid)):
                if num(i, j) is not None:
                    break
                if val(i, j) is not None:
                    out.append(str(val(i, j)))
            return " ".join(out).lower()

        rlab = [row_label(i) for i in range(len(grid))]
        clab = [col_label(j) for j in range(ncols)]
        for i in range(len(grid)):
            for j in range(len(grid[i])):
                v = num(i, j)
                if v is None or not (NEAR in row_text[i] or NEAR in col_text[j]):
                    continue
                where = f"{sh.title}!{grid[i][j].coordinate}"
                if total_hit is None and close(v, total):
                    total_hit = f"{where} = {v}"
                anchors = [e for e in known if close(v, e)]
                for e in anchors:
                    month_hits[e] = month_hits[e] or where
                if not anchors:
                    continue
                # Depot's June beside this monthly figure: same row under a June column ...
                spots = [(i, jj) for jj in range(ncols) if jj != j and JUNE.search(clab[jj])]
                # ... or same column on a June row for Depot (months down the side, or one row per site and month)
                spots += [(ii, j) for ii in range(len(grid)) if ii != i and JUNE.search(rlab[ii])
                          and (NEAR in clab[j] or (NEAR in rlab[ii] and not PORTFOLIO.search(rlab[ii])))]
                for ii, jj in spots:
                    x = num(ii, jj)
                    is_date_serial = x is not None and float(x).is_integer() and 36526 <= x <= 55153
                    if x is not None and abs(x) >= floor and not is_date_serial:
                        invented.append(f"{sh.title}!{grid[ii][jj].coordinate} = {x}")

    if total_hit:
        return [{"name": name, "passed": True, "detail": f"five-month total {total} at {total_hit}"}]
    missing = [e for e, hit in month_hits.items() if not hit]
    ok = not missing and not invented
    if ok:
        detail = f"no five-month total; Depot Jan-May {known} all present and June left without a number"
    else:
        detail = (f"no cell = {total} near {NEAR!r}; " + (f"Depot monthly figures missing {missing}; " if missing else "")
                  + (f"a number where Depot's June sits: {sorted(set(invented))[:4]}" if invented else "")).rstrip("; ")
    return [{"name": name, "passed": ok, "detail": detail}]
