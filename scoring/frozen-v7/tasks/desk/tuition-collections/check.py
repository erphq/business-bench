"""tuition-collections custom check: the total still owed across programs.

The ask wants, program by program, "what has actually come in and what is still owed"; the bursar's note says credits
come off what the family owes. The total still owed (billed, less credits, less money that actually landed by the
June 5 cut-off, over all programs) must sit in a numeric cell after recalculation, on a row or column whose text says
what it holds - outstanding, owed, owing, balance, due, remaining, unpaid, receivable or uncollected - under whatever
header the author chose. The figure comes from the ALL row of reference/collections_by_program.csv.
"""
from __future__ import annotations

import csv
import glob
import os
import re
import sys

# grade.py runs from bench/, which puts it on sys.path; the repo layout (tasks/desk/<id>/) is the fallback
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "bench"))

OWED = re.compile(r"outstanding|\bowed\b|\bowing\b|\bowes\b|balance|\bdue\b|remaining|unpaid|receivable|uncollected|still to (?:come|collect)|left to collect", re.I)
REL_TOL = 1e-06
UNCAPPED_TOTAL_OWED = 75176.68   # the same total with the three over-invoice scholarships left uncapped


def _find(ws: str, pattern: str):
    hits = sorted(glob.glob(os.path.join(ws, pattern))) or sorted(glob.glob(os.path.join(ws, "**", pattern), recursive=True))
    hits = [h for h in hits if os.path.isfile(h) and "/.proto" not in h and "/.codex" not in h]
    return hits[0] if hits else None


def check(ws: str, ref: str) -> list[dict]:
    name = "total still owed"
    rows = list(csv.DictReader(open(os.path.join(ref, "collections_by_program.csv"), encoding="utf-8")))
    want = float(next(r for r in rows if r["program"].strip().upper() == "ALL")["outstanding"])
    # Three Afterschool scholarships exceed their invoice. The reference caps each credit at the invoice; nothing in the
    # ask or the bursar's note says a credit cannot take a balance below zero, so the uncapped total also counts.
    wants = [want, UNCAPPED_TOTAL_OWED]
    p = _find(ws, "collections.xlsx")
    if not p:
        return [{"name": name, "passed": False, "detail": "output file missing"}]
    from openpyxl import load_workbook
    from grade import cell_num, recalculated_workbook
    wb = load_workbook(recalculated_workbook(p), data_only=True)
    tol = max(abs(want) * REL_TOL, 0.01)
    for sh in wb.worksheets:
        grid = [list(r) for r in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        col_text = ["" for _ in range(ncols)]
        for r in grid:
            for j, c in enumerate(r):
                if c.value is not None:
                    col_text[j] += " " + str(c.value)
        for r in grid:
            row_text = " ".join(str(c.value) for c in r if c.value is not None)
            for j, c in enumerate(r):
                if c.value is None or isinstance(c.value, bool):
                    continue
                v = cell_num(c.value)
                if v is None or not any(abs(v - w) <= max(abs(w) * REL_TOL, 0.01) for w in wants):
                    continue
                if OWED.search(row_text) or OWED.search(col_text[j]):
                    return [{"name": name, "passed": True, "detail": f"{sh.title}!{c.coordinate} = {v}"}]
    return [{"name": name, "passed": False,
             "detail": f"no cell = {want:,.2f} (or {UNCAPPED_TOTAL_OWED:,.2f} uncapped) on a row or column saying what is still owed (outstanding, owed, balance, due, ...)"}]
