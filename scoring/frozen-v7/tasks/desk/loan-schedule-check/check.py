"""Interest per payment, one line.

The note's interest for each of the eight 2026 payments must sit on one row or one column of the workbook after
recalculation (each within 0.01, each matched to a different cell), so payments-down-the-side and payments-across
layouts both pass. Figures come from reference/notes.json.
"""
import glob
import json
import os
import sys


def _bench():
    root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    sys.path.insert(0, os.path.join(root, "bench"))
    import grade  # noqa: E402
    return grade


def _lines(wb, cell_num):
    for sh in wb.worksheets:
        grid = [[cell_num(c.value) for c in row] for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        for r in grid:
            yield [v for v in r if v is not None]
        for j in range(ncols):
            yield [r[j] for r in grid if j < len(r) and r[j] is not None]


def _covers(line, want, tol):
    pool = list(line)
    for w in want:
        k = next((i for i, v in enumerate(pool) if abs(v - w) <= tol), None)
        if k is None:
            return False
        pool.pop(k)
    return True


def check(ws, ref):
    name = "interest per payment, one line"
    hits = sorted(glob.glob(os.path.join(ws, "loan_check.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "loan_check.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "loan_check.xlsx not found"}]
    try:
        want = [float(v) for v in json.load(open(os.path.join(ref, "notes.json")))["interest_by_payment"]]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        grade = _bench()
        from openpyxl import load_workbook
        wb = load_workbook(grade.recalculated_workbook(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    best = 0
    for line in _lines(wb, grade.cell_num):
        if _covers(line, want, 0.011):
            return [{"name": name, "passed": True, "detail": f"all {len(want)} interest figures sit on one line"}]
        best = max(best, sum(1 for w in want if any(abs(v - w) <= 0.011 for v in line)))
    return [{"name": name, "passed": False, "detail": f"no row or column holds all {len(want)} interest figures (best line matched {best})"}]
