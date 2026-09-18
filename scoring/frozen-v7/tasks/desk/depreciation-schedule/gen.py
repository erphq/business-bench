#!/usr/bin/env python3
"""depreciation-schedule: a print shop's fixed asset register to its 2026 straight-line depreciation schedule.

    python gen.py [--seed N] [--naive DIR]

Business: Quill & Ink, a commercial print shop with presses, finishing gear, two vans and a leasehold HVAC unit.
The office manager keeps the asset register; the outside accountant's email carries the policy: straight line by
month to salvage, start the month after an asset goes into service, stop the month an asset is sold, never go below
salvage, and leave anything under the $2,500 capitalization threshold off the schedule.

Traps (each caught by a check, see task.yaml):
  * an asset placed in service mid-2026 starts the month after                   (check: partial first year)
  * salvage values come off the depreciable base                                 (check: press with salvage)
  * the old van was sold mid-year: no depreciation in or after the sale month, and
    a gain or loss against book value then                                       (checks: sold van; gain on the sale)
  * the prepress server reaches the end of its life in spring 2026 and stops     (check: fully depreciated server)
  * three items under the $2,500 threshold sit in the register                   (check: total 2026 depreciation)
  * lives are written as years or months, costs and salvage are text, and an old
    2024 copy of the register is in the folder                                   (check: depreciation, one line per month)
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "lib"))
from bizgen import *  # noqa: E402,F403
from openpyxl.utils import get_column_letter  # noqa: E402

FY = 2026
LABELS = [f"{FY}-{m:02d}" for m in range(1, 13)]
THRESHOLD = 250000  # cents

# key, description, category, life months, cost range (dollars), salvage share, in-service window, location
TEMPLATES = [
    ("press", "Heidelberg Speedmaster SM 52 two-color offset press", "Production equipment", 120, (52000, 78000), (0.08, 0.10),
     (date(2019, 3, 1), date(2021, 6, 30)), "Press room"),
    ("digital", "Konica Minolta AccurioPress C4080 digital press", "Production equipment", 60, (38000, 46000), (0.05, 0.06),
     (date(2023, 1, 5), date(2024, 5, 30)), "Press room"),
    ("wide", "Epson SureColor wide-format printer", "Production equipment", 60, (9000, 14000), (0.0, 0.0),
     (date(2024, 2, 1), date(2025, 6, 30)), "Sign shop"),
    ("cutter", "Polar 66 guillotine paper cutter", "Finishing equipment", 84, (18000, 26000), (0.05, 0.05),
     (date(2020, 1, 6), date(2022, 6, 30)), "Bindery"),
    ("folder", "MBO K70 buckle folder", "Finishing equipment", 84, (12000, 17000), (0.0, 0.0),
     (date(2021, 2, 1), date(2023, 3, 31)), "Bindery"),
    ("hvac", "Rooftop HVAC unit (leasehold improvement)", "Leasehold improvements", 180, (21000, 28000), (0.0, 0.0),
     (date(2022, 4, 1), date(2022, 9, 30)), "Building"),
    ("pos", "Front counter POS terminals and office PCs", "Computers", 36, (3400, 4800), (0.0, 0.0),
     (date(2025, 1, 6), date(2025, 7, 31)), "Front office"),
    ("van_new", "2026 Ford Transit 250 cargo van", "Vehicles", 60, (46000, 52000), (0.16, 0.20),
     (date(2026, 6, 1), date(2026, 7, 31)), "Lot"),
    ("inserter", "Mailing inserter and folder-sealer", "Finishing equipment", 84, (14000, 19000), (0.0, 0.0),
     (date(2026, 12, 3), date(2026, 12, 18)), "Mail room"),
    # the trap assets
    ("laminator", "GBC Catena roll laminator", "Finishing equipment", 60, (5200, 8800), (0.0, 0.0),
     (date(2026, 3, 2), date(2026, 5, 28)), "Bindery"),
    ("van_old", "2021 Ford Transit Connect delivery van", "Vehicles", 60, (26000, 31000), (0.18, 0.24),
     (date(2022, 2, 1), date(2022, 8, 31)), "Lot"),
    ("prepress", "Prepress workstation and RIP server", "Computers", 36, (6000, 9500), (0.0, 0.0),
     (date(2023, 2, 6), date(2023, 6, 28)), "Prepress"),
]
SMALL = [("chairs", "Ergonomic office chairs (6)", "Furniture", 60, (1500, 2100), "Front office"),
         ("label", "Zebra thermal label printer", "Computers", 36, (480, 790), "Mail room"),
         ("shelving", "Paper storage shelving", "Furniture", 84, (1900, 2450), "Warehouse")]
BUYERS = ["Harper Auto Sales", "Midway Motors", "a private buyer", "Lakeshore Fleet Resale"]


def midx(d: date) -> int:
    return d.year * 12 + d.month - 1


def build(seed: int) -> dict:
    r = rng(seed)
    assets = []
    tags = sorted(r.sample(range(101, 199), len(TEMPLATES) + len(SMALL)))
    order = list(range(len(TEMPLATES) + len(SMALL)))
    for k, (key, desc, cat, life, (lo, hi), (s_lo, s_hi), (d0, d1), loc) in enumerate(TEMPLATES):
        c0 = r.uniform(lo, hi)
        salvage = int(round(c0 * r.uniform(s_lo, s_hi) / 50.0)) * 50 * 100
        monthly = int(round((c0 * 100 - salvage) / life))
        cost = monthly * life + salvage
        assets.append({"key": key, "desc": desc, "cat": cat, "life": life, "cost": cost, "salvage": salvage, "monthly": monthly,
                       "in_service": day_in(r, d0, d1, weekday_only=True), "loc": loc, "sale": None, "proceeds": None, "small": False})
    for key, desc, cat, life, (lo, hi), loc in SMALL:
        cost = int(round(r.uniform(lo, hi) * 100))
        assets.append({"key": key, "desc": desc, "cat": cat, "life": life, "cost": cost, "salvage": 0, "monthly": 0,
                       "in_service": day_in(r, date(2024, 3, 1), date(2026, 8, 30), weekday_only=True), "loc": loc,
                       "sale": None, "proceeds": None, "small": True})
    r.shuffle(order)
    for a, i in zip(assets, order):
        a["tag"] = f"FA-{tags[i]:03d}"
    van = next(a for a in assets if a["key"] == "van_old")
    van["sale"] = date(2026, r.choice([6, 7, 8, 9]), r.randint(3, 27))
    van["buyer"] = r.choice(BUYERS)

    # ---- truth
    jan = FY * 12
    for a in assets:
        if a["small"]:
            continue
        start = midx(a["in_service"]) + 1
        end = start + a["life"] - 1
        stop = midx(a["sale"]) - 1 if a["sale"] else end
        last = min(end, stop)

        def dep(k, a=a, start=start, last=last):
            return a["monthly"] if start <= k <= last else 0
        a["accum_before"] = sum(dep(k) for k in range(start, jan))
        a["fy"] = [dep(jan + i) for i in range(12)]
        a["dep_2026"] = sum(a["fy"])
        a["accum_end"] = a["accum_before"] + a["dep_2026"]
        assert a["accum_end"] <= a["cost"] - a["salvage"]
        a["book"] = a["cost"] - a["accum_end"]
    van["proceeds"] = int(round((van["book"] / 100) * r.uniform(1.06, 1.28) / 50.0)) * 50 * 100  # a gain, so the sign is not a matter of labeling
    van["gain"] = van["proceeds"] - van["book"]
    kept = [a for a in assets if not a["small"]]
    kept.sort(key=lambda a: a["tag"])
    monthly_tot = [sum(a["fy"][i] for a in kept) for i in range(12)]
    return {"assets": assets, "kept": kept, "monthly": monthly_tot, "total": sum(monthly_tot), "by_key": {a["key"]: a for a in assets}}


def naive(d: dict) -> list[dict]:
    """In-service month counts, no salvage, no stop at end of life or at the sale, every register row kept."""
    out = []
    for a in d["assets"]:
        m = int(round(a["cost"] / a["life"]))
        start = midx(a["in_service"])
        jan = FY * 12
        fy = [m if jan + i >= start else 0 for i in range(12)]
        before = m * max(0, jan - start)
        b = dict(a, monthly=m, salvage=a["salvage"], fy=fy, dep_2026=sum(fy), accum_before=before, accum_end=before + sum(fy))
        b["book"] = a["cost"] - b["accum_end"]
        if a["proceeds"] is not None:
            b["gain"] = a["proceeds"] - b["book"]
        out.append(b)
    out.sort(key=lambda a: a["tag"])
    return out


def acceptable(d: dict) -> bool:
    by = d["by_key"]
    nv = {a["key"]: a for a in naive(d)}
    van = by["van_old"]
    if not (0 < van["fy"].count(0) < 12) or by["prepress"]["dep_2026"] == 0 or by["prepress"]["fy"][-1] != 0:
        return False
    if abs(van["gain"]) < 20000:
        return False
    for key, field in (("laminator", "dep_2026"), ("press", "dep_2026"), ("van_old", "dep_2026"), ("prepress", "dep_2026"),
                       ("van_old", "gain")):
        a = by[key]
        v = a[field]
        # every other figure an agent's row carries (accumulated at year end equals 2026 depreciation for a new asset by definition)
        row = [a["cost"], a["salvage"], a["life"] * 100, a["monthly"], a["accum_before"], a["book"], a["proceeds"] or 0] + a["fy"]
        row += [a["dep_2026"], a["accum_end"]] if field == "gain" else [a.get("gain") or 0]
        if any(abs(x - v) <= 100 for x in row):
            return False
        if abs(nv[key][field] - v) < 500:
            return False
    if abs(sum(a["dep_2026"] for a in naive(d)) - d["total"]) < 10000:
        return False
    if len(set(d["monthly"])) < 4:
        return False
    return True


# --------------------------------------------------------------------------- deliverable

def workbook(kept: list[dict]) -> dict:
    n = len(kept)
    rows = []
    for i, a in enumerate(kept, start=2):
        line = [a["tag"], a["desc"], a["in_service"], a["cost"] / 100, a["salvage"] / 100, a["life"],
                f"=ROUND((D{i}-E{i})/F{i},2)", a["accum_before"] / 100] + [v / 100 for v in a["fy"]]
        line += [f"=SUM(I{i}:T{i})", f"=H{i}+U{i}", f"=D{i}-V{i}"]
        if a.get("proceeds") is not None:
            line += [a["proceeds"] / 100, f"=X{i}-W{i}"]
        rows.append(line)
    tr = n + 2
    rows.append(["Total", "", "", f"=SUM(D2:D{n + 1})", f"=SUM(E2:E{n + 1})", "", "", f"=SUM(H2:H{n + 1})"]
                + [f"=SUM({get_column_letter(j)}2:{get_column_letter(j)}{n + 1})" for j in range(9, 24)])
    monthly = [[LABELS[i], f"=Assets!{get_column_letter(9 + i)}{tr}"] for i in range(12)]
    monthly.append(["Total 2026", "=SUM(B2:B13)"])
    return {
        "Monthly": {"header": ["Month", "Depreciation expense"], "rows": monthly, "number_formats": {"B": "#,##0.00"},
                    "widths": {"A": 12, "B": 22}},
        "Assets": {"header": ["Asset tag", "Description", "In service", "Cost", "Salvage", "Life (months)", "Monthly depreciation",
                              "Accumulated at 2025-12-31"] + LABELS + ["2026 depreciation", "Accumulated at 2026-12-31 or sale",
                                                                     "Book value (at sale if sold)", "Sale proceeds", "Gain (loss) on sale"],
                   "rows": rows, "number_formats": {get_column_letter(j): "#,##0.00" for j in list(range(4, 6)) + list(range(7, 26))},
                   "widths": {"B": 40, "C": 12}, "freeze": "C2"},
    }


def life_text(months: int, style: int) -> str:
    if months % 12 == 0 and style % 3 != 2:
        y = months // 12
        return f"{y} yrs" if style % 3 == 0 else f"{y} years"
    return f"{months} months"


def emit(seed: int, naive_dir: str | None) -> None:
    d = build(seed)
    if naive_dir:
        os.makedirs(naive_dir, exist_ok=True)
        write_xlsx(os.path.join(naive_dir, "depreciation.xlsx"), workbook(naive(d)), creator="naive")
        return
    ws, ref, sol = task_dirs(HERE)
    r = rng(seed + 5)
    by = d["by_key"]
    van = by["van_old"]

    # ---- workspace: the register
    reg = []
    for a in sorted(d["assets"], key=lambda a: (a["in_service"], a["tag"])):
        k = sum(ord(ch) for ch in a["tag"])
        in_service = a["in_service"] if k % 4 else a["in_service"].strftime("%m/%d/%Y")
        salvage = "" if a["salvage"] == 0 and k % 2 else money_str(a["salvage"] / 100, 1)
        note = ""
        if a["sale"]:
            note = f"Sold {a['sale'].strftime('%m/%d/%Y')} to {a['buyer']} for {money_str(a['proceeds'] / 100, 1)}"
        elif a["key"] == "pos":
            note = "3 terminals + 2 PCs, one invoice"
        reg.append([a["tag"], a["desc"], a["cat"], a["loc"], in_service, money_str(a["cost"] / 100, 1), salvage,
                    life_text(a["life"], k), "Sold" if a["sale"] else "In use", note])
    write_xlsx(os.path.join(ws, "fixed_asset_register.xlsx"), {"Register": {
        "merged_title": "Quill & Ink - Fixed Asset Register",
        "preamble": [["Maintained by: Bev Castillo (office)", "", "", "", "", "", "", "", "", "Last updated 01/05/2027"]],
        "header": ["Tag", "Description", "Category", "Location", "Placed in service", "Cost", "Salvage value", "Useful life", "Status", "Notes"],
        "rows": reg, "widths": {"B": 46, "C": 22, "D": 14, "E": 16, "F": 14, "G": 14, "H": 12, "J": 48}, "freeze": "A4"}},
        creator="Bev Castillo")
    old = [row[:] for row, a in zip(reg, sorted(d["assets"], key=lambda a: (a["in_service"], a["tag"]))) if a["in_service"] < date(2024, 7, 1)]
    for row in old:
        row[8], row[9] = "In use", ""
    write_xlsx(os.path.join(ws, "fixed_asset_register_2024_backup.xlsx"), {"Register": {
        "merged_title": "Quill & Ink - Fixed Asset Register",
        "preamble": [["Maintained by: Bev Castillo (office)", "", "", "", "", "", "", "", "", "Last updated 06/28/2024"]],
        "header": ["Tag", "Description", "Category", "Location", "Placed in service", "Cost", "Salvage value", "Useful life", "Status", "Notes"],
        "rows": old, "widths": {"B": 46, "C": 22, "J": 30}}}, creator="Bev Castillo")

    # ---- workspace: the accountant's email
    write_email_thread(os.path.join(ws, "email_from_pat_accountant.txt"), [
        {"from": "Pat Lindqvist <pat@lindqvistbooks.com>", "to": "Hollis Grant <hollis@quillandink.print>", "date": "Wed, 6 Jan 2027 09:20",
         "subject": "2026 book depreciation",
         "body": ("Hi Hollis,\n\nI need the 2026 book depreciation schedule before I can close the year. Bev's register should have "
                  "everything. How we do it (same as always):\n\n"
                  "1. Straight line by month: cost less salvage value, divided by the useful life in months.\n"
                  "2. An asset starts depreciating the month after it is placed in service, whatever day of the month that was.\n"
                  "3. When something is sold or scrapped there is no depreciation in that month or after it. Work out the gain or loss "
                  "on the sale against its book value at that point.\n"
                  "4. Book value never goes below salvage. Once an asset is fully depreciated it stops.\n"
                  "5. Anything that cost under $2,500 we expense when we buy it. It does not belong on the schedule even if it made it "
                  "into the register.\n\n"
                  "What I need back: 2026 depreciation expense by month, one line per month with the year total, and the asset-by-asset "
                  "detail behind it (cost, salvage, life, monthly amount, accumulated depreciation at the start and end of the year, book "
                  "value, and the gain or loss on anything sold). Please leave the formulas in so I can follow the numbers.\n\nThanks,\nPat")},
        {"from": "Hollis Grant <hollis@quillandink.print>", "to": "Pat Lindqvist <pat@lindqvistbooks.com>", "date": "Wed, 6 Jan 2027 10:02",
         "subject": "RE: 2026 book depreciation",
         "body": ("Bev says the register is current as of this week, it is fixed_asset_register.xlsx. The 2024 file is just her old backup.\n\n"
                  "Hollis")}])

    # ---- reference
    write_csv(os.path.join(ref, "assets.csv"), ["tag", "description", "monthly", "accumulated_2025_12_31", "depreciation_2026",
                                                "accumulated_2026_12_31", "book_value", "gain_loss"],
              [[a["tag"], a["desc"], f"{a['monthly'] / 100:.2f}", f"{a['accum_before'] / 100:.2f}", f"{a['dep_2026'] / 100:.2f}",
                f"{a['accum_end'] / 100:.2f}", f"{a['book'] / 100:.2f}", f"{a['gain'] / 100:.2f}" if a.get("gain") is not None else ""]
               for a in d["kept"]])
    write_json(os.path.join(ref, "notes.json"), {
        "depreciation_by_month": {LABELS[i]: round(d["monthly"][i] / 100, 2) for i in range(12)},
        "total_2026": round(d["total"] / 100, 2), "excluded_under_threshold": [a["tag"] for a in d["assets"] if a["small"]],
        "pinned": {k: by[k]["tag"] for k in ("laminator", "press", "van_old", "prepress")}})

    # ---- reference solution
    write_xlsx(os.path.join(sol, "depreciation.xlsx"), workbook(d["kept"]), creator="Quill & Ink")

    def pin(name, key, field="dep_2026"):
        a = by[key]
        return {"type": "xlsx_value_present", "name": name, "path": "depreciation.xlsx", "expected": round(a[field] / 100, 2),
                "rel_tol": 0.000001, "near_text": a["tag"].lower()}

    lam, press, pre = by["laminator"], by["press"], by["prepress"]
    small = [a for a in d["assets"] if a["small"]]
    write_task_yaml(HERE, {
        "id": "depreciation-schedule", "track": "desk", "category": "bookkeeping",
        "title": "2026 depreciation schedule from the fixed asset register",
        "ask": "Pat needs our 2026 depreciation schedule so she can close the year; her email explains how she wants it. Save it as depreciation.xlsx.\n",
        "followup": None, "timeout_s": 1200,
        "traps": [
            f"the laminator {lam['tag']} went into service {lam['in_service'].isoformat()} and starts depreciating the month after, so it "
            f"has {sum(1 for v in lam['fy'] if v)} months in 2026, not {sum(1 for v in lam['fy'] if v) + 1} (check: partial first year)",
            f"the press {press['tag']} has a {press['salvage'] / 100:,.2f} salvage value that comes off the base; cost over life overstates "
            "every month (check: press with salvage)",
            f"the old van {van['tag']} was sold {van['sale'].isoformat()} (only in the Notes column); depreciation stops the month before the "
            "sale and the gain or loss is proceeds less book value at that point, not a full year of depreciation "
            "(checks: sold van; gain on the van sale)",
            f"the prepress server {pre['tag']} went into service {pre['in_service'].isoformat()} on a 36-month life and is fully "
            "depreciated partway through 2026; running it all year takes it below zero (check: fully depreciated server)",
            f"{len(small)} register rows ({', '.join(a['tag'] for a in small)}) cost under the $2,500 threshold and must be left off "
            "(check: total 2026 depreciation)",
            "lives are written '10 yrs', '7 years' and '36 months', cost and salvage are text, some in-service dates are text, blank "
            "salvage means none, a December 2026 addition earns nothing until 2027, and a 2024 backup of the register sits beside the "
            "current one (checks: depreciation, one line per month; total 2026 depreciation)",
        ],
        "checks": [
            {"type": "xlsx_has_formulas", "name": "live formulas", "path": "depreciation.xlsx", "min_count": 12},
            {"type": "xlsx_no_errors", "name": "no formula errors", "path": "depreciation.xlsx"},
            {"type": "custom", "name": "depreciation, one line per month", "module": "check.py"},
            {"type": "xlsx_value_present", "name": "total 2026 depreciation", "path": "depreciation.xlsx",
             "expected": round(d["total"] / 100, 2), "rel_tol": 0.000001, "near_text": "total"},
            pin("partial first year", "laminator"),
            pin("press with salvage", "press"),
            pin("sold van", "van_old"),
            pin("gain on the van sale", "van_old", "gain"),
            pin("fully depreciated server", "prepress"),
        ],
    })
    print(f"seed={seed} assets={len(d['kept'])} total={d['total'] / 100:.2f} monthly={[v / 100 for v in d['monthly']]}")
    for k in ("laminator", "press", "van_old", "prepress"):
        print(k, by[k]["tag"], by[k]["dep_2026"] / 100, by[k].get("gain"))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--naive", default=None)
    a_ = ap.parse_args()
    for attempt in range(400):
        if acceptable(build(a_.seed * 1000 + attempt)):
            break
    else:
        raise SystemExit("no acceptable draw")
    emit(a_.seed * 1000 + attempt, a_.naive)
