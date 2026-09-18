"""The on-leave rep's April cell must read 0, not blank.

xlsx_value_present cannot tell a blank from a zero at one row/column intersection (a 0 anywhere on a row
mentioning the rep would pass), so this check finds the rep's label cell and the April header cell on the same
sheet and reads the cell where that row and column cross, after recalculation. Both layouts pass: reps down the
side with months across, or months down the side with reps across. The rep and month come from
reference/notes.json, so nothing is pinned to a seed.
"""
import glob
import importlib.util
import json
import os
import re
from datetime import date, datetime

NAME = "Patel April reads 0"


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _is_month(v, year, month):
    if isinstance(v, (datetime, date)):
        return v.year == year and v.month == month
    if not isinstance(v, str):
        return False
    s = v.strip().lower()
    names = {1: "jan", 2: "feb", 3: "mar", 4: "apr", 5: "may", 6: "jun", 7: "jul", 8: "aug", 9: "sep", 10: "oct", 11: "nov", 12: "dec"}
    pats = [rf"\b{year}-{month:02d}\b", rf"\b0?{month}/{year}\b", rf"\b{year}/0?{month}\b", rf"^{names[month]}[a-z]*\.?(\s+'?(\d{{2}}|\d{{4}}))?$"]
    return any(re.search(p, s) for p in pats)


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("-", "$ -", "$-"):
        return 0.0
    m = re.fullmatch(r"\(?-?[$]?\s?(\d[\d,]*(?:\.\d+)?)\)?", s)
    return float(m.group(1).replace(",", "")) if m else None


def check(ws, ref):
    hits = sorted(glob.glob(os.path.join(ws, "sales_by_rep.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "sales_by_rep.xlsx"), recursive=True))
    if not hits:
        return [{"name": NAME, "passed": False, "detail": "sales_by_rep.xlsx not found"}]
    try:
        z = json.load(open(os.path.join(ref, "notes.json")))["zero_cell"]
        surname = z["surname"].lower()
        year, month = int(z["month"][:4]), int(z["month"][5:7])
    except Exception as e:
        return [{"name": NAME, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": NAME, "passed": False, "detail": f"workbook unreadable: {e}"}]

    seen = []
    for sh in wb.worksheets:
        grid = [[c.value for c in row] for row in sh.iter_rows()]
        if not grid:
            continue
        rep_cells = [(i, j) for i, row in enumerate(grid) for j, v in enumerate(row)
                     if isinstance(v, str) and re.search(rf"\b{re.escape(surname)}\b", v.lower())]
        month_cells = [(i, j) for i, row in enumerate(grid) for j, v in enumerate(row) if _is_month(v, year, month)]
        for ri, rj in rep_cells:
            for mi, mj in month_cells:
                # reps down / months across: header above the rep row, in another column
                if mi < ri and mj != rj:
                    v = grid[ri][mj] if mj < len(grid[ri]) else None
                    seen.append(f"{sh.title}!r{ri + 1}c{mj + 1}={v!r}")
                    n = _num(v)
                    if n is not None and abs(n) < 0.005:
                        return [{"name": NAME, "passed": True, "detail": f"{sh.title} row {ri + 1}, column {mj + 1} reads {v!r}"}]
                # months down / reps across: rep header above the month row
                if ri < mi and mj != rj:
                    v = grid[mi][rj] if rj < len(grid[mi]) else None
                    seen.append(f"{sh.title}!r{mi + 1}c{rj + 1}={v!r}")
                    n = _num(v)
                    if n is not None and abs(n) < 0.005:
                        return [{"name": NAME, "passed": True, "detail": f"{sh.title} row {mi + 1}, column {rj + 1} reads {v!r}"}]
    detail = ("no cell where the rep's row meets the April column reads 0: " + "; ".join(seen[:6])) if seen else \
        "could not find a row labelled with the rep crossing an April column"
    return [{"name": NAME, "passed": False, "detail": detail}]
