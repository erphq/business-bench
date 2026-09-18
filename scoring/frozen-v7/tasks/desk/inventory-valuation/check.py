"""Every stocked item with no usable cost is listed in valuation.xlsx on a row that flags it.

A flag is any text on the SKU's row saying the cost is missing ("no cost", "missing cost", "cost TBD", "needs cost",
"chase", "flag", "not valued", "unknown cost" and similar). The SKUs come from reference/notes.json and are matched
case-insensitively with surrounding spaces ignored, on any sheet, so nothing here is pinned to a seed.
"""
import glob
import json
import os
import re

FLAG = re.compile(
    r"(\bno\b[^|]{0,20}\b(cost|price)|\bmissing\b|\bwithout\b[^|]{0,12}\b(cost|price)|\bunknown\b|\bneeds?\b[^|]{0,12}\b(cost|price)"
    r"|\bchase\b|\bflag|\bnot valued\b|\bexcluded\b|\bno receipt|\btbd\b|\bn/?a\b|\bunpriced\b|\bnot costed\b|\bcost not\b|\bno data\b)",
    re.I)


def _recalculated(path):
    """Reuse the grader's own recalculation (LibreOffice, else the `formulas` engine)."""
    try:
        import importlib.util
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def check(ws, ref):
    name = "items without a cost are flagged"
    hits = sorted(glob.glob(os.path.join(ws, "valuation.xlsx"))) or sorted(glob.glob(os.path.join(ws, "**", "valuation.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "valuation.xlsx not found"}]
    try:
        skus = json.load(open(os.path.join(ref, "notes.json")))["no_cost_skus"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)   # a flag produced by a formula counts by its result
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    flagged = set()
    for sh in wb.worksheets:
        for row in sh.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None]
            on_row = {c.upper() for c in cells}
            for sku in skus:
                if sku.upper() in on_row and any(FLAG.search(t) for t in cells if t.upper() != sku.upper()):
                    flagged.add(sku)
    missing = [s for s in skus if s not in flagged]
    return [{"name": name, "passed": not missing,
             "detail": f"all {len(skus)} flagged" if not missing else f"not flagged on their row: {missing}"}]
