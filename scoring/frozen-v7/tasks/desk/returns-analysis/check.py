"""Return rates, accepted as a fraction (0.0816) or a percentage (8.16 / "8.16%").

Rates are legitimately written either way, so this check recalculates the workbook with the grader's own
engine and accepts both forms. Expected rates come from reference/notes.json (nothing pinned to a seed):
the overall quarter rate on a row or column mentioning "total", and the exchange-heavy jacket's rate on a
row or column mentioning its product name.
"""
import glob
import importlib.util
import json
import os
import re


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*(\d[\d,]*(?:\.\d+)?)\s?%?\s*", str(v))
    return float(m.group(1).replace(",", "")) if m else None


def check(ws, ref):
    name = "return rates (fraction or percent)"
    hits = sorted(glob.glob(os.path.join(ws, "returns.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "returns.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "returns.xlsx not found"}]
    try:
        rates = json.load(open(os.path.join(ref, "notes.json")))["rates"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    cells = []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        if not grid:
            continue
        col_text = [""] * max(len(r) for r in grid)
        for r in grid:
            for j, c in enumerate(r):
                if c.value is not None:
                    col_text[j] += " " + str(c.value).lower()
        for r in grid:
            row_text = " ".join(str(c.value).lower() for c in r if c.value is not None)
            for j, c in enumerate(r):
                v = _num(c.value)
                if v is not None:
                    cells.append((row_text, col_text[j], v))
    out = []
    for label, want in sorted(rates.items()):
        ok = any((label in rt or label in ct) and (abs(v - want) <= 0.0015 or abs(v - want * 100) <= 0.15)
                 for rt, ct, v in cells)
        out.append({"name": f"return rate {label}", "passed": ok,
                    "detail": f"expected {want:.4f} (or {want * 100:.2f}%)" + ("" if ok else " not found")})
    return out
