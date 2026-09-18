"""Utilization per person and for the firm, accepted as a fraction (0.8123) or a percentage (81.23 / "81.23%").

xlsx_value_present pins one numeric form and a rate is legitimately written either way, so this check recalculates the
workbook with the grader's own engine and accepts both. A person's rate must sit on a row or column that mentions the
person's first name; the firm rate on one that mentions "total" or "firm". Expected rates come from reference/notes.json,
so nothing here is pinned to a seed.
"""
import glob
import importlib.util
import json
import os
import re


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s?%?\s*", str(v))
    return float(m.group(1)) if m else None


def check(ws, ref):
    name = "utilization per person"
    hits = sorted(glob.glob(os.path.join(ws, "utilization.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "utilization.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "utilization.xlsx not found"}]
    try:
        notes = json.load(open(os.path.join(ref, "notes.json")))
        want = {n.split()[0].lower(): v for n, v in notes["utilization"].items()}
        firm = notes["firm_utilization"]
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]
    cells = []
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows()]
        if not grid:
            continue
        ncols = max(len(r) for r in grid)
        col_text = [""] * ncols
        for r in grid:
            for j, c in enumerate(r):
                if c.value is not None:
                    col_text[j] += " " + str(c.value).lower()
        for r in grid:
            row_text = " ".join(str(c.value).lower() for c in r if c.value is not None)
            for j, c in enumerate(r):
                v = _num(c.value)
                if v is not None:
                    cells.append((row_text, col_text[j], v))

    def found(keys, target):
        for row_text, col_text, v in cells:
            if not any(k in row_text or k in col_text for k in keys):
                continue
            if abs(v - target) <= max(target * 0.005, 0.0015) or abs(v - target * 100) <= max(target * 0.5, 0.15):
                return True
        return False

    bad = [f"{k}: {t:.4f}" for k, t in sorted(want.items()) if not found([k], t)]
    if not found(["total", "firm"], firm):
        bad.append(f"firm: {firm:.4f}")
    return [{"name": name, "passed": not bad,
             "detail": f"{len(want)} people and the firm match" if not bad else "missing or wrong: " + "; ".join(bad)}]
