"""Active subscribers at each month end and cancellations effective in each month, January to June 2026.

Months may run down the side or across the top, labelled 2026-03, Mar 2026, March, 03/2026 or as a date cell.
For every month, some row or column that names the month must hold both the active count and the cancellation
count. Expected figures come from reference/monthly_counts.csv, so nothing here is pinned to a seed.
"""
import csv
import glob
import importlib.util
import os
import re

NAMES = ["jan", "feb", "mar", "apr", "may", "jun"]
FULL = ["january", "february", "march", "april", "may", "june"]


def _recalculated(path):
    try:
        root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        spec = importlib.util.spec_from_file_location("bench_grade_for_check", os.path.join(root, "bench", "grade.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.recalculated_workbook(path)
    except Exception:
        return path


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r"\s*(\d[\d,]*(?:\.\d+)?)\s*", str(v))
    return float(m.group(1).replace(",", "")) if m else None


def _month_pattern(i):
    n = i + 1
    return re.compile(rf"(\b2026-0?{n}\b|\b0?{n}/2026\b|\b{NAMES[i]}(?:[a-z]*)?\b[\s\-'.]*(?:2026|26)?)", re.I)


def check(ws, ref):
    name = "active and cancelled by month"
    hits = sorted(glob.glob(os.path.join(ws, "subscriptions.xlsx"))) or sorted(
        glob.glob(os.path.join(ws, "**", "subscriptions.xlsx"), recursive=True))
    if not hits:
        return [{"name": name, "passed": False, "detail": "subscriptions.xlsx not found"}]
    try:
        with open(os.path.join(ref, "monthly_counts.csv"), encoding="utf-8") as f:
            truth = list(csv.DictReader(f))
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"reference unreadable: {e}"}]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(_recalculated(hits[0]), data_only=True)
        # labels come from the workbook as saved: recalculation engines can turn date cells into serial numbers
        raw = load_workbook(hits[0])
    except Exception as e:
        return [{"name": name, "passed": False, "detail": f"workbook unreadable: {e}"}]

    def label(v):
        if v is None or (isinstance(v, str) and v.startswith("=")):
            return None
        return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v)

    groups = []
    raw_by_title = {t.lower(): raw[t] for t in raw.sheetnames}   # engines may change the case of sheet names
    for sh in wb.worksheets:
        grid = [list(row) for row in sh.iter_rows(values_only=True)]
        if not grid:
            continue
        rsh = raw_by_title.get(sh.title.lower())
        rgrid = [list(row) for row in rsh.iter_rows(values_only=True)] if rsh is not None else grid
        ncols = max(len(r) for r in grid)
        cols_text = [[] for _ in range(ncols)]
        cols_nums = [[] for _ in range(ncols)]
        for i, row in enumerate(grid):
            rrow = rgrid[i] if i < len(rgrid) else []
            texts, nums = [], []
            for j, v in enumerate(row):
                rv = rrow[j] if j < len(rrow) else v
                t = label(rv) if rv is not None else label(v)
                n = _num(v) if not hasattr(rv, "strftime") else None
                if t is not None:
                    texts.append(t); cols_text[j].append(t)
                if n is not None:
                    nums.append(n); cols_nums[j].append(n)
            groups.append((" ".join(texts), nums))
        for j in range(ncols):
            groups.append((" ".join(cols_text[j]), cols_nums[j]))
    bad = []
    for i, t in enumerate(truth):
        pat = _month_pattern(i)
        want_a, want_c = float(t["active"]), float(t["cancelled"])
        ok = False
        for text, nums in groups:
            # a group naming several months (a month column or a header row) cannot pin one month's figures
            named = sum(1 for k in range(6) if _month_pattern(k).search(text))
            if named != 1 or not pat.search(text):
                continue
            if any(abs(n - want_a) < 0.01 for n in nums) and any(abs(n - want_c) < 0.01 for n in nums):
                ok = True
                break
        if not ok:
            bad.append(f"{t['month']}: expected {int(want_a)} active and {int(want_c)} cancelled")
    return [{"name": name, "passed": not bad, "detail": "; ".join(bad) if bad else "all six months match"}]
