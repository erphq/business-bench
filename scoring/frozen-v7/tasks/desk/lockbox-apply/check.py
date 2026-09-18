"""lockbox-apply: amounts applied and written off per invoice, and the lines of the multi-invoice checks.

Invoice numbers are compared on their digits (INV-48213, 48213 and INV-048213 are one invoice); blank write_off reads
as 0. Every invoice the reference applies to must carry the right applied total and write-off; the holding-company
check and the two checks that split one invoice must each carry a line for every invoice they pay.
Figures come from reference/lockbox.json.
"""
import csv
import glob
import io
import json
import os
import re


def _find(ws, name):
    hits = sorted(glob.glob(os.path.join(ws, name))) or sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    return hits[0] if hits else None


def _norm(c):
    return re.sub(r"[^a-z0-9]+", "_", str(c or "").strip().lower()).strip("_")


def read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        vals = [["" if v is None else str(v) for v in r] for r in load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    else:
        text = open(path, "rb").read().decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
        vals = list(csv.reader(io.StringIO(text)))
    if not vals:
        return []
    hdr = [_norm(h) for h in vals[0]]
    return [dict(zip(hdr, row)) for row in vals[1:] if any(str(x).strip() for x in row)]


def inv_key(v):
    dg = re.sub(r"\D", "", str(v))
    return str(int(dg)) if dg else ""


def num(v):
    s = str(v).strip()
    if s in ("", "-"):
        return 0.0
    neg = s.startswith("(") or s.startswith("-")
    try:
        x = float(re.sub(r"[^0-9.]", "", s))
    except ValueError:
        return None
    return -x if neg else x


def chk_key(v):
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s.lstrip("0") or s


def check(ws, ref):
    p = _find(ws, "applied.csv")
    if not p:
        return [{"name": "applied.csv present", "passed": False, "detail": "not found"}]
    rows = read_rows(p)
    need = {"check_number", "invoice_number", "amount_applied"}
    if not rows or not need <= set(rows[0]):
        return [{"name": "applied.csv columns", "passed": False, "detail": f"need {sorted(need)}; have {sorted(rows[0]) if rows else []}"}]
    lb = json.load(open(os.path.join(ref, "lockbox.json")))
    got, lines, bad = {}, {}, 0
    for r in rows:
        k = inv_key(r["invoice_number"])
        a, w = num(r.get("amount_applied", "")), num(r.get("write_off", ""))
        if a is None or w is None:
            bad += 1
            continue
        e = got.setdefault(k, {"applied": 0.0, "write_off": 0.0})
        e["applied"] += a
        e["write_off"] += w
        lines.setdefault(chk_key(r["check_number"]), set()).add(k)
    want = lb["per_invoice"]
    traps = {str(t) for t in lb["trap_invoices"]}
    wrong, trap_wrong = [], []
    for k, v in want.items():
        g = got.get(k, {"applied": 0.0, "write_off": 0.0})
        if abs(g["applied"] - v["applied"]) > 0.011 or abs(g["write_off"] - v["write_off"]) > 0.011:
            wrong.append(f"INV-{k}: applied {g['applied']:.2f}/{v['applied']:.2f}, write-off {g['write_off']:.2f}/{v['write_off']:.2f}")
            if k in traps:
                trap_wrong.append(k)
    out = [{"name": "every stub invoice applied with the right amount and write-off", "passed": not wrong and bad == 0,
            "detail": f"{len(want) - len(wrong)}/{len(want)} invoices right" + (f"; wrong: {wrong[:6]}" if wrong else "")
                      + (f"; trap invoices wrong: {trap_wrong}" if trap_wrong else "") + (f"; {bad} unreadable amounts" if bad else "")}]
    miss = []
    for chk, invs in lb["multi_line_checks"].items():
        have = lines.get(chk_key(chk), set())
        if not {str(i) for i in invs} <= have:
            miss.append(f"check {chk} -> {sorted(invs)} (have {sorted(have)})")
    out.append({"name": "holding-company and split checks carry a line per invoice", "passed": not miss,
                "detail": "all present" if not miss else f"missing: {miss}"})
    return out
