"""Custom grader for quote-comparison: the memo prices the vendor it recommends.

The pick itself is a judgment call (Bluepine on warranty, Northstar on price), so the memo is
not graded on which vendor it names. It is graded on stating the correct delivered total of
the vendor it does recommend: a memo that says "go with Bluepine" and never prices Bluepine,
or prices it at the fee-less 48,870, fails. The recommended vendor is the one named right
after a recommendation verb in a sentence that does not negate it; when no such sentence
exists the memo must state at least two of the three delivered totals.
"""
from __future__ import annotations

import csv
import glob
import os
import re

VERBS = r"(recommend(ed|ation)?|pick|go with|choose|chose|select(ed)?|winner|best value|best option|my (pick|choice)|should (buy|order|take)|award)"
NEG = r"(\bnot\b|\bn't\b|\bavoid\b|\brather than\b|\binstead of\b|\bdo not\b|\bwould not\b|\bdon't\b|\bunless\b)"


def _find(ws: str, name: str):
    hits = sorted(glob.glob(os.path.join(ws, name)))
    if not hits:
        hits = sorted(glob.glob(os.path.join(ws, "**", name), recursive=True))
    hits = [h for h in hits if os.path.isfile(h) and "/.proto" not in h and "/.codex" not in h]
    return hits[0] if hits else None


def _num(s: str):
    try:
        return float(re.sub(r"[^0-9.\-]", "", s))
    except Exception:
        return None


def _numbers(text: str) -> list[float]:
    out = []
    for m in re.finditer(r"\(?-?[$€£]?\s?\d[\d,]*(?:\.\d+)?\)?", text):
        v = _num(m.group(0))
        if v is not None:
            out.append(v)
    return out


def _has(nums: list[float], want: float, rel: float = 0.005) -> bool:
    return any(abs(v - want) <= max(abs(want) * rel, 0.01) for v in nums)


# Qualitative terms per vendor, seed 0 (see task.yaml traps). A cell under a vendor's header is wrong
# when it carries another vendor's distinctive terms and none of its own: a workbook that puts
# Northstar's 24-month depot warranty under Kestrel is wrong even when every total is right. A cell
# that matches nothing recognisable is left alone; this check is about misattribution, not wording.
QUAL = {
    "warranty": {
        "bluepine": r"\b36\b|3 ?years|three years|on-?site|next[- ]business[- ]day|\bnbd\b",
        "kestrel": r"\b12\b|1 ?year|one year|12-month|bring-?in|carry-?in|hamburg|return to kestrel",
        "northstar": r"\b24\b|2 ?years|two years|depot|mail-?in",
    },
    "lead": {
        "bluepine": r"\b8\b|eight weeks|built to order",
        "kestrel": r"\b3\s*[-–]\s*4\b|3\.5|three to four|\b3 weeks|\b4 weeks",
        "northstar": r"\b10 ?business|\b10\b(?=[^|]*day)|two weeks|\b2 weeks|1\.5|ships within 10",
    },
}
ROW_LABEL = {"warranty": r"warrant", "lead": r"lead[\s-]*time|delivery time|ships? (in|within)|\blead\b"}


def _qualitative_rows(ws: str) -> dict:
    """Every warranty and lead-time row under a vendor header, in every sheet: {(sheet,row,kind): {vendor: [text]}}."""
    import openpyxl  # grader env
    p = _find(ws, "comparison.xlsx")
    if not p:
        return {}
    wb = openpyxl.load_workbook(p, data_only=True)
    found: dict = {}
    for sh in wb.worksheets:
        rows = [[c.value for c in r] for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 80), max_col=min(sh.max_column, 16))]
        header = None
        for i, r in enumerate(rows):
            # A header row names one vendor per cell in distinct columns; a sentence that mentions
            # all three vendors in one cell (a note, a source list) is not a header.
            cols = {}
            for j, v in enumerate(r):
                low = str(v or "").lower()
                named = [vend for vend in QUAL["warranty"] if vend in low]
                if j >= 1 and len(named) == 1 and len(low) <= 60:
                    cols.setdefault(named[0], []).append(j)
            if len(cols) >= 2:
                header = cols
                continue
            if not header:
                continue
            label = " ".join(str(v) for v in r[:2] if v is not None).lower()
            for kind, pat in ROW_LABEL.items():
                if re.search(pat, label) and not (kind == "lead" and "warrant" in label):
                    # Table cells are short; a paragraph in a notes column is prose, not a vendor cell.
                    cells = {}
                    for vend, js in header.items():
                        for j in js:
                            if j < len(r) and r[j] not in (None, "") and len(str(r[j])) <= 90:
                                cells.setdefault(vend, []).append(str(r[j]))
                    if len(cells) >= 2:
                        found[(sh.title, i + 1, kind)] = cells
    return found


def check_qualitative(ws: str) -> dict:
    name = "workbook attributes warranty and lead time to the right vendor"
    try:
        rows = _qualitative_rows(ws)
    except Exception as e:  # a workbook openpyxl cannot read fails the xlsx checks already
        return {"name": name, "passed": True, "detail": f"not checked: {e}"}
    if not rows:
        return {"name": name, "passed": True, "detail": "no warranty or lead-time rows under vendor headers; nothing to misattribute"}
    bad = []
    checked = 0
    NEG = re.compile(r"\b(no|not|none|without|excludes?|excluding|lacks?|never)\b[^.;,]{0,25}$", re.I)
    def _matches(pattern: str, text: str) -> bool:
        # A term preceded by a negation within the clause is not that vendor's term: "no on-site service".
        for m in re.finditer(pattern, text, re.I):
            if not NEG.search(text[max(0, m.start() - 30):m.start()]):
                return True
        return False
    for (sheet, rownum, kind), cells in rows.items():
        for vend, texts in cells.items():
            for t in texts:
                low = t.lower()
                own = _matches(QUAL[kind][vend], low)
                others = [o for o in QUAL[kind] if o != vend and _matches(QUAL[kind][o], low)]
                if own or others:
                    checked += 1
                if not own and others:
                    bad.append(f"{sheet}!row {rownum} {kind}: the {vend} column says {t[:50]!r}, which is {others[0]}'s")
    return {"name": name, "passed": not bad, "detail": "; ".join(bad[:4]) if bad else f"{checked} vendor cell(s) in {len(rows)} qualitative row(s) carry their own vendor's terms"}


def check(ws: str, ref: str) -> list[dict]:
    return [check_memo(ws, ref), check_qualitative(ws)]


def check_memo(ws: str, ref: str) -> dict:
    name = "memo states the delivered total of the vendor it recommends"
    p = _find(ws, "recommendation.md")
    if not p:
        return {"name": name, "passed": False, "detail": "no recommendation.md in workspace"}
    text = open(p, encoding="utf-8", errors="replace").read()
    totals: dict[str, float] = {}
    with open(os.path.join(ref, "totals.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            totals[r["vendor"].split()[0].lower()] = float(r["total_usd"])  # northstar, kestrel, bluepine
    vendors = list(totals)
    sentences = [x.strip() for x in re.split(r"(?<=[.!?])\s+|\n+", text) if x.strip()]
    votes: dict[str, int] = {}
    for sent in sentences:
        low = sent.lower()
        m = re.search(VERBS, low)
        if not m:
            continue
        before = low[: m.start()]
        if re.search(NEG, before[-40:]):
            continue
        after = low[m.end():]
        first = None
        for v in vendors:
            i = after.find(v)
            if i >= 0 and (first is None or i < first[1]):
                first = (v, i)
        if first and first[1] <= 80:
            votes[first[0]] = votes.get(first[0], 0) + 1
    nums = _numbers(text)
    if votes:
        pick = max(votes, key=votes.get)
        ok = _has(nums, totals[pick])
        return {"name": name, "passed": ok,
                "detail": f"recommends {pick} (votes {votes}); its total {totals[pick]:,.2f} "
                          f"{'stated' if ok else 'NOT stated'} in the memo"}
    present = [v for v in vendors if _has(nums, totals[v])]
    return {"name": name, "passed": len(present) >= 2,
            "detail": f"no recommendation sentence found; totals stated for {present} (need 2 of 3)"}
